"""Trucks bulk import pipeline (no HTTP): parsers, template, export,
preview, commit."""
import io
import uuid

import openpyxl
import pytest
from sqlalchemy import func, select

from serversherpa.db.models import (
    AuditLog, Container, Initiative, Site, Truck, TruckContainer,
)
from serversherpa.trucks import bulk_import as bi


async def preview(db, rows):
    return await bi.preview_rows(db, bi.number_json_rows(rows))


async def one(db, row):
    return (await preview(db, [row]))["rows"][0]


async def commit(db, actor, rows, approved=(), source="test.csv"):
    return await bi.commit_rows(db, actor.id, bi.number_json_rows(rows),
                                approved_updates=set(approved), source_label=source)


async def mk_truck(db, name, **fields):
    containers = fields.pop("containers", [])
    truck = Truck(name=name, **fields)
    db.add(truck)
    await db.flush()
    for c in containers:
        db.add(TruckContainer(truck_id=truck.id, container_id=c.id))
    await db.commit()
    return truck


async def mk_initiative(db, name):
    """A minimal initiative row. The model's NOT-NULL-without-default
    columns are `name` and `initiative_type` (see
    tests/test_initiative_assets_api.py::_move, which posts
    initiative_type="move"); the brief's `kind=` guess doesn't match the
    real column name, so this uses `initiative_type=` instead."""
    init = Initiative(name=name, initiative_type="move")
    db.add(init)
    await db.commit()
    return init


# ── shape / parsers / template ──────────────────────────────────────

def test_columns_match_canonical_shape():
    assert bi.COLUMNS == [
        "name", "status", "driver_name", "co_driver_name", "team_drive",
        "contact_info", "load_number", "seal_id", "tracking_type",
        "tracking_update_type", "tracker_id", "initiative", "start_site",
        "end_site", "containers"]


def test_parse_bool_and_split_names():
    assert bi.parse_bool("Yes") is True and bi.parse_bool("TRUE") is True
    assert bi.parse_bool("1") is True and bi.parse_bool("y") is True
    assert bi.parse_bool("no") is False and bi.parse_bool("0") is False
    assert bi.parse_bool("maybe") is None and bi.parse_bool("") is None
    assert bi.split_names(" Crate A ; Crate B;;") == ["Crate A", "Crate B"]
    assert bi.split_names("") == []


def test_csv_and_json_normalize_identically():
    from_csv = bi.parse_upload("t.csv", bi.build_template_csv().encode())
    from_json = bi.number_json_rows(bi.SAMPLE_ROWS)
    assert [r for _, r in from_csv] == [r for _, r in from_json]
    assert [n for n, _ in from_csv] == [2, 3]


def test_xlsx_template_round_trips_with_reference_blocks():
    blob = bi.build_template_xlsx(["created", "active"], ["Move A"], ["DC-East"])
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    assert wb.sheetnames == ["Trucks", "Reference"]
    ref = [row[0].value for row in wb["Reference"].iter_rows()]
    assert ref == ["Valid statuses", "created", "active", None,
                   "Initiative names", "Move A", None, "Site names", "DC-East"]
    rows = bi.parse_upload("t.xlsx", blob)
    assert [r for _, r in rows] == [r for _, r in bi.number_json_rows(bi.SAMPLE_ROWS)]


# ── export ──────────────────────────────────────────────────────────

async def test_export_rows_shape_and_round_trip(db, seeded_user):
    site_a, site_b = Site(name="DC-East"), Site(name="DC-West")
    crate_a, crate_b = Container(name="Crate A"), Container(name="Crate B")
    db.add_all([site_a, site_b, crate_a, crate_b])
    await db.flush()
    move = await mk_initiative(db, "Move A")
    await mk_truck(db, "Zulu", status="in_transit", driver_name="Marcus", team_drive=True,
                   contact_info="+1 555", load_number="L-1", seal_id="S-1",
                   tracking_type={"type": "gps", "update_type": "API", "tracker_id": "T-1"},
                   initiative_id=move.id, start_site_id=site_a.id, end_site_id=site_b.id,
                   containers=[crate_b, crate_a])
    await mk_truck(db, "Alpha")
    await mk_truck(db, "Gone", archived_at=func.now())
    rows = await bi.export_rows(db)
    assert [r["name"] for r in rows] == ["Alpha", "Zulu"]
    assert set(rows[0]) == set(bi.COLUMNS)
    assert rows[0]["status"] == "created" and rows[0]["team_drive"] == "no"
    z = rows[1]
    assert z["team_drive"] == "yes" and z["tracking_type"] == "gps"
    assert z["tracking_update_type"] == "API" and z["tracker_id"] == "T-1"
    assert z["initiative"] == "Move A" and z["start_site"] == "DC-East"
    assert z["end_site"] == "DC-West" and z["containers"] == "Crate A; Crate B"

    csv_text = bi.build_rows_csv(rows)
    out = await preview(db, [r for _, r in bi.parse_upload("e.csv", csv_text.encode())])
    assert [r["action"] for r in out["rows"]] == ["unchanged", "unchanged"]


async def test_reference_lists(db, seeded_user):
    db.add(Site(name="Bee Site"))
    db.add(Site(name="Ant Site"))
    await db.flush()
    await mk_initiative(db, "Move Z")
    statuses, initiatives, sites = await bi.reference_lists(db)
    assert statuses == ["created", "active", "in_transit", "at_destination",
                        "inactive", "historical"]
    assert initiatives == ["Move Z"]
    assert sites == ["Ant Site", "Bee Site"]


# ── preview: validation ─────────────────────────────────────────────

async def test_validation_errors(db, seeded_user):
    db.add(Site(name="Twin Site"))
    db.add(Site(name="twin site"))
    db.add(Container(name="Crate A"))
    await db.commit()
    out = await preview(db, [
        {"name": "", "driver_name": "x"},  # non-blank companion cell so the
                                            # shared core doesn't treat this
                                            # as a fully-blank skipped row
        {"name": "A", "status": "flying"},
        {"name": "B", "team_drive": "maybe"},
        {"name": "C", "seal_id": "S" * 25},
        {"name": "D", "initiative": "Nowhere"},
        {"name": "E", "start_site": "Twin Site"},
        {"name": "F", "end_site": "Nowhere"},
        {"name": "G", "containers": "Crate A; Crate Z"},
        {"name": "Dup"},
        {"name": "dup"},
    ])
    errs = {r["row"]: r["errors"] for r in out["rows"]}
    assert errs[1] == ["name is required"]
    assert errs[2] == ["unknown status 'flying'"]
    assert errs[3] == ["team_drive must be yes or no"]
    assert errs[4] == ["seal_id is longer than 24 characters"]
    assert errs[5] == ["unknown initiative 'Nowhere'"]
    assert errs[6] == ["ambiguous site 'Twin Site'"]
    assert errs[7] == ["unknown site 'Nowhere'"]
    assert errs[8] == ["unknown container 'Crate Z'"]
    assert errs[9] == errs[10] == ["duplicate name 'Dup' within the import"] or \
        errs[10] == ["duplicate name 'dup' within the import"]
    assert out["can_commit"] is False


async def test_create_row_normalizes_and_defaults(db, seeded_user):
    site = Site(name="DC-East")
    crate = Container(name="Crate A")
    db.add_all([site, crate])
    await db.flush()
    await mk_initiative(db, "Move A")
    out = await preview(db, [{
        "name": "  Truck 1 ", "team_drive": "YES", "initiative": "move a",
        "start_site": "dc-east", "containers": "crate a", "tracking_type": "gps"}])
    row = out["rows"][0]
    assert row["action"] == "create" and row["matched_by"] is None
    assert row["name"] == "Truck 1"
    assert row["data"]["status"] == "created" and row["data"]["team_drive"] is True
    assert row["data"]["initiative"] == "Move A" and row["data"]["start_site"] == "DC-East"
    assert row["data"]["containers"] == ["Crate A"]
    assert row["cells"]["status"] == "" and row["cells"]["initiative"] == "move a"
    assert out["can_commit"] is True


# ── preview: matching ───────────────────────────────────────────────

async def test_match_by_name_and_ambiguity(db, seeded_user):
    await mk_truck(db, "Truck 7", driver_name="Old Driver")
    await mk_truck(db, "Twin")
    await mk_truck(db, "twin")
    await mk_truck(db, "Retired", archived_at=func.now())
    out = await preview(db, [
        {"name": "truck 7", "driver_name": "New Driver"},
        {"name": "Twin", "driver_name": "x"},
        {"name": "Retired"},
    ])
    rows = out["rows"]
    assert rows[0]["action"] == "update" and rows[0]["matched_by"] == "name"
    assert rows[0]["matched_name"] == "Truck 7"
    assert rows[0]["diff"]["driver_name"] == {"old": "Old Driver", "new": "New Driver"}
    assert rows[0]["diff"]["name"] == {"old": "Truck 7", "new": "truck 7"}
    assert rows[1]["errors"] == ["multiple existing trucks named 'Twin'"]
    assert rows[2]["action"] == "create"          # archived trucks never match


async def test_two_rows_on_one_truck_are_errors(db, seeded_user):
    await mk_truck(db, "Truck 7")
    out = await preview(db, [
        {"name": "Truck 7", "driver_name": "A"},
        {"name": "TRUCK 7", "driver_name": "B"},
    ])
    # both rows collide on the in-upload name key before the target check
    assert all("within the import" in r["errors"][0] for r in out["rows"])


async def test_update_diff_every_column_kind(db, seeded_user):
    site_a, site_b = Site(name="DC-East"), Site(name="DC-West")
    crate_a, crate_b, crate_c = Container(name="Crate A"), Container(name="Crate B"), Container(name="Crate C")
    db.add_all([site_a, site_b, crate_a, crate_b, crate_c])
    await db.flush()
    move_a = await mk_initiative(db, "Move A")
    move_b = await mk_initiative(db, "Move B")
    await mk_truck(db, "Truck 1", status="created", team_drive=False, contact_info="",
                   tracking_type={"type": "gps", "tracker_id": "T-1"},
                   initiative_id=move_a.id, start_site_id=site_a.id,
                   containers=[crate_a, crate_b])
    row = await one(db, {
        "name": "Truck 1", "status": "in_transit", "team_drive": "yes",
        "contact_info": "", "load_number": "L-9", "tracking_type": "cell",
        "tracking_update_type": "manual", "tracker_id": "T-1",
        "initiative": "Move B", "start_site": "DC-East", "end_site": "DC-West",
        "containers": "Crate B; Crate C"})
    assert row["action"] == "update"
    assert row["diff"] == {
        "status": {"old": "created", "new": "in_transit"},
        "team_drive": {"old": False, "new": True},
        "load_number": {"old": None, "new": "L-9"},
        "tracking_type": {"old": "gps", "new": "cell"},
        "tracking_update_type": {"old": None, "new": "manual"},
        "initiative": {"old": "Move A", "new": "Move B"},
        "end_site": {"old": None, "new": "DC-West"},
        "containers": {"add": ["Crate C"], "remove": ["Crate A"]},
    }


async def test_blank_cells_are_no_change_on_update(db, seeded_user):
    crate = Container(name="Crate A")
    db.add(crate)
    await db.flush()
    await mk_truck(db, "Truck 1", status="active", team_drive=True,
                   contact_info="call me", tracking_type={"type": "gps"},
                   containers=[crate])
    row = await one(db, {"name": "Truck 1"})
    assert row["action"] == "unchanged" and row["diff"] is None


# ── commit ──────────────────────────────────────────────────────────

async def test_commit_creates_truck_with_links_and_audit(db, seeded_user):
    site = Site(name="DC-East")
    crate = Container(name="Crate A")
    db.add_all([site, crate])
    await db.flush()
    move = await mk_initiative(db, "Move A")
    out = await commit(db, seeded_user, [{
        "name": "Truck 1", "status": "active", "team_drive": "yes",
        "driver_name": "Marcus", "tracking_type": "gps", "tracker_id": "T-1",
        "initiative": "Move A", "start_site": "DC-East", "containers": "Crate A"}],
        source="fleet.xlsx")
    assert (out["created"], out["updated"], out["skipped"], out["unchanged"]) == (1, 0, 0, 0)
    row = out["rows"][0]
    assert row["action"] == "created" and row["name"] == "Truck 1" and row["diff"] is None
    truck = await db.get(Truck, uuid.UUID(row["truck_id"]))
    assert truck.status == "active" and truck.team_drive is True
    assert truck.driver_name == "Marcus" and truck.contact_info == ""
    assert truck.tracking_type == {"type": "gps", "tracker_id": "T-1"}
    assert truck.initiative_id == move.id and truck.start_site_id == site.id
    assert truck.created_by == seeded_user.id
    assert set(await db.scalars(select(TruckContainer.container_id).where(
        TruckContainer.truck_id == truck.id))) == {crate.id}
    actions = sorted(await db.scalars(select(AuditLog.action).where(
        AuditLog.entity_type == "truck")))
    assert actions == ["bulk_import", "create"]
    bulk_row = await db.scalar(select(AuditLog).where(AuditLog.action == "bulk_import"))
    assert bulk_row.changes == {"created": 1, "updated": 0, "skipped": 0,
                                "unchanged": 0, "source": "fleet.xlsx"}


async def test_commit_updates_approved_skips_unapproved(db, seeded_user):
    crate_a, crate_b = Container(name="Crate A"), Container(name="Crate B")
    db.add_all([crate_a, crate_b])
    await db.flush()
    a = await mk_truck(db, "Truck A", tracking_type={"type": "gps", "tracker_id": "T-A"},
                       containers=[crate_a])
    b = await mk_truck(db, "Truck B", driver_name="Keep Me")
    await mk_truck(db, "Truck C")
    out = await commit(db, seeded_user, [
        {"name": "Truck A", "status": "in_transit", "tracking_type": "cell",
         "containers": "Crate B"},
        {"name": "Truck B", "driver_name": "Changed"},
        {"name": "Truck C"},
        {"name": "Truck D"},
    ], approved=[str(a.id)])
    assert (out["created"], out["updated"], out["skipped"], out["unchanged"]) == (1, 1, 1, 1)
    by_name = {r["name"]: r for r in out["rows"]}
    assert by_name["Truck A"]["action"] == "updated"
    assert by_name["Truck A"]["diff"]["containers"] == {"add": ["Crate B"], "remove": ["Crate A"]}
    assert by_name["Truck B"]["action"] == "skipped"
    assert by_name["Truck B"]["diff"] == {"driver_name": {"old": "Keep Me", "new": "Changed"}}
    assert by_name["Truck C"]["action"] == "unchanged"
    assert by_name["Truck D"]["action"] == "created"
    await db.refresh(a)
    await db.refresh(b)
    assert a.status == "in_transit"
    assert a.tracking_type == {"type": "cell", "tracker_id": "T-A"}      # merged, not replaced
    assert set(await db.scalars(select(TruckContainer.container_id).where(
        TruckContainer.truck_id == a.id))) == {crate_b.id}
    assert b.driver_name == "Keep Me"                                    # skipped row untouched
    update_audit = await db.scalar(select(AuditLog).where(
        AuditLog.action == "update", AuditLog.entity_type == "truck"))
    assert update_audit.entity_id == str(a.id)
    assert update_audit.changes["containers"] == {"add": ["Crate B"], "remove": ["Crate A"]}


async def test_commit_blank_cells_never_clear(db, seeded_user):
    t = await mk_truck(db, "Truck 1", status="active", team_drive=True,
                       contact_info="keep", tracking_type={"type": "gps"})
    out = await commit(db, seeded_user, [
        {"name": "Truck 1", "status": "", "team_drive": "", "contact_info": "",
         "tracking_type": "", "load_number": "L-1"}], approved=[str(t.id)])
    assert out["updated"] == 1
    await db.refresh(t)
    assert t.status == "active" and t.team_drive is True
    assert t.contact_info == "keep" and t.tracking_type == {"type": "gps"}
    assert t.load_number == "L-1"


async def test_commit_is_all_or_nothing(db, seeded_user):
    with pytest.raises(bi.BulkImportError) as exc:
        # the second row needs a non-blank companion cell so the shared
        # core doesn't treat it as a fully-blank skipped line (see the
        # same gotcha in test_validation_errors above)
        await commit(db, seeded_user, [
            {"name": "Good"}, {"name": "", "driver_name": "Bad"}])
    assert exc.value.code == "rows_invalid"
    assert [r["action"] for r in exc.value.extra["rows"]] == ["create", "error"]
    assert await db.scalar(select(func.count()).select_from(Truck)) == 0
    with pytest.raises(bi.BulkImportError):
        await commit(db, seeded_user, [])
