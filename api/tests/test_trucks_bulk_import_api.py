"""Trucks bulk-import endpoints: rank gating, formats, preview, commit."""
import io

import openpyxl
import pytest
from sqlalchemy import select

from serversherpa.db.models import (
    Container, PermissionOverride, Person, PersonRole, Truck, TruckContainer,
)
from serversherpa.trucks import bulk_import as bi
from tests.test_sites_api import login, make_login


@pytest.fixture
async def admin_hdrs(db, client):
    person = Person(first_name="Ada", last_name="Admin", email="ada@test.example.com")
    db.add(person)
    await db.flush()
    db.add(PersonRole(person_id=person.id, role="admin"))
    await db.commit()
    return await make_login(db, client, person, "ada@test.example.com")


async def test_staff_rank_forbidden_on_all_four(client, seeded_user):
    hdrs = await login(client)
    assert (await client.get("/trucks/bulk-import/template?format=csv", headers=hdrs)).status_code == 403
    assert (await client.get("/trucks/bulk-import/export?format=csv", headers=hdrs)).status_code == 403
    assert (await client.post("/trucks/bulk-import/preview", headers=hdrs,
                              json={"rows": [{"name": "X"}]})).status_code == 403
    assert (await client.post("/trucks/bulk-import/commit", headers=hdrs,
                              json={"rows": [{"name": "X"}]})).status_code == 403


async def test_template_and_export_formats(client, db, seeded_user, admin_hdrs):
    csv_resp = await client.get("/trucks/bulk-import/template?format=csv", headers=admin_hdrs)
    assert csv_resp.status_code == 200
    assert csv_resp.headers["content-disposition"] == 'attachment; filename="trucks-template.csv"'
    assert csv_resp.text.splitlines()[0] == ",".join(bi.COLUMNS)
    xlsx_resp = await client.get("/trucks/bulk-import/template?format=xlsx", headers=admin_hdrs)
    assert xlsx_resp.headers["content-disposition"] == 'attachment; filename="trucks-template.xlsx"'
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert wb.sheetnames == ["Trucks", "Reference"]
    ref_cells = [row[0].value for row in wb["Reference"].iter_rows()]
    assert "in_transit" in ref_cells and "Site names" in ref_cells
    assert (await client.get("/trucks/bulk-import/template?format=doc",
                             headers=admin_hdrs)).status_code == 422

    db.add(Truck(name="Exported", driver_name="Dee"))
    await db.commit()
    csv_resp = await client.get("/trucks/bulk-import/export?format=csv", headers=admin_hdrs)
    assert csv_resp.headers["content-disposition"] == 'attachment; filename="trucks-export.csv"'
    assert csv_resp.text.splitlines()[1].startswith("Exported,created,Dee,,no,")
    xlsx_resp = await client.get("/trucks/bulk-import/export?format=xlsx", headers=admin_hdrs)
    assert openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))["Trucks"]["A2"].value == "Exported"


async def test_preview_json_and_file_paths(client, db, seeded_user, admin_hdrs):
    json_resp = await client.post("/trucks/bulk-import/preview", headers=admin_hdrs,
                                  json={"rows": [{"name": "Jay"}]})
    assert json_resp.status_code == 200
    assert json_resp.json()["rows"][0]["action"] == "create"
    csv_bytes = b"name,status\nCee,flying\n"
    file_resp = await client.post("/trucks/bulk-import/preview", headers=admin_hdrs,
                                  files={"file": ("fleet.csv", csv_bytes, "text/csv")})
    row = file_resp.json()["rows"][0]
    assert row["row"] == 2 and row["errors"] == ["unknown status 'flying'"]
    bad = await client.post("/trucks/bulk-import/preview", headers=admin_hdrs,
                            json={"rows": [{"nope": 1}]})
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "unknown_columns"


async def test_commit_end_to_end_with_approved_and_skipped(client, db, seeded_user, admin_hdrs):
    crate = Container(name="Crate A")
    existing = Truck(name="Truck A")
    other = Truck(name="Truck B", driver_name="Keep")
    db.add_all([crate, existing, other])
    await db.commit()
    rows = [
        {"name": "Truck A", "status": "active", "containers": "Crate A"},
        {"name": "Truck B", "driver_name": "Changed"},
        {"name": "Truck C", "team_drive": "yes"},
    ]
    preview = (await client.post("/trucks/bulk-import/preview", headers=admin_hdrs,
                                 json={"rows": rows})).json()
    assert [r["action"] for r in preview["rows"]] == ["update", "update", "create"]
    resp = await client.post("/trucks/bulk-import/commit", headers=admin_hdrs, json={
        "rows": [r["cells"] for r in preview["rows"]],
        "approved_updates": [str(existing.id)], "source": "fleet.csv"})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert (body["created"], body["updated"], body["skipped"], body["unchanged"]) == (1, 1, 1, 0)
    assert [r["action"] for r in body["rows"]] == ["updated", "skipped", "created"]
    assert body["rows"][0]["diff"]["containers"] == {"add": ["Crate A"], "remove": []}
    await db.refresh(existing)
    await db.refresh(other)
    assert existing.status == "active" and other.driver_name == "Keep"
    assert set(await db.scalars(select(TruckContainer.container_id).where(
        TruckContainer.truck_id == existing.id))) == {crate.id}
    bad = await client.post("/trucks/bulk-import/commit", headers=admin_hdrs, json={
        "rows": [{"name": ""}], "approved_updates": []})
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "rows_invalid"


async def test_commit_also_requires_trucks_change(client, db, seeded_user, admin_hdrs):
    ada = await db.scalar(select(Person).where(Person.email == "ada@test.example.com"))
    db.add(PermissionOverride(person_id=ada.id, resource="trucks", action="change", allow=False))
    await db.commit()
    assert (await client.post("/trucks/bulk-import/preview", headers=admin_hdrs,
                              json={"rows": [{"name": "X"}]})).status_code == 200
    resp = await client.post("/trucks/bulk-import/commit", headers=admin_hdrs,
                             json={"rows": [{"name": "X"}], "approved_updates": []})
    assert resp.status_code == 403
    assert await db.scalar(select(Truck).where(Truck.name == "X")) is None
