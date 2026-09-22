"""Workers bulk-import endpoints: rank gating, template/export formats,
preview and commit through HTTP. seeded_user (staff, rank 40) holds
workers:add but is below the bulk bar; admin (60) clears it."""
import io

import openpyxl
import pytest
from sqlalchemy import select

from serversherpa.db.models import (
    PermissionOverride, Person, PersonRole, UserAccount, WorkerProfile,
)
from serversherpa.people import bulk_import as bi
from tests.test_sites_api import login, make_login


@pytest.fixture
async def admin_hdrs(db, client):
    person = Person(first_name="Ada", last_name="Admin", email="ada@test.example.com")
    db.add(person)
    await db.flush()
    db.add(PersonRole(person_id=person.id, role="admin"))
    await db.commit()
    return await make_login(db, client, person, "ada@test.example.com")


async def test_staff_rank_forbidden_on_all_four(client, seeded_user):
    hdrs = await login(client)
    assert (await client.get("/workers/bulk-import/template?format=csv",
                             headers=hdrs)).status_code == 403
    assert (await client.get("/workers/bulk-import/export?format=csv",
                             headers=hdrs)).status_code == 403
    assert (await client.post("/workers/bulk-import/preview", headers=hdrs,
                              json={"rows": [{"first_name": "X"}]})).status_code == 403
    assert (await client.post("/workers/bulk-import/commit", headers=hdrs,
                              json={"rows": [{"first_name": "X"}]})).status_code == 403


async def test_template_formats(client, db, seeded_user, admin_hdrs):
    csv_resp = await client.get("/workers/bulk-import/template?format=csv", headers=admin_hdrs)
    assert csv_resp.status_code == 200
    assert csv_resp.headers["content-type"].startswith("text/csv")
    assert csv_resp.headers["content-disposition"] == 'attachment; filename="workers-template.csv"'
    assert csv_resp.text.splitlines()[0] == ",".join(bi.COLUMNS)

    xlsx_resp = await client.get("/workers/bulk-import/template?format=xlsx", headers=admin_hdrs)
    assert xlsx_resp.status_code == 200
    assert xlsx_resp.headers["content-disposition"] == (
        'attachment; filename="workers-template.xlsx"')
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert wb.sheetnames == ["Workers", "Reference"]
    ref_cells = [row[0].value for row in wb["Reference"].iter_rows()]
    assert "L1" in ref_cells and "blacklist" in ref_cells and "Partner names" in ref_cells

    assert (await client.get("/workers/bulk-import/template?format=doc",
                             headers=admin_hdrs)).status_code == 422


async def test_export_formats(client, db, seeded_user, admin_hdrs):
    person = Person(first_name="Exported", last_name="Worker", phone="555-000-9999")
    db.add(person)
    await db.flush()
    db.add(PersonRole(person_id=person.id, role="worker"))
    await db.commit()
    csv_resp = await client.get("/workers/bulk-import/export?format=csv", headers=admin_hdrs)
    assert csv_resp.status_code == 200
    assert csv_resp.headers["content-disposition"] == 'attachment; filename="workers-export.csv"'
    lines = csv_resp.text.splitlines()
    assert lines[0] == ",".join(bi.COLUMNS)
    assert lines[1].startswith("Exported,Worker,,,555-000-9999")
    xlsx_resp = await client.get("/workers/bulk-import/export?format=xlsx", headers=admin_hdrs)
    assert xlsx_resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert wb["Workers"]["A2"].value == "Exported"
    assert (await client.get("/workers/bulk-import/export?format=doc",
                             headers=admin_hdrs)).status_code == 422


async def test_preview_json_and_file_paths(client, db, seeded_user, admin_hdrs):
    json_resp = await client.post("/workers/bulk-import/preview", headers=admin_hdrs,
                                  json={"rows": [{"first_name": "Jay", "last_name": "Son"}]})
    assert json_resp.status_code == 200
    assert json_resp.json()["rows"][0]["action"] == "create"
    assert json_resp.json()["rows"][0]["row"] == 1

    csv_bytes = b"first_name,last_name,level\nCee,Ess,L9\n"
    file_resp = await client.post("/workers/bulk-import/preview", headers=admin_hdrs,
                                  files={"file": ("crew.csv", csv_bytes, "text/csv")})
    assert file_resp.status_code == 200
    row = file_resp.json()["rows"][0]
    assert row["row"] == 2 and row["errors"] == ["unknown level 'L9'"]

    bad = await client.post("/workers/bulk-import/preview", headers=admin_hdrs,
                            json={"rows": [{"nope": 1}]})
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "unknown_columns"
    missing = await client.post("/workers/bulk-import/preview", headers=admin_hdrs,
                                files={"other": ("x.csv", b"a", "text/csv")})
    assert missing.status_code == 422 and missing.json()["detail"]["code"] == "missing_file"


async def test_commit_also_requires_workers_change(client, db, seeded_user, admin_hdrs):
    """Template/export/preview ride on workers:add; the commit edits people
    who already exist, so it must hold workers:change as well. The admin's
    grant is removed with a deny override — no seeded role splits the two."""
    admin = await db.scalar(select(Person).where(Person.email == "ada@test.example.com"))
    db.add(PermissionOverride(person_id=admin.id, resource="workers",
                              action="change", allow=False))
    await db.commit()
    rows = [{"first_name": "Jay", "last_name": "Son"}]
    ok = await client.post("/workers/bulk-import/preview", headers=admin_hdrs,
                           json={"rows": rows})
    assert ok.status_code == 200
    denied = await client.post("/workers/bulk-import/commit", headers=admin_hdrs,
                               json={"rows": rows, "approved_updates": []})
    assert denied.status_code == 403
    assert await db.scalar(select(Person).where(Person.last_name == "Son")) is None


async def test_commit_refuses_a_row_that_outranks_the_actor(client, db, seeded_user,
                                                            admin_hdrs):
    """A hand-rolled commit cannot slip past the preview's rank guard: the
    re-preview inside commit_rows refuses the whole payload."""
    boss = Person(first_name="Big", last_name="Boss", email="boss@test.example.com")
    db.add(boss)
    await db.flush()
    db.add(PersonRole(person_id=boss.id, role="worker"))
    db.add(PersonRole(person_id=boss.id, role="developer"))     # rank above admin's 60
    db.add(WorkerProfile(person_id=boss.id, status="active"))
    db.add(UserAccount(person_id=boss.id, email="boss.login@test.example.com",
                       password_hash="x"))
    await db.commit()

    resp = await client.post("/workers/bulk-import/commit", headers=admin_hdrs, json={
        "rows": [{"first_name": "Big", "last_name": "Boss", "status": "standby",
                  "city": "Reno"}],
        "approved_updates": [str(boss.id)], "source": "crew.csv"})
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert detail["code"] == "rows_invalid"
    assert detail["rows"][0]["errors"] == ["rank too low to edit this person"]
    profile = await db.get(WorkerProfile, boss.id)
    await db.refresh(profile)
    await db.refresh(boss)
    assert profile.status == "active" and boss.city is None


async def test_commit_end_to_end_with_approved_and_skipped(client, db, seeded_user, admin_hdrs):
    existing = Person(first_name="Robert", last_name="Smith", email="bob@test.example.com")
    other = Person(first_name="Sara", last_name="Jones", email="sara@test.example.com")
    db.add_all([existing, other])
    await db.flush()
    db.add(PersonRole(person_id=existing.id, role="worker"))
    db.add(PersonRole(person_id=other.id, role="worker"))
    await db.commit()

    rows = [
        {"first_name": "Robert", "last_name": "Smith", "email": "bob@test.example.com",
         "trade": "Cable"},
        {"first_name": "Sara", "last_name": "Jones", "city": "Austin"},
        {"first_name": "Maria", "last_name": "Lopez", "phone": "555-987-6543"},
    ]
    preview = (await client.post("/workers/bulk-import/preview", headers=admin_hdrs,
                                 json={"rows": rows})).json()
    assert [r["action"] for r in preview["rows"]] == ["update", "update", "create"]
    assert preview["rows"][0]["matched_by"] == "email, name"

    resp = await client.post("/workers/bulk-import/commit", headers=admin_hdrs, json={
        "rows": [r["cells"] for r in preview["rows"]],
        "approved_updates": [str(existing.id)], "source": "crew.csv"})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert (body["created"], body["updated"], body["skipped"], body["unchanged"]) == (1, 1, 1, 0)
    assert [r["action"] for r in body["rows"]] == ["updated", "skipped", "created"]
    assert body["rows"][0]["diff"] == {"trade": {"old": None, "new": "Cable"}}
    assert (await db.get(WorkerProfile, existing.id)).trade == "Cable"
    await db.refresh(other)
    assert other.city is None
    created = await db.scalar(select(Person).where(Person.last_name == "Lopez"))
    assert created.source_ref == "crew.csv"

    # a row error blocks the whole commit
    bad = await client.post("/workers/bulk-import/commit", headers=admin_hdrs, json={
        "rows": [{"first_name": "", "last_name": "Nope"}], "approved_updates": []})
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "rows_invalid"
    assert bad.json()["detail"]["rows"][0]["errors"] == ["first_name is required"]
