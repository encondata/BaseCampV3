"""Parsing + validation pipeline for sites bulk import (no HTTP)."""
import pytest

from serversherpa.sites import bulk_import as bi


def test_columns_match_canonical_shape():
    assert bi.COLUMNS == [
        "name", "code", "type", "status", "address_line1", "address_line2",
        "city", "region", "postal_code", "country", "latitude", "longitude",
        "timezone", "dc_provider", "partner", "clients", "notes"]


def test_csv_and_json_normalize_identically():
    csv_text = bi.build_template_csv()
    from_csv = bi.parse_upload("t.csv", csv_text.encode())
    from_json = bi.number_json_rows(bi.SAMPLE_ROWS)
    assert [r for _, r in from_csv] == [r for _, r in from_json]
    assert [n for n, _ in from_csv] == [2, 3]      # header is row 1
    assert [n for n, _ in from_json] == [1, 2]


def test_xlsx_template_round_trips():
    blob = bi.build_template_xlsx(["datacenter"], ["active"])
    rows = bi.parse_upload("t.xlsx", blob)
    assert [r for _, r in rows] == [r for _, r in bi.number_json_rows(bi.SAMPLE_ROWS)]


def test_csv_bom_and_numeric_cells():
    csv_text = "﻿name,postal_code\nBom DC,89501\n"
    rows = bi.parse_upload("t.csv", csv_text.encode())
    assert rows[0][1]["name"] == "Bom DC"
    assert rows[0][1]["postal_code"] == "89501"
    # xlsx numeric cell coerces without a trailing .0
    import io

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sites"
    ws.append(["name", "postal_code"])
    ws.append(["Num DC", 89501])
    buf = io.BytesIO()
    wb.save(buf)
    rows = bi.parse_upload("n.xlsx", buf.getvalue())
    assert rows[0][1]["postal_code"] == "89501"


def test_xlsx_blank_header_cell_does_not_shift_columns():
    import io

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sites"
    ws.append(["name", "", "city"])
    ws.append(["Blank DC", "spacer", "Reno"])
    buf = io.BytesIO()
    wb.save(buf)
    [(_, row)] = bi.parse_upload("b.xlsx", buf.getvalue())
    assert row["name"] == "Blank DC"
    assert row["city"] == "Reno"          # not shifted into the spacer


def test_unknown_column_rejected():
    with pytest.raises(bi.BulkImportError) as exc:
        bi.number_json_rows([{"name": "A", "citty": "Reno"}])
    assert exc.value.code == "unknown_columns"
    assert exc.value.extra["columns"] == ["citty"]


def test_too_many_rows_rejected():
    with pytest.raises(bi.BulkImportError) as exc:
        bi.number_json_rows([{"name": str(i)} for i in range(1001)])
    assert exc.value.code == "too_many_rows"


def test_bad_file_extension_and_broken_payloads():
    with pytest.raises(bi.BulkImportError) as exc:
        bi.parse_upload("notes.txt", b"hello")
    assert exc.value.code == "unsupported_file"
    with pytest.raises(bi.BulkImportError) as exc:
        bi.parse_upload("x.json", b"{not json")
    assert exc.value.code == "invalid_json"
    with pytest.raises(bi.BulkImportError) as exc:
        bi.number_json_rows("just a string")  # type: ignore[arg-type]
    assert exc.value.code == "invalid_json"
    with pytest.raises(bi.BulkImportError) as exc:
        bi.number_json_rows([{"name": "A"}, "not a row"])  # type: ignore[list-item]
    assert exc.value.code == "invalid_json"
    with pytest.raises(bi.BulkImportError) as exc:
        bi.parse_upload("big.csv", b"x" * (bi.MAX_BYTES + 1))
    assert exc.value.code == "file_too_large"


def test_single_object_wraps_into_one_row():
    """Pasting ONE site as a bare JSON object (no array brackets) must work —
    the array wrapper is a formality users will forget."""
    rows = bi.number_json_rows({"name": "HO1", "city": "Houston"})
    assert [n for n, _ in rows] == [1]
    assert rows[0][1]["name"] == "HO1"

    via_file = bi.parse_upload("x.json", b'{"name": "HO1", "city": "Houston"}')
    assert [r for _, r in via_file] == [r for _, r in rows]

    # pasting the API envelope by mistake gets the clearer column error
    with pytest.raises(bi.BulkImportError) as exc:
        bi.number_json_rows({"rows": []})
    assert exc.value.code == "unknown_columns"
    assert exc.value.extra["columns"] == ["rows"]


async def test_preview_missing_name_and_payload_dupes(db, seeded_user):
    # a fully blank row is dropped by the parser, so give row 1 some content
    rows = bi.number_json_rows([
        {"name": "", "city": "Reno"}, {"name": "Twin"}, {"name": "twin"}])
    out = await bi.preview_rows(db, rows)
    by_row = {r["row"]: r for r in out["rows"]}
    assert by_row[1]["action"] == "error" and "name" in by_row[1]["errors"][0]
    assert by_row[2]["action"] == "error"     # in-payload duplicate (both rows)
    assert by_row[3]["action"] == "error"
    assert out["can_commit"] is False


async def test_preview_validates_lookups_coords_orgs(db, seeded_user):
    rows = bi.number_json_rows([
        {"name": "A", "type": "spaceport"},
        {"name": "B", "status": "haunted"},
        {"name": "C", "latitude": "95", "longitude": "0"},
        {"name": "D", "latitude": "40"},
        {"name": "E", "partner": "Nobody"},
        {"name": "F", "clients": "Ghost Co"},
    ])
    out = await bi.preview_rows(db, rows)
    assert all(r["action"] == "error" for r in out["rows"])
    assert out["can_commit"] is False


async def test_preview_good_rows_normalize_defaults(db, seeded_user):
    rows = bi.number_json_rows([{"name": "  Fresh DC  ", "city": "Reno"}])
    out = await bi.preview_rows(db, rows)
    row = out["rows"][0]
    assert row["action"] == "create" and out["can_commit"] is True
    assert row["data"]["name"] == "Fresh DC"          # trimmed
    assert row["data"]["status"] == "active"          # default applied
    assert row["data"]["country"] == "US"
    assert "_blank" not in row["data"]


async def test_duplicate_name_is_an_update_with_diff(db, seeded_user):
    from serversherpa.db.models import Site
    db.add(Site(name="Exists", city="Old Town", country="US", status="active"))
    await db.commit()
    rows = bi.number_json_rows([{"name": "exists", "city": "New Town"}])
    out = await bi.preview_rows(db, rows)
    row = out["rows"][0]
    assert row["action"] == "update" and row["site_id"]
    assert row["matched_by"] == "name" and row["matched_name"] == "Exists"
    assert row["diff"]["city"] == {"old": "Old Town", "new": "New Town"}
    assert "country" not in row["diff"]               # blank = no change
    assert "status" not in row["diff"]


async def test_duplicate_with_no_changes_is_unchanged(db, seeded_user):
    from serversherpa.db.models import Site
    db.add(Site(name="Same", city="Reno", country="US", status="active"))
    await db.commit()
    rows = bi.number_json_rows([{"name": "Same", "city": "Reno"}])
    out = await bi.preview_rows(db, rows)
    assert out["rows"][0]["action"] == "unchanged"
    assert out["rows"][0]["diff"] is None
    assert out["can_commit"] is True


async def test_ambiguous_existing_name_is_error(db, seeded_user):
    from serversherpa.db.models import Site
    db.add(Site(name="Dup A", country="US", status="active"))
    db.add(Site(name="dup a", country="US", status="active"))
    await db.commit()
    rows = bi.number_json_rows([{"name": "Dup A", "city": "X"}])
    out = await bi.preview_rows(db, rows)
    assert out["rows"][0]["action"] == "error"
    assert "multiple existing sites" in out["rows"][0]["errors"][0]


async def test_clients_diff_add_remove_and_partner(db, seeded_user):
    from serversherpa.db.models import Client, Partner, Site, SiteClient
    old_co = Client(name="Old Co")
    new_co = Client(name="New Co")
    pt = Partner(name="Haul It")
    db.add_all([old_co, new_co, pt])
    await db.flush()
    site = Site(name="Linked", country="US", status="active")
    db.add(site)
    await db.flush()
    db.add(SiteClient(site_id=site.id, client_id=old_co.id))
    await db.commit()

    rows = bi.number_json_rows([
        {"name": "Linked", "clients": "New Co", "partner": "haul it"}])
    out = await bi.preview_rows(db, rows)
    row = out["rows"][0]
    assert row["action"] == "update"
    assert row["diff"]["clients"] == {"add": ["New Co"], "remove": ["Old Co"]}
    assert row["diff"]["partner"] == {"old": None, "new": "Haul It"}


def test_normalize_address():
    assert bi.normalize_address("  607 14th St. NW, Suite 660 ") == "607 14th st nw suite 660"
    assert bi.normalize_address("100 Server-Way") == "100 server way"
    assert bi.normalize_address("") == ""


async def test_address_match_renames_and_reports_matched_by(db, seeded_user):
    from serversherpa.db.models import Site
    db.add(Site(name="Old Name", address_line1="100 Server Way", country="US", status="active"))
    await db.commit()
    rows = bi.number_json_rows([{"name": "New Name", "address_line1": "100 server-way"}])
    out = await bi.preview_rows(db, rows)
    row = out["rows"][0]
    assert row["action"] == "update"
    assert row["matched_by"] == "address" and row["matched_name"] == "Old Name"
    assert row["diff"]["name"] == {"old": "Old Name", "new": "New Name"}
    # the address matched on its normalized key, but the cleanup itself is an
    # ordinary field change the approver gets to see
    assert row["diff"]["address_line1"] == {"old": "100 Server Way",
                                            "new": "100 server-way"}


async def test_new_row_reports_no_match(db, seeded_user):
    rows = bi.number_json_rows([{"name": "Brand New", "address_line1": "1 Nowhere Rd"}])
    out = await bi.preview_rows(db, rows)
    assert out["rows"][0]["action"] == "create"
    assert out["rows"][0]["matched_by"] is None and out["rows"][0]["matched_name"] is None


async def test_name_and_address_pointing_at_different_sites_is_error(db, seeded_user):
    from serversherpa.db.models import Site
    db.add(Site(name="Site A", address_line1="1 First St", country="US", status="active"))
    db.add(Site(name="Site B", address_line1="2 Second St", country="US", status="active"))
    await db.commit()
    rows = bi.number_json_rows([{"name": "Site A", "address_line1": "2 Second St"}])
    out = await bi.preview_rows(db, rows)
    assert out["rows"][0]["action"] == "error"
    assert out["rows"][0]["errors"] == ["name matches 'Site A' but address matches 'Site B'"]


async def test_ambiguous_address_is_error(db, seeded_user):
    from serversherpa.db.models import Site
    db.add(Site(name="Twin 1", address_line1="9 Same Ave", country="US", status="active"))
    db.add(Site(name="Twin 2", address_line1="9 same ave.", country="US", status="active"))
    await db.commit()
    rows = bi.number_json_rows([{"name": "Third", "address_line1": "9 Same Ave"}])
    out = await bi.preview_rows(db, rows)
    assert out["rows"][0]["action"] == "error"
    assert out["rows"][0]["errors"] == ["multiple existing sites at that address: Twin 1, Twin 2"]


async def test_duplicate_address_within_import_is_error(db, seeded_user):
    rows = bi.number_json_rows([{"name": "One", "address_line1": "5 Dup Ln"},
                                {"name": "Two", "address_line1": "5 dup ln"}])
    out = await bi.preview_rows(db, rows)
    assert all(r["action"] == "error" for r in out["rows"])
    assert all("duplicate address within the import" in r["errors"] for r in out["rows"])


async def test_archived_sites_are_not_matched(db, seeded_user):
    from datetime import UTC, datetime
    from serversherpa.db.models import Site
    db.add(Site(name="Gone", address_line1="7 Past Rd", country="US", status="active",
                archived_at=datetime.now(UTC)))
    await db.commit()
    rows = bi.number_json_rows([{"name": "Gone", "address_line1": "7 Past Rd"}])
    out = await bi.preview_rows(db, rows)
    assert out["rows"][0]["action"] == "create"


# ── commit_rows ─────────────────────────────────────────────────────

async def _count_sites(db):
    from sqlalchemy import func
    from serversherpa.db.models import Site
    return await db.scalar(select(func.count()).select_from(Site))


from sqlalchemy import select  # noqa: E402  (test-file convenience)


async def test_commit_creates_sites_links_and_audit(db, seeded_user):
    from serversherpa.db.models import AuditLog, Client, Partner, Site, SiteClient
    acme = Client(name="Acme Co")
    pt = Partner(name="Haul It")
    db.add_all([acme, pt])
    await db.commit()

    rows = bi.number_json_rows([
        {"name": "  BC One  ", "city": "Reno"},
        {"name": "BC Two", "clients": "Acme Co", "partner": "Haul It",
         "type": "datacenter"},
    ])
    out = await bi.commit_rows(db, seeded_user.id, rows,
                               approved_updates=set(), source_label="paste")
    assert (out["created"], out["updated"], out["unchanged"]) == (2, 0, 0)
    assert [r["action"] for r in out["rows"]] == ["created", "created"]
    assert all(r["site_id"] for r in out["rows"])

    one = await db.scalar(select(Site).where(Site.name == "BC One"))
    assert one.status == "active" and one.country == "US"
    two = await db.scalar(select(Site).where(Site.name == "BC Two"))
    assert two.partner_id == pt.id and two.site_type == "datacenter"
    link = await db.scalar(select(SiteClient).where(
        SiteClient.site_id == two.id, SiteClient.client_id == acme.id))
    assert link is not None

    creates = (await db.scalars(select(AuditLog).where(
        AuditLog.entity_type == "site", AuditLog.action == "create"))).all()
    assert len(creates) == 2
    summary = await db.scalar(select(AuditLog).where(
        AuditLog.entity_type == "site_bulk_import"))
    assert summary.changes["created"] == 2
    assert summary.changes["source"] == "paste"


async def test_commit_returns_per_row_results(db, seeded_user):
    from serversherpa.db.models import Site
    db.add(Site(name="Keep Me", city="Old", country="US", status="active"))
    await db.commit()
    rows = bi.number_json_rows([{"name": "Keep Me", "city": "New"},
                                {"name": "Fresh One"},
                                {"name": "Keep Me 2"}])
    keep = (await bi.preview_rows(db, rows))["rows"][0]["site_id"]
    out = await bi.commit_rows(db, seeded_user.id, rows, approved_updates={keep},
                               source_label="t.csv")
    assert (out["created"], out["updated"], out["unchanged"]) == (2, 1, 0)
    assert [r["action"] for r in out["rows"]] == ["updated", "created", "created"]
    assert out["rows"][0] == {"row": 1, "name": "Keep Me", "site_id": keep,
                              "action": "updated", "diff": {"city": {"old": "Old", "new": "New"}}}
    assert all(r["site_id"] for r in out["rows"])
    assert out["rows"][1]["diff"] is None


async def test_commit_all_or_nothing(db, seeded_user):
    from serversherpa.db.models import AuditLog
    before = await _count_sites(db)
    rows = bi.number_json_rows([
        {"name": "Good Row"}, {"name": "Bad Row", "type": "spaceport"}])
    with pytest.raises(bi.BulkImportError) as exc:
        await bi.commit_rows(db, seeded_user.id, rows,
                             approved_updates=set(), source_label="paste")
    assert exc.value.code == "rows_invalid"
    actions = {r["action"] for r in exc.value.extra["rows"]}
    assert "error" in actions
    assert await _count_sites(db) == before
    assert await db.scalar(select(AuditLog.id).where(
        AuditLog.entity_type == "site_bulk_import")) is None


async def test_commit_update_requires_approval(db, seeded_user):
    from serversherpa.db.models import AuditLog, Site
    site = Site(name="Approve Me", city="Old", country="US", status="active")
    db.add(site)
    await db.commit()
    rows = bi.number_json_rows([{"name": "Approve Me", "city": "New"}])

    with pytest.raises(bi.BulkImportError) as exc:
        await bi.commit_rows(db, seeded_user.id, rows,
                             approved_updates=set(), source_label="paste")
    assert exc.value.code == "rows_invalid"
    assert any("update not approved" in e
               for r in exc.value.extra["rows"] for e in r["errors"])

    out = await bi.commit_rows(db, seeded_user.id, rows,
                               approved_updates={str(site.id)},
                               source_label="paste")
    assert out["updated"] == 1
    await db.refresh(site)
    assert site.city == "New"
    upd = await db.scalar(select(AuditLog).where(
        AuditLog.entity_type == "site", AuditLog.action == "update"))
    assert upd is not None and upd.changes["city"]["to"] == "New"


async def test_commit_unchanged_rows_skipped(db, seeded_user):
    from serversherpa.db.models import AuditLog, Site
    db.add(Site(name="Static", city="Reno", country="US", status="active"))
    await db.commit()
    rows = bi.number_json_rows([{"name": "Static", "city": "Reno"}])
    out = await bi.commit_rows(db, seeded_user.id, rows,
                               approved_updates=set(), source_label="paste")
    assert (out["created"], out["updated"], out["unchanged"]) == (0, 0, 1)
    assert out["rows"] == [{"row": 1, "name": "Static", "site_id": str(
        (await db.scalar(select(Site).where(Site.name == "Static"))).id),
        "action": "unchanged", "diff": None}]
    assert await db.scalar(select(AuditLog.id).where(
        AuditLog.entity_type == "site", AuditLog.action == "update")) is None


async def test_commit_update_links_clients(db, seeded_user):
    from serversherpa.db.models import Client, Site, SiteClient
    old_co, new_co = Client(name="Old Co"), Client(name="New Co")
    db.add_all([old_co, new_co])
    await db.flush()
    site = Site(name="Relink", country="US", status="active")
    db.add(site)
    await db.flush()
    db.add(SiteClient(site_id=site.id, client_id=old_co.id))
    await db.commit()

    rows = bi.number_json_rows([{"name": "Relink", "clients": "New Co"}])
    out = await bi.commit_rows(db, seeded_user.id, rows,
                               approved_updates={str(site.id)},
                               source_label="paste")
    assert out["updated"] == 1
    linked = set(await db.scalars(select(SiteClient.client_id).where(
        SiteClient.site_id == site.id)))
    assert linked == {new_co.id}


async def test_export_rows_round_trip_as_unchanged(db, seeded_user):
    from serversherpa.db.models import Client, Partner, Site, SiteClient
    p = Partner(name="ColoCo")
    c = Client(name="Acme")
    db.add_all([p, c])
    await db.flush()
    s = Site(name="Export Me", code="EXP", site_type="datacenter", status="active",
             address_line1="1 Export Way", city="Reno", region="NV", postal_code="89501",
             country="US", latitude=39.5296, longitude=-119.8138,
             timezone="America/Los_Angeles", dc_provider="Switch", partner_id=p.id,
             notes="hi")
    db.add(s)
    await db.flush()
    db.add(SiteClient(site_id=s.id, client_id=c.id))
    await db.commit()
    rows = await bi.export_rows(db)
    mine = next(r for r in rows if r["name"] == "Export Me")
    assert list(mine) == bi.COLUMNS
    assert mine["type"] == "datacenter" and mine["partner"] == "ColoCo"
    assert mine["clients"] == "Acme" and mine["latitude"] == "39.5296"
    assert mine["address_line2"] == ""
    # export -> upload previews as unchanged
    out = await bi.preview_rows(db, bi.number_json_rows(rows))
    assert {r["action"] for r in out["rows"]} == {"unchanged"}
    # the csv/xlsx writers accept the same rows
    assert bi.build_rows_csv(rows).splitlines()[0] == ",".join(bi.COLUMNS)
    parsed = bi.parse_upload("e.xlsx", bi.build_rows_xlsx(rows, ["datacenter"], ["active"]))
    assert [r for _, r in parsed][0]["name"] == rows[0]["name"]


async def test_two_rows_matching_the_same_site_are_errors(db, seeded_user):
    """One by name, one by address — applying both would silently let the
    last row win, so neither is applied."""
    from serversherpa.db.models import Site
    db.add(Site(name="Alpha", address_line1="100 Main", country="US",
                status="active"))
    await db.commit()
    rows = bi.number_json_rows([{"name": "Alpha", "city": "Reno"},
                                {"name": "Beta", "address_line1": "100 Main"}])
    out = await bi.preview_rows(db, rows)
    assert [r["action"] for r in out["rows"]] == ["error", "error"]
    assert all("two rows match the same existing site 'Alpha'" in r["errors"]
               for r in out["rows"])
    assert out["can_commit"] is False


async def test_preview_cells_keep_the_uploaded_blanks(db, seeded_user):
    rows = bi.number_json_rows([{"name": "Cellophane", "city": "Reno"}])
    row = (await bi.preview_rows(db, rows))["rows"][0]
    assert row["cells"]["status"] == "" and row["cells"]["country"] == ""
    assert row["data"]["status"] == "active"      # defaults live in data only
    assert row["cells"]["clients"] == ""          # unsplit, template-shaped


async def test_commit_blank_status_country_never_written(db, seeded_user):
    """Replaying the preview's `cells` (what the portal sends) must not carry
    the create-only status/country defaults onto an existing site."""
    from serversherpa.db.models import Site
    site = Site(name="Planned One", city="Old", country="CH", status="planned")
    db.add(site)
    await db.commit()
    site_id = str(site.id)

    rows = bi.number_json_rows([{"name": "Planned One", "city": "New"}])
    preview = await bi.preview_rows(db, rows)
    row = preview["rows"][0]
    assert row["action"] == "update" and row["site_id"] == site_id
    assert set(row["diff"]) == {"city"}

    out = await bi.commit_rows(db, seeded_user.id,
                               bi.number_json_rows([row["cells"]]),
                               approved_updates={site_id},
                               source_label="upload")
    assert out["updated"] == 1
    await db.refresh(site)
    assert site.city == "New"
    assert site.status == "planned" and site.country == "CH"

    # and a name-only row stays a no-op through preview AND commit
    bare = await bi.preview_rows(db, bi.number_json_rows([{"name": "Planned One"}]))
    assert bare["rows"][0]["action"] == "unchanged"
    again = await bi.commit_rows(db, seeded_user.id,
                                 bi.number_json_rows([bare["rows"][0]["cells"]]),
                                 approved_updates=set(), source_label="upload")
    assert (again["created"], again["updated"], again["unchanged"]) == (0, 0, 1)
    await db.refresh(site)
    assert site.status == "planned" and site.country == "CH"


async def test_export_guards_formula_cells_and_round_trips(db, seeded_user):
    import io

    import openpyxl
    from serversherpa.db.models import Site
    evil = '=HYPERLINK("http://evil","x")'
    db.add(Site(name=evil, country="US", status="active",
                latitude=39.5296, longitude=-119.8138))
    await db.commit()
    rows = await bi.export_rows(db)
    mine = next(r for r in rows if r["name"] == evil)

    csv_text = bi.build_rows_csv(rows)
    # guarded, then csv-quoted (the inner quotes double) — never a bare "="
    assert '"\'=HYPERLINK(""http://evil"",""x"")"' in csv_text
    assert "-119.8138" in csv_text and "'-119.8138" not in csv_text

    blob = bi.build_rows_xlsx(rows, ["datacenter"], ["active"])
    ws = openpyxl.load_workbook(io.BytesIO(blob))["Sites"]
    cells = {c.value: c for line in ws.iter_rows() for c in line
             if isinstance(c.value, str)}
    assert cells[evil].data_type == "s"           # text, never a formula
    assert cells["-119.8138"].data_type == "s"

    # both guarded exports re-upload as the row they came from
    from_csv = bi.parse_upload("e.csv", csv_text.encode())
    assert dict(from_csv[0][1])["name"] == evil
    from_xlsx = bi.parse_upload("e.xlsx", blob)
    assert dict(from_xlsx[0][1])["name"] == evil
    out = await bi.preview_rows(db, from_csv)
    assert {r["action"] for r in out["rows"]} == {"unchanged"}
    assert mine["longitude"] == "-119.8138"


def test_coord_text_keeps_whole_degrees():
    assert bi._coord_text(40) == "40"
    assert bi._coord_text(39.5296) == "39.5296"
    assert bi._coord_text(-119.8) == "-119.8"
    assert bi._coord_text(None) == ""


async def test_export_rows_excludes_archived_sites(db, seeded_user):
    from datetime import UTC, datetime
    from serversherpa.db.models import Site
    db.add(Site(name="Live Export", country="US", status="active"))
    db.add(Site(name="Archived Export", country="US", status="active",
                archived_at=datetime.now(UTC)))
    await db.commit()
    names = {r["name"] for r in await bi.export_rows(db)}
    assert "Live Export" in names and "Archived Export" not in names
