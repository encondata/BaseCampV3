"""Move Scan History options validation + build()'s run/definition option
resolution. New file per the Task 1 review — Task 2 owns `__init__.py`
itself (pdf branch, preview endpoint, routes/schemas), so this only
imports and exercises its already-landed xlsx-path contract: the module
under test is unmodified by this file."""

import re
from datetime import UTC, datetime, timedelta
from io import BytesIO

import openpyxl
import pytest

from serversherpa.db.models import (
    Asset, Client, Initiative, InitiativeAsset, Person, ReportDefinition, ReportRun,
)
from serversherpa.reports import move_scan_history
from serversherpa.reports.move_scan_history.gather import PIPELINE_STATUS_KEYS
from serversherpa.reports.registry import OptionsError


# ── validate_options (definition-level) ──────────────────────────────

def test_validate_options_normalizes_defaults_for_missing_keys():
    assert move_scan_history.validate_options({}) == {
        "default_format": "xlsx", "status_columns": "pipeline"}
    assert move_scan_history.validate_options({"default_format": "pdf"}) == {
        "default_format": "pdf", "status_columns": "pipeline"}


def test_validate_options_rejects_unknown_key():
    with pytest.raises(OptionsError):
        move_scan_history.validate_options({"bogus": 1})


def test_validate_options_rejects_bad_domain_values():
    with pytest.raises(OptionsError):
        move_scan_history.validate_options({"default_format": "csv"})
    with pytest.raises(OptionsError):
        move_scan_history.validate_options({"status_columns": "some"})


# ── validate_run_options (run-level) ─────────────────────────────────

def test_validate_run_options_returns_only_the_keys_given():
    # No default_format/format key injected when the run only overrides
    # status_columns — build() relies on distinguishing "the run didn't
    # mention this key" from "the run chose the default".
    assert move_scan_history.validate_run_options({"status_columns": "all"}) == {
        "status_columns": "all"}
    assert move_scan_history.validate_run_options({}) == {}
    assert move_scan_history.validate_run_options({"format": "pdf"}) == {"format": "pdf"}


def test_validate_run_options_rejects_unknown_key():
    with pytest.raises(OptionsError):
        move_scan_history.validate_run_options({"bogus": 1})


def test_validate_run_options_rejects_bad_format():
    with pytest.raises(OptionsError):
        move_scan_history.validate_run_options({"format": "csv"})


def test_validate_run_options_rejects_bad_status_columns():
    with pytest.raises(OptionsError):
        move_scan_history.validate_run_options({"status_columns": "some"})


# ── build(): option resolution + filename ────────────────────────────

async def _seed(db, *, definition_options: dict, run_options: dict):
    person = Person(first_name="Rae", last_name="Requester")
    client = Client(name="Acme")
    db.add_all([person, client])
    await db.flush()
    ini = Initiative(name="NAP11 Move", initiative_type="move", status="planned",
                     client_id=client.id)
    definition = ReportDefinition(name="Move Scan History", report_type="move_scan_history",
                                  options=definition_options, is_system=True)
    db.add_all([ini, definition])
    await db.flush()

    asset = Asset(legacy_id=9101, serial_number="SN-9101", name="Widget")
    db.add(asset)
    await db.flush()
    db.add(InitiativeAsset(initiative_id=ini.id, asset_id=asset.id))
    await db.flush()

    run = ReportRun(definition_id=definition.id, report_type="move_scan_history",
                    initiative_id=ini.id, options=run_options, requested_by=person.id,
                    requested_rank=0)
    db.add(run)
    await db.commit()
    return run


async def test_build_falls_back_to_definition_options_when_run_omits_them(db):
    run = await _seed(db, definition_options={"default_format": "xlsx",
                                              "status_columns": "pipeline"},
                      run_options={})
    result = await move_scan_history.build(db, run)
    assert result.filename.endswith(".xlsx")

    wb = openpyxl.load_workbook(BytesIO(result.content))
    ws = wb["Overview"]
    header_row = 11   # 9 block rows + blank separator
    headers = [ws.cell(row=header_row, column=c).value
              for c in range(4, 4 + len(PIPELINE_STATUS_KEYS) + 1)]
    # pipeline mode (from the definition): exactly the 14 pipeline columns
    assert headers[len(PIPELINE_STATUS_KEYS)] is None  # nothing past column 17


async def test_build_run_status_columns_overrides_definition_and_widens_header(db):
    run = await _seed(db, definition_options={"default_format": "xlsx",
                                              "status_columns": "pipeline"},
                      run_options={"status_columns": "all"})
    result = await move_scan_history.build(db, run)

    wb = openpyxl.load_workbook(BytesIO(result.content))
    ws = wb["Overview"]
    header_row = 11   # 9 block rows + blank separator
    # "Asset ID"/"Serial Number"/"Asset Name" + every status column
    header_count = 0
    col = 1
    while ws.cell(row=header_row, column=col).value is not None:
        header_count += 1
        col += 1
    # "all" mode = 14 pipeline columns + every other ACTIVE asset status
    # (the canonical seed has well over a dozen of those) — strictly more
    # than pipeline mode's fixed 3 + 14 = 17 columns.
    assert header_count > 3 + len(PIPELINE_STATUS_KEYS)


async def test_build_filename_carries_tz_local_date_hourminute_stamp(db):
    run = await _seed(db, definition_options={"default_format": "xlsx",
                                              "status_columns": "pipeline"},
                      run_options={})
    before = datetime.now(UTC).astimezone(move_scan_history.report_timezone())
    result = await move_scan_history.build(db, run)
    after = datetime.now(UTC).astimezone(move_scan_history.report_timezone())

    match = re.search(r" - (\d{4}-\d{2}-\d{2} \d{4})\.xlsx$", result.filename)
    assert match, result.filename
    stamped = datetime.strptime(match.group(1), "%Y-%m-%d %H%M").replace(
        tzinfo=move_scan_history.report_timezone())
    # tz-local, bracketed by "before" and "after" (both truncated to the
    # minute, since the stamp itself has minute resolution) — proves the
    # stamp is a real local-time conversion, not raw UTC or a fixed string
    assert before.replace(second=0, microsecond=0) <= stamped
    assert stamped <= after.replace(second=0, microsecond=0) + timedelta(minutes=1)
