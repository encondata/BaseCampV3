"""Fill-engine tests for the Site & Move Survey xlsx: placeholder
resolution, asset-row expansion (condensed and per-asset), the transport-
sheet toggle, and the real Champagne fixture end to end. Pure Python — no
DB. Ported behavior from V2's api/reports/site_move_survey.py
(BaseCampV2-reference, read-only); see
docs/superpowers/specs/2026-09-11-site-move-survey-design.md § Report
module and docs/site-move-survey/template-annotation-guide.md."""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

from serversherpa.reports.site_move_survey.assets import AssetRowInput, asset_rows
from serversherpa.reports.site_move_survey.fill import (
    PLACEHOLDER_RE, WHOLE_CELL_RE, expand_asset_rows, fill_workbook, resolve_path,
    substitute,
)

FIXTURE = Path(__file__).parent / "fixtures" / "champagne_annotated_template.xlsx"


def _champagne():
    return openpyxl.load_workbook(FIXTURE)


def _new_wb():
    wb = openpyxl.Workbook()
    return wb, wb.active


# ---------------------------------------------------------------------------
# regexes
# ---------------------------------------------------------------------------

def test_placeholder_re_finds_dotted_paths_ignoring_inner_whitespace():
    assert PLACEHOLDER_RE.findall("{{origin.city}}, {{ destination.state }}") == [
        "origin.city", "destination.state"]


def test_whole_cell_re_matches_only_a_single_full_cell_placeholder():
    assert WHOLE_CELL_RE.match("{{move.asset_count}}")
    assert WHOLE_CELL_RE.match("  {{ move.asset_count }}  ")
    assert WHOLE_CELL_RE.match("prefix {{move.asset_count}}") is None


# ---------------------------------------------------------------------------
# resolve_path / substitute
# ---------------------------------------------------------------------------

def test_resolve_path_missing_key_returns_empty_string():
    assert resolve_path({"origin": {}}, "origin.city") == ""


def test_resolve_path_none_value_returns_empty_string():
    assert resolve_path({"origin": {"city": None}}, "origin.city") == ""


def test_resolve_path_non_dict_node_returns_empty_string():
    assert resolve_path({"origin": "not a dict"}, "origin.city") == ""


def test_resolve_path_strips_whitespace_around_each_part():
    assert resolve_path({"origin": {"city": "Seattle"}}, " origin . city ") == "Seattle"


def test_substitute_non_string_value_passes_through_unchanged():
    assert substitute(42, {}) == 42
    assert substitute(None, {}) is None


def test_substitute_string_without_placeholder_passes_through():
    assert substitute("plain text", {}) == "plain text"


def test_substitute_whole_cell_preserves_raw_type_int():
    ctx = {"move": {"asset_count": 4}}
    result = substitute("{{move.asset_count}}", ctx)
    assert result == 4 and isinstance(result, int)


def test_substitute_whole_cell_none_becomes_empty_string_not_the_word_none():
    assert substitute("{{origin.city}}", {"origin": {"city": None}}) == ""


def test_substitute_mixed_text_renders_as_string_with_missing_literal():
    # Both origin.city and origin.state are missing/empty — a mixed cell
    # keeps its literal text, rendering ", " rather than disappearing.
    assert substitute("{{origin.city}}, {{origin.state}}", {"origin": {}}) == ", "


def test_substitute_mixed_text_fills_present_values():
    ctx = {"origin": {"city": "Seattle", "state": "WA"}}
    assert substitute("{{origin.city}}, {{origin.state}}", ctx) == "Seattle, WA"


# ---------------------------------------------------------------------------
# fill_workbook — merged cells, plain substitution
# ---------------------------------------------------------------------------

def test_fill_workbook_skips_merged_cells_without_erroring():
    wb, ws = _new_wb()
    ws["A1"] = "{{origin.name}}"
    ws["B1"] = "{{origin.name}}"
    ws.merge_cells("A1:B1")

    fill_workbook(wb, {"origin": {"name": "Datacenter West"}}, [], None,
                  include_transportation_standards=True)

    assert ws["A1"].value == "Datacenter West"
    assert ws["B1"].value is None   # the merged (non-anchor) cell is never touched


def test_fill_workbook_substitutes_every_placeholder_cell():
    wb, ws = _new_wb()
    ws["A1"] = "{{partner.name}}"
    ws["A2"] = "no placeholder here"
    ws["A3"] = "{{move.asset_count}}"

    ctx = {"partner": {"name": "Champagne Logistics"}, "move": {"asset_count": 3}}
    fill_workbook(wb, ctx, [], None, include_transportation_standards=True)

    assert ws["A1"].value == "Champagne Logistics"
    assert ws["A2"].value == "no placeholder here"
    assert ws["A3"].value == 3


# ---------------------------------------------------------------------------
# transport sheet toggle
# ---------------------------------------------------------------------------

def test_fill_workbook_removes_transport_sheet_when_toggle_off():
    wb, ws = _new_wb()
    ws.title = "Main"
    wb.create_sheet("Transportation Standards")

    fill_workbook(wb, {}, [], None, include_transportation_standards=False)

    assert "Transportation Standards" not in wb.sheetnames
    assert "Main" in wb.sheetnames


def test_fill_workbook_keeps_transport_sheet_when_toggle_on():
    wb, ws = _new_wb()
    ws.title = "Main"
    wb.create_sheet("Transportation Standards")

    fill_workbook(wb, {}, [], None, include_transportation_standards=True)

    assert "Transportation Standards" in wb.sheetnames


def test_fill_workbook_never_removes_the_only_sheet():
    wb, ws = _new_wb()
    ws.title = "Transportation Standards"

    fill_workbook(wb, {}, [], None, include_transportation_standards=False)

    assert wb.sheetnames == ["Transportation Standards"]


# ---------------------------------------------------------------------------
# expand_asset_rows — template row detection, style copy, trailing clear
# ---------------------------------------------------------------------------

def _template_row(ws, row=8):
    ws.cell(row=row, column=1).value = "{{asset.index}}"
    ws.cell(row=row, column=2).value = "{{asset.manufacturer}}"
    ws.cell(row=row, column=1).font = Font(bold=True, name="Trebuchet MS")
    ws.cell(row=row, column=1).border = Border(bottom=Side(style="thin"))
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
    return ws


def test_fill_workbook_detects_template_row_by_asset_placeholder():
    wb, ws = _new_wb()
    ws["A1"] = "not a template row"
    _template_row(ws, row=8)

    assets = [{"index": 1, "manufacturer": "Dell"}]
    fill_workbook(wb, {}, assets, None, include_transportation_standards=True)

    assert ws["A8"].value == 1
    assert ws["B8"].value == "Dell"
    assert ws["A1"].value == "not a template row"   # untouched by asset expansion


def test_expand_asset_rows_writes_one_row_per_asset_with_copied_style():
    wb, ws = _new_wb()
    _template_row(ws, row=8)

    assets = [{"index": 1, "manufacturer": "Dell"}, {"index": 2, "manufacturer": "HP"},
              {"index": 3, "manufacturer": "Cisco"}]
    expand_asset_rows(ws, 8, {}, assets, None)

    assert [ws.cell(row=r, column=2).value for r in (8, 9, 10)] == ["Dell", "HP", "Cisco"]
    assert [ws.cell(row=r, column=1).value for r in (8, 9, 10)] == [1, 2, 3]
    # Style copied onto the NEW rows (9, 10) from the template row (8).
    for r in (9, 10):
        cell = ws.cell(row=r, column=1)
        assert cell.font.bold is True
        assert cell.font.name == "Trebuchet MS"
        assert cell.alignment.horizontal == "center"


def test_expand_asset_rows_clears_trailing_rows_until_first_empty_row():
    wb, ws = _new_wb()
    _template_row(ws, row=8)
    ws.cell(row=9, column=1).value = "stale"
    ws.cell(row=9, column=2).value = "stale too"
    ws.cell(row=10, column=1).value = "also stale"
    # row 11 already blank in the template's columns — clearing stops there
    ws.cell(row=12, column=1).value = "beyond the stop row, never touched"

    expand_asset_rows(ws, 8, {}, [{"index": 1, "manufacturer": "Dell"}], None)

    assert ws.cell(row=8, column=2).value == "Dell"
    assert ws.cell(row=9, column=1).value is None
    assert ws.cell(row=9, column=2).value is None
    assert ws.cell(row=10, column=1).value is None
    assert ws.cell(row=12, column=1).value == "beyond the stop row, never touched"


def test_expand_asset_rows_zero_assets_writes_notes_into_last_placeholder_column():
    wb, ws = _new_wb()
    ws.cell(row=8, column=1).value = "{{asset.index}}"
    ws.cell(row=8, column=2).value = "{{asset.manufacturer}}"
    ws.cell(row=8, column=3).value = "{{asset.comments}}"   # last placeholder column

    expand_asset_rows(ws, 8, {}, [], "Equipment list will be provided separately")

    assert ws.cell(row=8, column=1).value == ""
    assert ws.cell(row=8, column=2).value == ""
    assert ws.cell(row=8, column=3).value == "Equipment list will be provided separately"


def test_expand_asset_rows_zero_assets_no_notes_leaves_placeholders_blank():
    wb, ws = _new_wb()
    ws.cell(row=8, column=1).value = "{{asset.index}}"
    ws.cell(row=8, column=2).value = "{{asset.comments}}"

    expand_asset_rows(ws, 8, {}, [], None)

    assert ws.cell(row=8, column=1).value == ""
    assert ws.cell(row=8, column=2).value == ""


# ---------------------------------------------------------------------------
# condensed vs per-asset rows (assets.py, fed straight into fill_workbook)
# ---------------------------------------------------------------------------

def _roster():
    return [
        AssetRowInput(make="Dell", model="R740", ru_size=2, weight=50, rack="R1",
                     ru_position=10, legacy_id=101, serial="SN1", name="web-01",
                     rfid="RF1", location="Cage A"),
        AssetRowInput(make="Dell", model="R740", ru_size=2, weight=50, rack="R1",
                     ru_position=12, legacy_id=102, serial="SN2", name="web-02",
                     rfid="RF2", location="Cage A"),
        AssetRowInput(make="Cisco", model="Nexus 9336C", ru_size=1, weight=20, rack="R2",
                     ru_position=20, legacy_id=103, serial="SN3", name="sw-01",
                     rfid="RF3", location="Cage B"),
    ]


def test_condensed_asset_rows_group_by_make_model_with_qty():
    rows = asset_rows(_roster(), condensed=True)

    assert len(rows) == 2   # (Dell, R740) and (Cisco, Nexus 9336C)
    dell = next(r for r in rows if r["make"] == "Dell")
    assert dell["qty"] == 2
    assert dell["model"] == "R740"
    assert dell["rack"] == ""          # a group can span racks — left blank
    assert dell["serial_number"] == ""


def test_per_asset_rows_one_row_per_asset_qty_always_one():
    rows = asset_rows(_roster(), condensed=False)

    assert len(rows) == 3
    assert all(r["qty"] == 1 for r in rows)
    assert rows[0]["serial_number"] == "SN1"
    assert rows[0]["rack"] == "R1 U10"    # source rack + RU combined
    assert rows[2]["make"] == "Cisco" and rows[2]["model"] == "Nexus 9336C"


def test_condensed_vs_per_asset_feed_straight_into_expand_asset_rows():
    roster = _roster()

    wb_c, ws_c = _new_wb()
    _template_row(ws_c, row=8)
    ws_c.cell(row=8, column=2).value = "{{asset.model}}"
    expand_asset_rows(ws_c, 8, {}, asset_rows(roster, condensed=True), None)
    assert [ws_c.cell(row=r, column=2).value for r in (8, 9)] == ["R740", "Nexus 9336C"]

    wb_p, ws_p = _new_wb()
    _template_row(ws_p, row=8)
    ws_p.cell(row=8, column=2).value = "{{asset.model}}"
    expand_asset_rows(ws_p, 8, {}, asset_rows(roster, condensed=False), None)
    assert [ws_p.cell(row=r, column=2).value for r in (8, 9, 10)] == [
        "R740", "R740", "Nexus 9336C"]


# ---------------------------------------------------------------------------
# the real Champagne fixture, end to end
# ---------------------------------------------------------------------------

def _champagne_context(**overrides):
    base = {
        "customer": {"company": "Cumulus Solutions Group", "contact_name": "Jimmy Henderson",
                    "address": "", "phone": "", "email": ""},
        "origin": {"name": "Datacenter West", "address": "300 Origin St",
                  "city": "Seattle", "state": "WA", "zip": "98101",
                  "contact_name": "", "contact_phone": "", "contact_email": "",
                  "survey": {}},
        "destination": {"name": "Datacenter East", "address": "400 Dest Ave",
                        "city": "Austin", "state": "TX", "zip": "78701",
                        "contact_name": "", "contact_phone": "", "contact_email": "",
                        "survey": {}},
        "move": {"scheduled_start_date": "2026-07-15", "scheduled_start_time": "08:30 AM"},
        "assets_notes": "", "asset_notes": "",
    }
    base.update(overrides)
    return base


def test_champagne_fixture_customer_and_site_information_cells():
    wb = _champagne()
    ctx = _champagne_context()

    fill_workbook(wb, ctx, [], None, include_transportation_standards=True)

    ws = wb["Customer and Site Information"]
    assert ws["C11"].value == "Cumulus Solutions Group"
    assert ws["C21"].value == "Seattle, WA"
    assert ws["C31"].value == "Austin, TX"


def test_champagne_fixture_equipment_listing_rows_with_copied_style_and_trailing_clear():
    import copy as copy_module

    wb = _champagne()
    ws = wb["Equipment Listing"]
    # `.border`/`.alignment`/`.font` return a StyleProxy that only compares
    # equal to another instance of the exact same (non-proxy) class, so
    # `copy.copy()` each side to unwrap before comparing — a well-known
    # openpyxl gotcha, not specific to this fill engine.
    original_font = copy_module.copy(ws["A8"].font)
    original_border = copy_module.copy(ws["A8"].border)
    original_alignment = copy_module.copy(ws["A8"].alignment)

    assets = asset_rows(_roster(), condensed=False)
    fill_workbook(wb, _champagne_context(), assets, None,
                  include_transportation_standards=True)

    n = len(assets)
    for i, row in enumerate(assets):
        r = 8 + i
        assert ws.cell(row=r, column=3).value == row["manufacturer"]
        assert ws.cell(row=r, column=4).value == row["model"]
        if r != 8:   # row 8 itself keeps its own (identical) original style
            cell = ws.cell(row=r, column=1)   # compare column A to column A
            assert copy_module.copy(cell.font) == original_font
            assert copy_module.copy(cell.border) == original_border
            assert copy_module.copy(cell.alignment) == original_alignment

    # The row right after the last written asset row is cleared.
    for col in range(1, 8):
        assert ws.cell(row=8 + n, column=col).value in (None, "")


def test_champagne_fixture_zero_assets_notes_land_in_last_placeholder_column_g8():
    wb = _champagne()
    ws = wb["Equipment Listing"]

    fill_workbook(wb, _champagne_context(), [], "Equipment list will be provided separately",
                  include_transportation_standards=True)

    assert ws["G8"].value == "Equipment list will be provided separately"
    assert ws["A8"].value == ""
    assert ws["C8"].value == ""


def test_champagne_fixture_has_no_transport_sheet_and_toggle_is_a_no_op():
    wb_on = _champagne()
    fill_workbook(wb_on, _champagne_context(), [], None, include_transportation_standards=True)
    wb_off = _champagne()
    fill_workbook(wb_off, _champagne_context(), [], None, include_transportation_standards=False)

    assert wb_on.sheetnames == wb_off.sheetnames == [
        "Customer and Site Information", "General Questions", "Equipment Listing"]
