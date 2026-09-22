"""Workers bulk import pipeline (no HTTP): keys, template, export, preview,
commit. Actor for preview/commit is an admin (rank 60) unless a test says
otherwise."""
import io
import uuid

import openpyxl
import pytest
from sqlalchemy import func, select

from serversherpa.db.models import (
    AuditLog, Partner, Person, PersonRole, UserAccount, WorkerProfile,
)
from serversherpa.people import bulk_import as bi

ADMIN_RANK = 60


@pytest.fixture
async def admin(db):
    person = Person(first_name="Ada", last_name="Admin", email="ada@test.example.com")
    db.add(person)
    await db.flush()
    db.add(PersonRole(person_id=person.id, role="admin"))
    await db.commit()
    return person


async def preview(db, admin, rows):
    return await bi.preview_rows(db, bi.number_json_rows(rows),
                                 actor_id=admin.id, actor_rank=ADMIN_RANK)


async def one(db, admin, row):
    """Preview a single row — matching scenarios must not share keys across rows."""
    return (await preview(db, admin, [row]))["rows"][0]


async def commit(db, admin, rows, approved=(), source="test.csv"):
    return await bi.commit_rows(db, bi.number_json_rows(rows), actor_id=admin.id,
                                actor_rank=ADMIN_RANK,
                                approved_updates=set(approved), source_label=source)


async def mk_worker(db, first="Robert", last="Smith", *, preferred=None,
                    email=None, phone=None, rfid=None, role=True, profile=None,
                    archived=False, account=False):
    from datetime import UTC, datetime
    person = Person(first_name=first, last_name=last, preferred_name=preferred,
                    email=email, phone=phone, rfid_tag=rfid,
                    archived_at=datetime.now(UTC) if archived else None)
    db.add(person)
    await db.flush()
    if role:
        db.add(PersonRole(person_id=person.id, role="worker"))
    if profile is not None:
        db.add(WorkerProfile(person_id=person.id, **profile))
    if account:
        db.add(UserAccount(person_id=person.id, email=email or f"{first}@test.example.com",
                           password_hash="x"))
    await db.commit()
    return person


# ── shape / keys / template ─────────────────────────────────────────

def test_columns_match_canonical_shape():
    assert bi.COLUMNS == [
        "first_name", "last_name", "preferred_name", "email", "phone",
        "job_title", "employee_number", "rfid_tag", "address_line1",
        "address_line2", "city", "region", "postal_code", "country",
        "partner", "trade", "level", "status", "status_note", "notes"]


def test_normalize_phone():
    assert bi.normalize_phone("(555) 123-4567") == "5551234567"
    assert bi.normalize_phone("1-555-123-4567") == "5551234567"
    assert bi.normalize_phone("+44 20 7946 0958") == "442079460958"   # not a US 1
    assert bi.normalize_phone("12345") == ""                            # too short
    assert bi.normalize_phone("") == ""


def test_name_keys_cover_first_and_preferred():
    assert bi.name_keys("Robert", "Smith", "Bob") == {"robert smith", "bob smith"}
    assert bi.name_keys("  Robert ", "SMITH", "") == {"robert smith"}
    assert bi.name_keys("", "Smith", "Bob") == {"bob smith"}
    assert bi.name_keys("", "", "") == set()


def test_csv_and_json_normalize_identically():
    from_csv = bi.parse_upload("t.csv", bi.build_template_csv().encode())
    from_json = bi.number_json_rows(bi.SAMPLE_ROWS)
    assert [r for _, r in from_csv] == [r for _, r in from_json]
    assert [n for n, _ in from_csv] == [2, 3]


def test_xlsx_template_round_trips_with_reference_blocks():
    blob = bi.build_template_xlsx(["L1", "L2"], ["active", "standby"], ["Haul It"])
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    assert wb.sheetnames == ["Workers", "Reference"]
    ref = [row[0].value for row in wb["Reference"].iter_rows()]
    assert ref == ["Valid levels", "L1", "L2", None,
                   "Valid statuses", "active", "standby", None,
                   "Partner names", "Haul It"]
    rows = bi.parse_upload("t.xlsx", blob)
    assert [r for _, r in rows] == [r for _, r in bi.number_json_rows(bi.SAMPLE_ROWS)]


# ── export ──────────────────────────────────────────────────────────

async def test_export_rows_shape_and_round_trip(db, admin):
    pt = Partner(name="Haul It")
    db.add(pt)
    await db.flush()
    await mk_worker(db, "Zed", "Zulu", email="zed@test.example.com", phone="555-000-1111",
                    profile={"partner_id": pt.id, "trade": "Cable", "level": "L2",
                             "status": "standby"})
    await mk_worker(db, "Amy", "Alpha")                         # no profile
    await mk_worker(db, "Not", "Worker", role=False)            # no worker role
    await mk_worker(db, "Old", "Gone", archived=True)
    rows = await bi.export_rows(db)
    assert [r["last_name"] for r in rows] == ["Alpha", "Zulu"]
    assert set(rows[0]) == set(bi.COLUMNS)
    assert rows[0]["status"] == "active"                # profile-less default
    assert rows[1]["partner"] == "Haul It" and rows[1]["level"] == "L2"
    assert rows[1]["phone"] == "555-000-1111"

    csv_text = bi.build_rows_csv(rows)
    out = await preview(db, admin, [r for _, r in bi.parse_upload("e.csv", csv_text.encode())])
    assert [r["action"] for r in out["rows"]] == ["unchanged", "unchanged"]


async def test_export_round_trips_with_two_workers_sharing_a_name(db, admin):
    """The promise that an export re-uploads clean is what the name rule buys:
    two Chris Lees round-trip because the export carries their emails."""
    await mk_worker(db, "Chris", "Lee", email="c1@test.example.com")
    await mk_worker(db, "Chris", "Lee", email="c2@test.example.com")
    await mk_worker(db, "Amy", "Alpha", phone="555-000-2222")
    csv_text = bi.build_rows_csv(await bi.export_rows(db))
    out = await preview(db, admin, [r for _, r in bi.parse_upload("e.csv", csv_text.encode())])
    assert [r["action"] for r in out["rows"]] == ["unchanged", "unchanged", "unchanged"]
    assert out["can_commit"] is True


async def test_reference_lists(db, admin):
    db.add(Partner(name="Bee Co"))
    db.add(Partner(name="Ant Co"))
    await db.commit()
    levels, statuses, partners = await bi.reference_lists(db)
    assert levels == ["L1", "L2", "L3", "L4", "L5", "L6"]
    assert statuses == ["active", "standby", "blacklist"]
    assert partners == ["Ant Co", "Bee Co"]


# ── preview: validation ─────────────────────────────────────────────

async def test_required_names_and_field_validation(db, admin):
    out = await preview(db, admin, [
        {"first_name": "", "last_name": "Solo"},
        {"first_name": "No", "last_name": ""},
        {"first_name": "Bad", "last_name": "Email", "email": "not-an-email"},
        {"first_name": "Short", "last_name": "Phone", "phone": "12345"},
        {"first_name": "Long", "last_name": "Country", "country": "USA"},
        {"first_name": "No", "last_name": "Partner", "partner": "Nobody"},
        {"first_name": "No", "last_name": "Level", "level": "L9"},
        {"first_name": "No", "last_name": "Status", "status": "haunted"},
        {"first_name": "No", "last_name": "Note", "status": "blacklist"},
    ])
    errs = {r["row"]: r["errors"] for r in out["rows"]}
    assert errs[1] == ["first_name is required"]
    assert errs[2] == ["last_name is required"]
    assert errs[3] == ["email 'not-an-email' is not valid"]
    assert errs[4] == ["phone needs at least 7 digits"]
    assert errs[5] == ["country must be a two-letter code"]
    assert errs[6] == ["unknown partner 'Nobody'"]
    assert errs[7] == ["unknown level 'L9'"]
    assert errs[8] == ["unknown status 'haunted'"]
    assert errs[9] == ["blacklist requires a status_note"]
    assert out["can_commit"] is False


async def test_create_row_normalizes_and_defaults(db, admin):
    db.add(Partner(name="Haul It"))
    await db.commit()
    out = await preview(db, admin, [{
        "first_name": " Robert ", "last_name": "Smith", "country": "us",
        "partner": "haul it", "level": "L3"}])
    row = out["rows"][0]
    assert row["action"] == "create" and row["matched_by"] is None
    assert row["name"] == "Robert Smith"
    assert row["data"]["first_name"] == "Robert"
    assert row["data"]["country"] == "US" and row["data"]["status"] == "active"
    assert row["data"]["partner"] == "Haul It"       # canonical partner name
    assert row["cells"]["country"] == "us" and row["cells"]["status"] == ""
    assert out["can_commit"] is True


async def test_duplicate_keys_within_the_upload(db, admin):
    out = await preview(db, admin, [
        {"first_name": "A", "last_name": "One", "email": "Dup@Example.com"},
        {"first_name": "B", "last_name": "Two", "email": "dup@example.com"},
        {"first_name": "C", "last_name": "Three", "phone": "555-111-2222"},
        {"first_name": "D", "last_name": "Four", "phone": "(555) 111 2222"},
        {"first_name": "Bob", "last_name": "Smith"},
        {"first_name": "Robert", "last_name": "Smith", "preferred_name": "Bob"},
        {"first_name": "E", "last_name": "Five", "rfid_tag": "TAG1"},
        {"first_name": "F", "last_name": "Six", "rfid_tag": "tag1"},
    ])
    errs = {r["row"]: r["errors"] for r in out["rows"]}
    assert errs[1] == errs[2] == ["duplicate email 'dup@example.com' within the import"]
    assert errs[3] == errs[4] == ["duplicate phone within the import"]
    assert errs[5] == ["duplicate name 'Bob Smith' within the import"]
    assert errs[6] == ["duplicate name 'Bob Smith' within the import"]
    assert errs[7] == errs[8] == ["duplicate rfid_tag 'tag1' within the import"]


# ── preview: matching ───────────────────────────────────────────────

async def test_match_by_each_key_alone(db, admin):
    await mk_worker(db, "Robert", "Smith", preferred="Bob",
                    email="bob@test.example.com", phone="555-123-4567")
    cases = [
        ({"first_name": "X", "last_name": "Y", "email": "BOB@test.example.com"}, "email", "update"),
        ({"first_name": "X", "last_name": "Y", "phone": "1 (555) 123-4567"}, "phone", "update"),
        ({"first_name": "Robert", "last_name": "Smith"}, "name", "unchanged"),
        ({"first_name": "Bob", "last_name": "Smith"}, "name", "update"),
    ]
    for row, key, action in cases:
        out = await one(db, admin, row)
        assert (out["matched_by"], out["action"], out["matched_name"]) == (key, action, "Bob Smith"), row
    bob = await one(db, admin, cases[3][0])
    assert bob["diff"]["first_name"] == {"old": "Robert", "new": "Bob"}


async def test_keys_that_agree_are_listed_and_keys_that_disagree_are_errors(db, admin):
    a = await mk_worker(db, "Robert", "Smith", email="a@test.example.com", phone="555-000-0001")
    await mk_worker(db, "Roberta", "Smith", email="b@test.example.com", phone="555-000-0002")
    good = await one(db, admin, {
        "first_name": "Robert", "last_name": "Smith", "email": "a@test.example.com",
        "phone": "555-000-0001"})
    assert good["action"] == "unchanged" and good["matched_by"] == "email, phone, name"
    assert good["person_id"] == str(a.id)
    bad = await one(db, admin, {
        "first_name": "Zed", "last_name": "Zulu", "email": "a@test.example.com",
        "phone": "555-000-0002"})
    assert bad["action"] == "error"
    assert bad["errors"] == ["email matches Robert Smith, phone matches Roberta Smith"]


async def test_ambiguous_name_is_an_error_only_without_a_stronger_key(db, admin):
    await mk_worker(db, "Chris", "Lee", email="c1@test.example.com")
    c2 = await mk_worker(db, "Chris", "Lee", email="c2@test.example.com")
    bare = await one(db, admin, {"first_name": "Chris", "last_name": "Lee"})
    assert bare["errors"] == ["two people share the name 'Chris Lee'"]
    # an email landing on exactly one of them decides the row: the shared
    # name drops out of the match instead of blocking it
    keyed = await one(db, admin, {"first_name": "Chris", "last_name": "Lee",
                                  "email": "c2@test.example.com"})
    assert keyed["errors"] == []
    assert (keyed["matched_by"], keyed["action"]) == ("email", "unchanged")
    assert keyed["person_id"] == str(c2.id)


async def test_two_rows_sharing_a_name_are_split_by_their_emails(db, admin):
    a = await mk_worker(db, "Chris", "Lee", email="c1@test.example.com")
    b = await mk_worker(db, "Chris", "Lee", email="c2@test.example.com")
    out = await preview(db, admin, [
        {"first_name": "Chris", "last_name": "Lee", "email": "c1@test.example.com"},
        {"first_name": "Chris", "last_name": "Lee", "email": "c2@test.example.com"},
    ])
    assert [r["action"] for r in out["rows"]] == ["unchanged", "unchanged"]
    assert [r["matched_by"] for r in out["rows"]] == ["email", "email"]
    assert [r["person_id"] for r in out["rows"]] == [str(a.id), str(b.id)]
    assert out["can_commit"] is True


async def test_two_rows_on_one_person_and_archived_never_match(db, admin):
    await mk_worker(db, "Robert", "Smith", email="bob@test.example.com")
    await mk_worker(db, "Old", "Timer", email="old@test.example.com", archived=True)
    out = await preview(db, admin, [
        {"first_name": "Robert", "last_name": "Smith"},
        {"first_name": "Zed", "last_name": "Zulu", "email": "bob@test.example.com"},
        {"first_name": "Old", "last_name": "Timer"},
        {"first_name": "New", "last_name": "Person", "email": "old@test.example.com"},
    ])
    rows = out["rows"]
    assert rows[0]["errors"] == ["two rows match the same existing person 'Robert Smith'"]
    assert rows[1]["errors"] == ["two rows match the same existing person 'Robert Smith'"]
    assert rows[2]["action"] == "create"
    assert rows[3]["errors"] == ["email 'old@test.example.com' belongs to an archived person"]


async def test_rfid_tag_collisions(db, admin):
    holder = await mk_worker(db, "Tag", "Holder", email="tag@test.example.com", rfid="ABC123")
    await mk_worker(db, "Gone", "Tag", rfid="OLD1", archived=True)
    other = await one(db, admin, {"first_name": "Other", "last_name": "Person", "rfid_tag": "abc123"})
    assert other["errors"] == ["rfid_tag 'abc123' belongs to Tag Holder"]
    own = await one(db, admin, {"first_name": "Tag", "last_name": "Holder",
                                "email": "tag@test.example.com", "rfid_tag": "ABC123"})
    assert own["action"] == "unchanged" and own["person_id"] == str(holder.id)
    stale = await one(db, admin, {"first_name": "Third", "last_name": "Person", "rfid_tag": "old1"})
    assert stale["errors"] == ["rfid_tag 'old1' belongs to an archived person"]


async def test_non_worker_user_matches_and_gets_the_role_in_the_diff(db, admin):
    user = await mk_worker(db, "Office", "User", email="ou@test.example.com", role=False)
    out = await preview(db, admin, [
        {"first_name": "Office", "last_name": "User", "trade": "Cable"}])
    row = out["rows"][0]
    assert row["action"] == "update" and row["person_id"] == str(user.id)
    assert row["diff"]["worker_role"] == {"old": None, "new": "granted"}
    assert row["diff"]["trade"] == {"old": None, "new": "Cable"}


async def test_update_diff_blank_means_no_change_and_phone_email_normalize(db, admin):
    pt = Partner(name="Haul It")
    db.add(pt)
    await db.flush()
    await mk_worker(db, "Robert", "Smith", email="Bob@test.example.com", phone="(555) 123-4567",
                    profile={"trade": "Cable", "level": "L2", "status": "standby",
                             "partner_id": pt.id})
    out = await preview(db, admin, [{
        "first_name": "Robert", "last_name": "Smith", "email": "bob@test.example.com",
        "phone": "555.123.4567", "status": "", "country": "", "city": "Reno",
        "level": "L3", "partner": ""}])
    row = out["rows"][0]
    assert row["action"] == "update"
    assert row["diff"] == {"city": {"old": None, "new": "Reno"},
                           "level": {"old": "L2", "new": "L3"}}


async def test_status_change_guards_rank_and_self(db, admin):
    boss = await mk_worker(db, "Big", "Boss", email="boss@test.example.com")
    db.add(PersonRole(person_id=boss.id, role="developer"))     # rank above admin's 60
    await db.commit()
    out = await preview(db, admin, [
        {"first_name": "Big", "last_name": "Boss", "status": "standby"},
        {"first_name": "Ada", "last_name": "Admin", "status": "standby"},
    ])
    rows = out["rows"]
    assert rows[0]["errors"] == ["rank too low to edit this person"]
    assert rows[1]["errors"] == ["cannot change your own status"]


async def test_rank_guard_covers_every_edit_of_an_account_holder(db, admin):
    """PATCH /workers/{id}/person refuses any edit to a higher-ranked person
    who can log in; the bulk path must refuse the same edit."""
    boss = await mk_worker(db, "Big", "Boss", email="boss@test.example.com",
                           account=True)
    db.add(PersonRole(person_id=boss.id, role="developer"))     # rank 100
    await db.commit()
    blocked = await one(db, admin, {"first_name": "Big", "last_name": "Boss",
                                    "email": "boss@test.example.com", "city": "Reno"})
    assert blocked["errors"] == ["rank too low to edit this person"]
    # an unchanged row never trips the guard — an export of the whole
    # workforce still previews clean
    same = await one(db, admin, {"first_name": "Big", "last_name": "Boss",
                                 "email": "boss@test.example.com"})
    assert same["errors"] == [] and same["action"] == "unchanged"
    # the guard follows the login account, exactly as the endpoint does
    plain = await mk_worker(db, "No", "Account", email="na@test.example.com")
    db.add(PersonRole(person_id=plain.id, role="developer"))
    await db.commit()
    ok = await one(db, admin, {"first_name": "No", "last_name": "Account",
                               "email": "na@test.example.com", "city": "Reno"})
    assert ok["errors"] == [] and ok["action"] == "update"


# ── commit ──────────────────────────────────────────────────────────

async def test_commit_creates_person_role_and_profile(db, admin):
    pt = Partner(name="Haul It")
    db.add(pt)
    await db.commit()
    out = await commit(db, admin, [{
        "first_name": "Maria", "last_name": "Lopez", "phone": "(555) 987-6543",
        "employee_number": "E7", "rfid_tag": "RF1", "partner": "haul it",
        "trade": "Cable", "level": "L2", "status": "standby", "city": "Reno"}],
        source="crew.xlsx")
    assert out["created"] == 1 and out["updated"] == out["skipped"] == out["unchanged"] == 0
    row = out["rows"][0]
    assert row["action"] == "created" and row["name"] == "Maria Lopez" and row["diff"] is None
    person = await db.get(Person, uuid.UUID(row["person_id"]))
    assert person.phone == "(555) 987-6543" and person.external_id == "E7"
    assert person.rfid_tag == "RF1" and person.country == "US"
    assert person.source == "import" and person.source_ref == "crew.xlsx"
    assert person.created_by == admin.id
    profile = await db.get(WorkerProfile, person.id)
    assert profile.partner_id == pt.id and profile.trade == "Cable"
    assert profile.level == "L2" and profile.status == "standby"
    assert await db.scalar(select(PersonRole.id).where(
        PersonRole.person_id == person.id, PersonRole.role == "worker",
        PersonRole.revoked_at.is_(None))) is not None
    actions = list(await db.scalars(select(AuditLog.action).where(
        AuditLog.entity_type == "worker")))
    assert sorted(actions) == ["bulk_import", "create"]


async def test_commit_updates_approved_skips_unapproved_counts_unchanged(db, admin):
    a = await mk_worker(db, "Robert", "Smith", email="a@test.example.com")
    b = await mk_worker(db, "Sara", "Jones", email="b@test.example.com")
    await mk_worker(db, "Same", "Person", email="s@test.example.com")
    out = await commit(db, admin, [
        {"first_name": "Robert", "last_name": "Smith", "city": "Reno"},
        {"first_name": "Sara", "last_name": "Jones", "city": "Austin"},
        {"first_name": "Same", "last_name": "Person"},
        {"first_name": "Brand", "last_name": "New"},
    ], approved=[str(a.id)])
    assert (out["created"], out["updated"], out["skipped"], out["unchanged"]) == (1, 1, 1, 1)
    by_name = {r["name"]: r for r in out["rows"]}
    assert by_name["Robert Smith"]["action"] == "updated"
    assert by_name["Robert Smith"]["diff"] == {"city": {"old": None, "new": "Reno"}}
    assert by_name["Sara Jones"]["action"] == "skipped"
    assert by_name["Sara Jones"]["diff"] == {"city": {"old": None, "new": "Austin"}}
    assert by_name["Same Person"]["action"] == "unchanged"
    assert by_name["Brand New"]["action"] == "created"
    await db.refresh(a)
    await db.refresh(b)
    assert a.city == "Reno" and b.city is None            # skipped row untouched
    bulk_row = await db.scalar(select(AuditLog).where(AuditLog.action == "bulk_import"))
    assert bulk_row.changes == {"created": 1, "updated": 1, "skipped": 1,
                                "unchanged": 1, "source": "test.csv"}


async def test_commit_grants_role_and_creates_profile_on_matched_non_worker(db, admin):
    user = await mk_worker(db, "Office", "User", email="ou@test.example.com", role=False)
    out = await commit(db, admin, [
        {"first_name": "Office", "last_name": "User", "trade": "Cable"}],
        approved=[str(user.id)])
    assert out["updated"] == 1
    assert await db.scalar(select(PersonRole.id).where(
        PersonRole.person_id == user.id, PersonRole.role == "worker",
        PersonRole.revoked_at.is_(None))) is not None
    profile = await db.get(WorkerProfile, user.id)
    assert profile.trade == "Cable" and profile.status == "active"


async def test_commit_blank_status_country_never_written_on_update(db, admin):
    w = await mk_worker(db, "Keep", "Country", email="k@test.example.com",
                        profile={"status": "standby"})
    w.country = "CH"
    await db.commit()
    out = await commit(db, admin, [
        {"first_name": "Keep", "last_name": "Country", "status": "", "country": "",
         "city": "Zurich"}], approved=[str(w.id)])
    assert out["updated"] == 1
    await db.refresh(w)
    assert w.country == "CH" and w.city == "Zurich"
    assert (await db.get(WorkerProfile, w.id)).status == "standby"


async def test_commit_blacklist_disables_account_and_unblacklist_restores(db, admin):
    w = await mk_worker(db, "Bad", "Actor", email="bad@test.example.com", account=True)
    out = await commit(db, admin, [
        {"first_name": "Bad", "last_name": "Actor", "status": "blacklist",
         "status_note": "no-show x3"}], approved=[str(w.id)])
    assert out["updated"] == 1
    account = await db.get(UserAccount, w.id)
    assert account.disabled_at is not None
    out = await commit(db, admin, [
        {"first_name": "Bad", "last_name": "Actor", "status": "active"}],
        approved=[str(w.id)])
    assert out["updated"] == 1
    await db.refresh(account)
    assert account.disabled_at is None


async def test_commit_is_all_or_nothing(db, admin):
    with pytest.raises(bi.BulkImportError) as exc:
        await commit(db, admin, [
            {"first_name": "Good", "last_name": "Row"},
            {"first_name": "", "last_name": "Bad"},
        ])
    assert exc.value.code == "rows_invalid"
    assert [r["action"] for r in exc.value.extra["rows"]] == ["create", "error"]
    assert await db.scalar(select(func.count()).select_from(Person).where(
        Person.last_name == "Row")) == 0
    with pytest.raises(bi.BulkImportError):
        await commit(db, admin, [])
