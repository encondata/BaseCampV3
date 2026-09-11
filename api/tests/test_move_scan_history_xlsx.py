"""build_workbook() — Overview + Scan History sheets, port of V2's
generate_xlsx() (api/reports/scan_history_report.py) plus a `build()`
end-to-end test through the registry. See docs/superpowers/specs/
2026-09-11-move-scan-history-design.md."""

from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo

import openpyxl

from serversherpa.db.models import (
    Asset, Client, Initiative, InitiativeAsset, Person, ProcessedScan, ReportDefinition,
    ReportRun,
)
from serversherpa.reports.move_scan_history.gather import (
    AssetRow, ScanHistoryData, ScanHit, StatusCol,
)
from serversherpa.reports.move_scan_history.xlsx import XLSX_MIME, build_workbook
from serversherpa.reports.registry import get_module

TZ = ZoneInfo("America/New_York")

COLUMNS = [
    StatusCol(key="pre_stage", label="Pre-Stage", color="#caa0a0", in_pipeline=True, scan_count=2),
    StatusCol(key="complete", label="Complete", color="#8e27f5", in_pipeline=True, scan_count=1),
]


def _populated_data() -> ScanHistoryData:
    return ScanHistoryData(
        initiative_id=uuid4(), name="NAP11 Move", client_name="Acme",
        scheduled_start=datetime(2026, 3, 15, 9, 0, tzinfo=UTC),
        source_name="DC-A", destination_name="DC-B",
        assets=[AssetRow(100, "SN-B", "Asset B"), AssetRow(200, "SN-C", "Asset C")],
        statuses=COLUMNS,
        # deliberately unsorted (both across and within assets) — proves
        # build_workbook does its own sort rather than trusting the caller
        scan_progress={
            200: [ScanHit("pre_stage", "Pre-Stage", datetime(2026, 3, 16, 10, 0, tzinfo=UTC))],
            100: [
                ScanHit("complete", "Complete", datetime(2026, 3, 16, 14, 30, tzinfo=UTC)),
                ScanHit("pre_stage", "Pre-Stage", datetime(2026, 3, 16, 8, 0, tzinfo=UTC)),
            ],
        },
        total_assets=2, scanned_assets=2, completed=1, completion_pct=50,
        last_scan_at=datetime(2026, 3, 16, 14, 30, tzinfo=UTC))


def _load(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(BytesIO(content))


def test_overview_block_rows_and_bold_labels():
    data = _populated_data()
    generated_at = datetime(2026, 3, 17, 8, 0, tzinfo=UTC)
    wb = _load(build_workbook(data, COLUMNS, generated_at, TZ))
    ws = wb["Overview"]

    expected = [
        ("Move Name", "NAP11 Move"),
        ("Client", "Acme"),
        ("Scheduled Start", "03/15/2026"),
        ("Source", "DC-A"),
        ("Destination", "DC-B"),
        ("Date Generated", "03/17/2026, 04:00:00 AM EDT"),
        ("Time Zone", "America/New_York (EDT, UTC-04:00)"),
        ("Total Assets", 2),
        ("Completion", "50%"),
    ]
    for row_idx, (label, value) in enumerate(expected, start=1):
        assert ws.cell(row=row_idx, column=1).value == label
        assert ws.cell(row=row_idx, column=1).font.bold is True
        assert ws.cell(row=row_idx, column=2).value == value
    assert ws.cell(row=10, column=1).value is None  # blank separator row


def test_overview_missing_client_and_site_names_render_na_and_not_scheduled():
    data = ScanHistoryData(
        initiative_id=uuid4(), name="Solo Move", client_name=None, scheduled_start=None,
        source_name=None, destination_name=None, assets=[], statuses=[], scan_progress={},
        total_assets=0, scanned_assets=0, completed=0, completion_pct=0, last_scan_at=None)
    wb = _load(build_workbook(data, [], datetime(2026, 3, 17, 8, 0, tzinfo=UTC), TZ))
    ws = wb["Overview"]
    assert ws.cell(row=2, column=2).value == "N/A"           # Client
    assert ws.cell(row=3, column=2).value == "Not scheduled"  # Scheduled Start
    assert ws.cell(row=4, column=2).value == "N/A"           # Source
    assert ws.cell(row=5, column=2).value == "N/A"           # Destination


def test_overview_header_row_and_asset_rows():
    data = _populated_data()
    generated_at = datetime(2026, 3, 17, 8, 0, tzinfo=UTC)
    wb = _load(build_workbook(data, COLUMNS, generated_at, TZ))
    ws = wb["Overview"]

    header_row = 11   # 9 block rows + blank separator
    assert [ws.cell(row=header_row, column=c).value for c in range(1, 5)] == (
        ["Asset ID", "Serial Number", "Asset Name", "Pre-Stage"])
    assert ws.cell(row=header_row, column=5).value == "Complete"
    for col in range(1, 6):
        assert ws.cell(row=header_row, column=col).font.bold is True

    # asset legacy 100 (Asset B): complete + pre_stage both filled
    row_100 = header_row + 1
    assert [ws.cell(row=row_100, column=c).value for c in range(1, 6)] == [
        100, "SN-B", "Asset B", "03/16/2026 04:00:00 AM", "03/16/2026 10:30:00 AM"]
    # asset legacy 200 (Asset C): only pre_stage filled, complete blank
    # (an empty-string cell round-trips through openpyxl save/load as None)
    row_200 = header_row + 2
    assert [ws.cell(row=row_200, column=c).value for c in range(1, 6)] == [
        200, "SN-C", "Asset C", "03/16/2026 06:00:00 AM", None]


def test_overview_no_assets_shows_placeholder_message():
    data = ScanHistoryData(
        initiative_id=uuid4(), name="Empty Move", client_name=None, scheduled_start=None,
        source_name=None, destination_name=None, assets=[], statuses=[], scan_progress={},
        total_assets=0, scanned_assets=0, completed=0, completion_pct=0, last_scan_at=None)
    wb = _load(build_workbook(data, [], datetime(2026, 3, 17, 8, 0, tzinfo=UTC), TZ))
    ws = wb["Overview"]
    assert ws.cell(row=11, column=1).value == "No assets found in this move"


def test_scan_history_sheet_header_and_detail_sort():
    data = _populated_data()
    generated_at = datetime(2026, 3, 17, 8, 0, tzinfo=UTC)
    wb = _load(build_workbook(data, COLUMNS, generated_at, TZ))
    ws = wb["Scan History"]

    assert [ws.cell(row=1, column=c).value for c in range(1, 6)] == [
        "Asset ID", "Serial Number", "Asset Name", "Status Name", "Timestamp (EDT)"]
    for col in range(1, 6):
        assert ws.cell(row=1, column=col).font.bold is True

    # sorted by asset_id then timestamp, regardless of insertion order
    rows = [[ws.cell(row=r, column=c).value for c in range(1, 6)] for r in (2, 3, 4)]
    assert rows == [
        [100, "SN-B", "Asset B", "Pre-Stage", "03/16/2026 04:00:00 AM"],
        [100, "SN-B", "Asset B", "Complete", "03/16/2026 10:30:00 AM"],
        [200, "SN-C", "Asset C", "Pre-Stage", "03/16/2026 06:00:00 AM"],
    ]
    assert ws.cell(row=5, column=1).value is None


def test_scan_history_sheet_no_scans_shows_placeholder_message():
    data = ScanHistoryData(
        initiative_id=uuid4(), name="Quiet Move", client_name="Acme", scheduled_start=None,
        source_name=None, destination_name=None,
        assets=[AssetRow(1, "SN-1", "Only Asset")], statuses=[], scan_progress={},
        total_assets=1, scanned_assets=0, completed=0, completion_pct=0, last_scan_at=None)
    wb = _load(build_workbook(data, [], datetime(2026, 3, 17, 8, 0, tzinfo=UTC), TZ))
    ws = wb["Scan History"]
    assert ws.cell(row=2, column=1).value == "No scan history found"


def test_column_width_auto_fit_capped_at_30():
    long_name = "A" * 50  # +2 padding would be 52, capped to 30
    data = ScanHistoryData(
        initiative_id=uuid4(), name="Widths Move", client_name="Acme", scheduled_start=None,
        source_name=None, destination_name=None,
        assets=[AssetRow(1, "SN-1", long_name)],
        statuses=[], scan_progress={}, total_assets=1, scanned_assets=0, completed=0,
        completion_pct=0, last_scan_at=None)
    wb = _load(build_workbook(data, [], datetime(2026, 3, 17, 8, 0, tzinfo=UTC), TZ))
    ws = wb["Overview"]
    assert ws.column_dimensions["C"].width == 30  # Asset Name column


async def test_build_produces_xlsx_through_the_registry(db):
    module = get_module("move_scan_history")
    assert module.report_type == "move_scan_history"

    person = Person(first_name="Rae", last_name="Requester")
    client = Client(name="Acme")
    db.add_all([person, client])
    await db.flush()
    ini = Initiative(name="NAP11 / Move #1", initiative_type="move", status="planned",
                     client_id=client.id)
    definition = ReportDefinition(name="Move Scan History", report_type="move_scan_history",
                                  options={"default_format": "xlsx", "status_columns": "pipeline"},
                                  is_system=True)
    db.add_all([ini, definition])
    await db.flush()

    asset = Asset(legacy_id=9001, serial_number="SN-9001", name="Widget")
    db.add(asset)
    await db.flush()
    db.add(InitiativeAsset(initiative_id=ini.id, asset_id=asset.id))
    db.add(ProcessedScan(
        scanned_value="EPC-9001", scan_type="rfid", scanned_at=datetime.now(UTC),
        processed_at=datetime.now(UTC), match_type="asset", asset_id=asset.id, status="complete"))
    await db.flush()

    run = ReportRun(definition_id=definition.id, report_type="move_scan_history",
                    initiative_id=ini.id, options={}, requested_by=person.id, requested_rank=0)
    db.add(run)
    await db.commit()

    result = await module.build(db, run)
    assert result.content_type == XLSX_MIME
    # V3 filename house style, slashes sanitized out of the initiative name
    assert result.filename.startswith("Move Scan History - NAP11 - Move #1 - ")
    assert result.filename.endswith(".xlsx")

    wb = _load(result.content)
    assert wb.sheetnames == ["Overview", "Scan History"]
