"""Update assets in bulk (no HTTP): parse, resolve each row to an asset and
each value to a record, preview, template, export — plus the shared model
catalog index the roster importer also uses."""
import io
import uuid

import openpyxl
import pytest
from sqlalchemy import func

from serversherpa.assets import bulk_update as bu
from serversherpa.assets.model_index import build_model_index, find_model
from serversherpa.db.models import (
    Asset,
    AssetModel,
    AssetModelAlias,
    Client,
    Site,
    StatusValue,
)
from serversherpa.imports.bulk import BulkImportError

PAD = "0" * 18


# ── fixtures ────────────────────────────────────────────────────────

async def mk_model(db, make, model, category=None, aliases=()):
    m = AssetModel(make=make, model=model, category=category)
    db.add(m)
    await db.flush()
    for a in aliases:
        db.add(AssetModelAlias(model_id=m.id, alias=a))
    await db.commit()
    return m


async def mk_client(db, name, archived=False):
    c = Client(name=name, archived_at=func.now() if archived else None)
    db.add(c)
    await db.commit()
    return c


async def mk_site(db, name, archived=False):
    s = Site(name=name, archived_at=func.now() if archived else None)
    db.add(s)
    await db.commit()
    return s


async def mk_asset(db, number, serial, name=None, archived=False, **fields):
    a = Asset(legacy_id=number, serial_number=serial, name=name,
              archived_at=func.now() if archived else None, **fields)
    db.add(a)
    await db.commit()
    await db.refresh(a)
    return a


async def preview(db, rows, **kw):
    return await bu.preview_rows(db, bu.number_json_rows(rows), **kw)


async def one(db, row, **kw):
    return (await preview(db, [row], **kw))["rows"][0]


# ── shape / parsing ─────────────────────────────────────────────────

def test_columns_sheet_and_limits():
    assert bu.COLUMNS == [
        "asset_id", "serial_number", "name", "new_serial_number", "rfid_tag",
        "make", "model", "client", "site", "location", "pod", "status", "has_rails"]
    assert bu.SHEET == "Assets"
    assert bu.MAX_ROWS == 15000
    assert bu.MAX_BYTES == 20 * 1024 * 1024


def test_parse_upload_accepts_15000_rows_and_refuses_15001():
    ok = ("asset_id\n" + "1\n" * 15000).encode()
    assert len(bu.parse_upload("a.csv", ok)) == 15000
    with pytest.raises(BulkImportError) as exc:
        bu.parse_upload("a.csv", ("asset_id\n" + "1\n" * 15001).encode())
    assert exc.value.code == "too_many_rows"
    assert exc.value.extra == {"limit": 15000}


def test_parse_upload_refuses_unknown_columns():
    with pytest.raises(BulkImportError) as exc:
        bu.parse_upload("a.csv", b"asset_id,colour\n1,red\n")
    assert exc.value.code == "unknown_columns"


def test_normalize_rfid():
    assert bu.normalize_rfid(" e2 80 11 ") == PAD + "E28011"
    assert bu.normalize_rfid("A" * 24) == "A" * 24
    assert bu.normalize_rfid("A" * 25) is None
    assert bu.normalize_rfid("E2-80") is None
    assert bu.normalize_rfid("É280") is None
    assert bu.normalize_rfid("   ") is None


def test_parse_overrides_and_row_lists():
    assert bu.parse_overrides(None) == {}
    assert bu.parse_overrides({"3": {"asset": "x", "status": "racked"}}) == {
        3: {"asset": "x", "status": "racked"}}
    for bad in ([], {"x": {}}, {"3": {"worker": "x"}}, {"3": {"model": ""}},
                {"3": {"model": 5}}, {"3": "x"}):
        with pytest.raises(BulkImportError) as exc:
            bu.parse_overrides(bad)
        assert exc.value.code == "invalid_overrides"
    assert bu.parse_row_list(None, "invalid_skip") == set()
    assert bu.parse_row_list([2, 3], "invalid_skip") == {2, 3}
    for bad in ("2", [True], ["2"], {"2": 1}):
        with pytest.raises(BulkImportError) as exc:
            bu.parse_row_list(bad, "invalid_skip")
        assert exc.value.code == "invalid_skip"


# ── matching the asset ──────────────────────────────────────────────

async def test_match_by_asset_id(db):
    a = await mk_asset(db, 100123, "SN-1", name="old-name")
    r = await one(db, {"asset_id": "100123", "name": "new-name"})
    assert r["action"] == "update"
    assert r["matched_by"] == "asset ID"
    assert r["asset_id"] == str(a.id)
    assert r["asset_number"] == 100123
    assert r["name"] == "old-name"
    assert r["diff"] == {"name": {"old": "old-name", "new": "new-name"}}
    assert r["changes"] == {"name": "new-name"}


async def test_asset_id_wins_over_serial(db):
    a = await mk_asset(db, 100123, "SN-1")
    await mk_asset(db, 100124, "SN-2")
    r = await one(db, {"asset_id": "100123", "serial_number": "SN-2", "pod": "P1"})
    assert r["asset_id"] == str(a.id) and r["matched_by"] == "asset ID"


async def test_match_by_serial_case_insensitively(db):
    a = await mk_asset(db, 100123, "SN-AbC")
    r = await one(db, {"serial_number": "sn-abc", "pod": "P7"})
    assert r["action"] == "update", r
    assert r["matched_by"] == "serial"
    assert r["asset_id"] == str(a.id)
    assert r["name"] == "SN-AbC"          # no asset name → the serial
    assert r["changes"] == {"pod_number": "P7"}


async def test_serial_ignores_archived_assets(db):
    live = await mk_asset(db, 100123, "SN-1")
    await mk_asset(db, 100124, "SN-1", archived=True)
    r = await one(db, {"serial_number": "SN-1", "pod": "P1"})
    assert r["action"] == "update" and r["asset_id"] == str(live.id)


async def test_duplicate_serial_needs_a_pick_and_the_pick_resolves(db):
    site = await mk_site(db, "DC West")
    a = await mk_asset(db, 100123, "DUP-1", name="alpha", site_id=site.id)
    b = await mk_asset(db, 100124, "DUP-1", name="beta")
    r = await one(db, {"serial_number": "dup-1", "status": "racked"})
    assert r["action"] == "attention"
    assert r["asset_id"] is None
    [issue] = r["issues"]
    assert issue["field"] == "asset" and issue["kind"] == "ambiguous"
    assert issue["value"] == "dup-1"
    assert sorted(issue["candidates"], key=lambda c: c["label"]) == [
        {"id": str(a.id), "label": "Asset 100123", "detail": "DUP-1 · alpha · DC West"},
        {"id": str(b.id), "label": "Asset 100124", "detail": "DUP-1 · beta"},
    ]
    picked = await one(db, {"serial_number": "dup-1", "status": "racked"},
                       overrides={1: {"asset": str(b.id)}})
    assert picked["action"] == "update"
    assert picked["matched_by"] == "your pick"
    assert picked["asset_id"] == str(b.id)
    assert picked["issues"] == []


async def test_asset_pick_without_that_serial_is_an_error(db):
    await mk_asset(db, 100123, "DUP-1")
    await mk_asset(db, 100124, "DUP-1")
    other = await mk_asset(db, 100125, "OTHER")
    r = await one(db, {"serial_number": "DUP-1", "pod": "P1"},
                  overrides={1: {"asset": str(other.id)}})
    assert r["action"] == "error"
    assert r["errors"] == [
        "The chosen asset is no longer a live asset with serial 'DUP-1'. Pick again."]


async def test_archived_asset_id_is_an_error(db):
    await mk_asset(db, 12345, "SN-1", archived=True)
    r = await one(db, {"asset_id": "12345", "pod": "P1"})
    assert r["action"] == "error"
    assert r["errors"] == ["Asset 12345 is archived."]


async def test_not_found_and_missing_keys_are_errors(db):
    p = await preview(db, [
        {"asset_id": "12345", "pod": "P1"},
        {"serial_number": "NOPE", "pod": "P1"},
        {"pod": "P1"},
        {"asset_id": "12a", "pod": "P1"},
    ])
    assert [r["action"] for r in p["rows"]] == ["error"] * 4
    assert [r["errors"] for r in p["rows"]] == [
        ["No asset with Asset ID 12345."],
        ["No live asset with serial 'NOPE'."],
        ["Each row needs an asset_id or a serial_number."],
        ["Asset ID must be a number."],
    ]
    assert p["can_commit"] is False


async def test_blank_cells_mean_no_change(db):
    client = await mk_client(db, "Acme")
    await mk_asset(db, 100123, "SN-1", name="n1", client_id=client.id, pod_number="P1",
                   location_detail="Cage 4", status="racked", has_rails=True)
    r = await one(db, {"asset_id": "100123"})
    assert r["action"] == "unchanged"
    # every cell filled but equal to what the asset already has
    p = await preview(db, [{"serial_number": "SN-1", "name": "n1", "client": "acme",
                            "pod": "P1", "location": "Cage 4", "status": "Racked",
                            "has_rails": "yes"}])
    [r] = p["rows"]
    assert r["action"] == "unchanged", r
    assert r["diff"] is None and r["changes"] == {}
    assert p["counts"]["unchanged"] == 1
    assert p["can_commit"] is True


# ── fields ──────────────────────────────────────────────────────────

async def test_every_field_changes_with_display_names(db):
    old_model = await mk_model(db, "Dell", "R640")
    new_model = await mk_model(db, "Dell", "R740")
    old_client = await mk_client(db, "Acme")
    new_client = await mk_client(db, "Globex")
    old_site = await mk_site(db, "DC West")
    new_site = await mk_site(db, "DC East")
    a = await mk_asset(db, 100123, "SN-1", name="old", rfid_tag=PAD + "AAAAAA",
                       model_id=old_model.id, client_id=old_client.id, site_id=old_site.id,
                       location_detail="Cage 1", pod_number="P1", status="unknown",
                       has_rails=False)
    r = await one(db, {
        "asset_id": "100123", "name": "new", "new_serial_number": "SN-2",
        "rfid_tag": "bbbbbb", "make": "dell", "model": "r740", "client": "GLOBEX",
        "site": "dc east", "location": "Cage 2", "pod": "P2", "status": "Racked",
        "has_rails": "Yes"})
    assert r["action"] == "update", r
    assert r["diff"] == {
        "name": {"old": "old", "new": "new"},
        "serial_number": {"old": "SN-1", "new": "SN-2"},
        "rfid_tag": {"old": PAD + "AAAAAA", "new": PAD + "BBBBBB"},
        "model": {"old": "Dell R640", "new": "Dell R740"},
        "client": {"old": "Acme", "new": "Globex"},
        "site": {"old": "DC West", "new": "DC East"},
        "location": {"old": "Cage 1", "new": "Cage 2"},
        "pod": {"old": "P1", "new": "P2"},
        "status": {"old": "Unknown", "new": "Racked"},
        "has_rails": {"old": "no", "new": "yes"},
    }
    assert r["changes"] == {
        "name": "new", "serial_number": "SN-2", "rfid_tag": PAD + "BBBBBB",
        "model_id": str(new_model.id), "client_id": str(new_client.id),
        "site_id": str(new_site.id), "location_detail": "Cage 2", "pod_number": "P2",
        "status": "racked", "has_rails": True,
    }
    assert r["asset_id"] == str(a.id)


async def test_diff_old_values_name_archived_and_missing_records(db):
    gone = await mk_client(db, "Old Co", archived=True)
    await mk_client(db, "New Co")
    await mk_asset(db, 100123, "SN-1", client_id=gone.id)
    r = await one(db, {"asset_id": "100123", "client": "New Co", "has_rails": "no"})
    assert r["diff"] == {"client": {"old": "Old Co", "new": "New Co"},
                         "has_rails": {"old": None, "new": "no"}}


async def test_new_serial_needs_an_asset_id(db):
    await mk_asset(db, 100123, "SN-1")
    r = await one(db, {"serial_number": "SN-1", "new_serial_number": "SN-2"})
    assert r["action"] == "error"
    assert r["errors"] == ["Change the serial only on rows with an asset_id."]


async def test_rfid_is_normalized(db):
    await mk_asset(db, 100123, "SN-1")
    r = await one(db, {"asset_id": "100123", "rfid_tag": " e2 80 11 "})
    assert r["changes"] == {"rfid_tag": PAD + "E28011"}
    await mk_asset(db, 100124, "SN-2", rfid_tag=PAD + "E28012")
    r = await one(db, {"asset_id": "100124", "rfid_tag": "e28012"})
    assert r["action"] == "unchanged"      # its own tag is not "taken"


async def test_bad_or_taken_rfid_is_an_error(db):
    await mk_asset(db, 100123, "SN-1")
    await mk_asset(db, 100124, "SN-2")
    await mk_asset(db, 100200, "SN-HOLDER", rfid_tag=PAD + "E28011")
    p = await preview(db, [
        {"asset_id": "100123", "rfid_tag": "E2-80"},
        {"serial_number": "SN-2", "rfid_tag": "e28011"},
    ])
    bad, taken = p["rows"]
    assert bad["errors"] == ["RFID tag 'E2-80' is not valid."]
    assert taken["errors"] == ["RFID tag e28011 is already on asset 100200."]


async def test_two_rows_setting_one_rfid_are_both_errors(db):
    await mk_asset(db, 100123, "SN-1")
    await mk_asset(db, 100124, "SN-2")
    p = await preview(db, [
        {"asset_id": "100123", "rfid_tag": "ABC1"},
        {"asset_id": "100124", "rfid_tag": " abc1"},
    ])
    assert [r["action"] for r in p["rows"]] == ["error", "error"]
    assert p["rows"][0]["errors"] == ["RFID tag ABC1 appears on more than one row (1, 2)."]
    assert p["rows"][1]["errors"] == ["RFID tag abc1 appears on more than one row (1, 2)."]


async def test_make_and_model_go_together(db):
    await mk_asset(db, 100123, "SN-1")
    await mk_asset(db, 100124, "SN-2")
    p = await preview(db, [
        {"asset_id": "100123", "make": "Dell"},
        {"asset_id": "100124", "model": "R740"},
    ])
    for r in p["rows"]:
        assert r["action"] == "error"
        assert r["errors"] == ["Fill both make and model, or neither."]


async def test_model_exact_alias_and_normalized_match(db):
    m = await mk_model(db, "DellEMC", "Isilon H5600 (Chassis)", aliases=["Isilon Box"])
    for n in (100123, 100124, 100125):
        await mk_asset(db, n, f"SN-{n}")
    p = await preview(db, [
        {"asset_id": "100123", "make": "dellemc", "model": "isilon h5600 (chassis)"},
        {"asset_id": "100124", "make": "Isilon", "model": "Box"},
        {"asset_id": "100125", "make": "DellEMC_Isilon", "model": "H5600 Chassis 4U"},
    ])
    for r in p["rows"]:
        assert r["action"] == "update", r
        assert r["changes"] == {"model_id": str(m.id)}
        assert r["diff"] == {"model": {"old": None, "new": "DellEMC Isilon H5600 (Chassis)"}}


async def test_unknown_model_needs_a_pick_and_the_pick_resolves(db):
    m = await mk_model(db, "Dell", "R740", category="server")
    await mk_asset(db, 100123, "SN-1")
    r = await one(db, {"asset_id": "100123", "make": "Dell", "model": "R999"})
    assert r["action"] == "attention"
    assert r["issues"] == [{"field": "model", "kind": "unknown", "value": "Dell R999",
                            "candidates": []}]
    picked = await one(db, {"asset_id": "100123", "make": "Dell", "model": "R999"},
                       overrides={1: {"model": str(m.id)}})
    assert picked["action"] == "update"
    assert picked["changes"] == {"model_id": str(m.id)}
    stale = await one(db, {"asset_id": "100123", "make": "Dell", "model": "R999"},
                      overrides={1: {"model": str(uuid.uuid4())}})
    assert stale["action"] == "error"
    assert stale["errors"] == ["The chosen model no longer exists. Pick again."]


async def test_ambiguous_normalized_model_offers_both(db):
    p1 = await mk_model(db, "Blank", "Panel 1U", category="network")
    p2 = await mk_model(db, "Blank", "Panel 2U")
    await mk_asset(db, 100123, "SN-1")
    r = await one(db, {"asset_id": "100123", "make": "Blank_Panel", "model": "3U"})
    [issue] = r["issues"]
    assert issue["kind"] == "ambiguous"
    assert issue["candidates"] == [
        {"id": str(p1.id), "label": "Blank Panel 1U", "detail": "network"},
        {"id": str(p2.id), "label": "Blank Panel 2U", "detail": ""},
    ]


async def test_status_by_label_or_key_and_unknown_or_inactive_needs_a_pick(db):
    for n in (100123, 100124, 100125, 100126):
        await mk_asset(db, n, f"SN-{n}")
    await db.execute(StatusValue.__table__.update()
                     .where(StatusValue.record_type == "asset",
                            StatusValue.key == "historical")
                     .values(is_active=False))
    await db.commit()
    p = await preview(db, [
        {"asset_id": "100123", "status": "in transit"},
        {"asset_id": "100124", "status": "IN_TRANSIT"},
        {"asset_id": "100125", "status": "Floating"},
        {"asset_id": "100126", "status": "historical"},
    ])
    by_label, by_key, unknown, inactive = p["rows"]
    assert by_label["changes"] == {"status": "in_transit"}
    assert by_key["changes"] == {"status": "in_transit"}
    assert by_label["diff"] == {"status": {"old": "Unknown", "new": "In Transit"}}
    assert unknown["action"] == inactive["action"] == "attention"
    assert unknown["issues"][0]["field"] == "status"
    picked = await one(db, {"asset_id": "100125", "status": "Floating"},
                       overrides={1: {"status": "racked"}})
    assert picked["changes"] == {"status": "racked"}


async def test_unknown_client_and_site_need_a_pick(db):
    await mk_client(db, "Old Co", archived=True)
    globex = await mk_client(db, "Globex")
    await mk_site(db, "Closed DC", archived=True)
    dup1 = await mk_site(db, "Twin")
    await mk_asset(db, 100123, "SN-1")
    r = await one(db, {"asset_id": "100123", "client": "Old Co", "site": "Closed DC"})
    assert r["action"] == "attention"
    assert [(i["field"], i["kind"], i["candidates"]) for i in r["issues"]] == [
        ("client", "unknown", []), ("site", "unknown", [])]
    r = await one(db, {"asset_id": "100123", "client": "Old Co", "site": "Closed DC"},
                  overrides={1: {"client": str(globex.id), "site": str(dup1.id)}})
    assert r["action"] == "update"
    assert r["changes"] == {"client_id": str(globex.id), "site_id": str(dup1.id)}


async def test_error_row_keeps_its_issues(db):
    r = await one(db, {"asset_id": "999", "status": "Floating"})
    assert r["action"] == "error"
    assert r["issues"][0]["field"] == "status"


async def test_has_rails_parsing(db):
    await mk_asset(db, 100123, "SN-1", has_rails=None)
    for text, value in (("yes", True), ("Y", True), ("true", True), ("1", True),
                        ("NO", False), ("n", False), ("False", False), ("0", False)):
        r = await one(db, {"asset_id": "100123", "has_rails": text})
        assert r["changes"] == {"has_rails": value}, text
    r = await one(db, {"asset_id": "100123", "has_rails": "maybe"})
    assert r["errors"] == ["has_rails must be yes or no."]


async def test_two_rows_on_one_asset_are_both_errors(db):
    await mk_asset(db, 12345, "SN-1")
    p = await preview(db, [
        {"asset_id": "12345", "pod": "P1"},
        {"serial_number": "sn-1", "pod": "P2"},
    ])
    for r in p["rows"]:
        assert r["action"] == "error"
        assert r["errors"] == ["Asset 12345 appears on more than one row (1, 2)."]


async def test_skip(db):
    await mk_asset(db, 100123, "SN-1")
    p = await preview(db, [{"asset_id": "100123", "pod": "P1"}, {"pod": "x"}],
                      skip={2})
    assert [r["action"] for r in p["rows"]] == ["update", "skipped"]
    assert p["counts"] == {"update": 1, "unchanged": 0, "attention": 0, "error": 0,
                           "skipped": 1}
    assert p["can_commit"] is True


async def test_listing_omits_unchanged_and_orders_rows(db):
    await mk_asset(db, 100123, "SN-1")
    await mk_asset(db, 100124, "SN-2")
    await mk_asset(db, 100125, "SN-3")
    p = await preview(db, [
        {"asset_id": "100123", "pod": "P1"},          # 1 update
        {"asset_id": "100124"},                       # 2 unchanged
        {"pod": "x"},                                 # 3 error
        {"asset_id": "100125", "status": "Floating"},  # 4 attention
        {"pod": "y"},                                 # 5 skipped
        {"pod": "z"},                                 # 6 error
    ], skip={5})
    out = bu.listing(p)
    assert [(r["row"], r["action"]) for r in out["rows"]] == [
        (4, "attention"), (3, "error"), (6, "error"), (1, "update"), (5, "skipped")]
    assert out["counts"]["unchanged"] == 1
    assert out["total"] == 6
    assert out["can_commit"] is False


# ── model index ─────────────────────────────────────────────────────

async def test_model_index_find_model(db):
    exact = await mk_model(db, "Dell", "R740", aliases=["PowerEdge 740"])
    p1 = await mk_model(db, "Blank", "Panel 1U")
    p2 = await mk_model(db, "Blank", "Panel 2U")
    index = await build_model_index(db)
    assert find_model(index, "DELL", "r740") == (exact, [])
    assert find_model(index, "PowerEdge", "740") == (exact, [])
    assert find_model(index, "Dell_R740", "2U") == (exact, [])
    assert find_model(index, "Blank", "Panel 2U") == (p2, [])
    match, candidates = find_model(index, "Blank_Panel", "")
    assert match is None and {c.id for c in candidates} == {p1.id, p2.id}
    assert find_model(index, "Nope", "Nothing") == (None, [])


# ── template / export ───────────────────────────────────────────────

async def test_template_round_trips_through_parse_upload(db):
    from_csv = bu.parse_upload("t.csv", bu.build_template_csv().encode())
    assert [n for n, _ in from_csv] == [2, 3]
    assert [r for _, r in from_csv] == [
        {c: str(s.get(c, "")) for c in bu.COLUMNS} for s in bu.SAMPLE_ROWS]
    assert from_csv[0][1]["asset_id"] and from_csv[0][1]["status"]
    assert from_csv[0][1]["site"]
    assert from_csv[1][1]["serial_number"] and from_csv[1][1]["make"]
    assert from_csv[1][1]["model"] and not from_csv[1][1]["asset_id"]

    await mk_model(db, "Dell", "R740")
    await mk_client(db, "Acme")
    await mk_client(db, "Gone", archived=True)
    await mk_site(db, "DC West")
    blob = await bu.build_template_xlsx(db)
    assert bu.parse_upload("t.xlsx", blob) == from_csv
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    assert wb.sheetnames == ["Assets", "Reference"]
    ref = [r[0] for r in wb["Reference"].iter_rows(values_only=True)]
    assert ref[0] == "Statuses"
    assert "racked — Racked" in ref
    assert ref[ref.index("Makes and models") + 1] == "Dell R740"
    assert ref[ref.index("Clients") + 1] == "Acme" and "Gone" not in ref
    assert ref[ref.index("Sites") + 1] == "DC West"


async def test_export_of_two_assets_previews_as_all_unchanged(db):
    m = await mk_model(db, "Dell", "R740")
    c = await mk_client(db, "Acme")
    s = await mk_site(db, "DC West")
    await mk_asset(db, 100124, "SN-2")
    await mk_asset(db, 100123, "SN-1", name="core-sw", rfid_tag=PAD + "E28011",
                   model_id=m.id, client_id=c.id, site_id=s.id, location_detail="Cage 4",
                   pod_number="P9", status="racked", has_rails=True)
    await mk_asset(db, 100125, "SN-3", archived=True)
    rows = await bu.export_rows(db)
    assert [str(r["asset_id"]) for r in rows] == ["100123", "100124"]
    assert rows[0] == {
        "asset_id": 100123, "serial_number": "SN-1", "name": "core-sw",
        "new_serial_number": "", "rfid_tag": PAD + "E28011", "make": "Dell",
        "model": "R740", "client": "Acme", "site": "DC West", "location": "Cage 4",
        "pod": "P9", "status": "racked", "has_rails": "yes"}
    assert rows[1]["has_rails"] == "" and rows[1]["status"] == "unknown"

    for numbered in (bu.parse_upload("e.csv", bu.build_rows_csv(rows).encode()),
                     bu.parse_upload("e.xlsx", await bu.build_export_xlsx(db))):
        p = await bu.preview_rows(db, numbered)
        assert p["counts"] == {"update": 0, "unchanged": 2, "attention": 0, "error": 0,
                               "skipped": 0}, p["rows"]
