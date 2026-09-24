"""Bulk-update-assets routes: rank gating, job lifecycle, preview/commit,
template/export formats, and that /{asset_id} still resolves (route order)."""
import io
import uuid

import openpyxl
import pytest
from sqlalchemy import select, text

from serversherpa.db.models import Asset, AuditLog, ImportJob, Person, PersonRole
from serversherpa.imports.worker import run_once
from tests.test_assets_api import login, make_login

CSV = "asset_id,serial_number,name,new_serial_number,rfid_tag,make,model,client,site,location,pod,status,has_rails\n"


def base(job_id) -> str:
    return f"/assets/bulk-update/{job_id}"


@pytest.fixture
async def admin_hdrs(db, client):
    person = Person(first_name="Ada", last_name="Admin", email="ada@test.example.com")
    db.add(person)
    await db.flush()
    db.add(PersonRole(person_id=person.id, role="admin"))
    await db.commit()
    return await make_login(db, client, person, "ada@test.example.com")


@pytest.fixture
async def other_admin_hdrs(db, client):
    person = Person(first_name="Owen", last_name="Other", email="owen@test.example.com")
    db.add(person)
    await db.flush()
    db.add(PersonRole(person_id=person.id, role="admin"))
    await db.commit()
    return await make_login(db, client, person, "owen@test.example.com")


async def mk_asset(db, number, serial, **fields):
    a = Asset(legacy_id=number, serial_number=serial, **fields)
    db.add(a)
    await db.commit()
    await db.refresh(a)
    return a


async def upload(client, headers, content):
    return await client.post(
        "/assets/bulk-update", headers=headers,
        files={"file": ("assets.csv", content.encode(), "text/csv")})


async def test_staff_forbidden_on_all_seven(client, db, seeded_user):
    hdrs = await login(client)
    job_id = uuid.uuid4()
    assert (await client.get(
        "/assets/bulk-update/template?format=csv", headers=hdrs)).status_code == 403
    assert (await client.get(
        "/assets/bulk-update/export?format=csv", headers=hdrs)).status_code == 403
    assert (await upload(client, hdrs, CSV)).status_code == 403
    assert (await client.post(
        f"{base(job_id)}/preview", headers=hdrs, json={})).status_code == 403
    assert (await client.post(
        f"{base(job_id)}/commit", headers=hdrs, json={})).status_code == 403
    assert (await client.get(base(job_id), headers=hdrs)).status_code == 403
    assert (await client.post(
        f"{base(job_id)}/cancel", headers=hdrs)).status_code == 403


async def test_upload_returns_job_and_preview_numbered_from_two(
        client, db, seeded_user, admin_hdrs):
    a = await mk_asset(db, 100, "SN-A", name="old-a")
    resp = await upload(client, admin_hdrs, CSV + f"{a.legacy_id},,new-a,,,,,,,,,,\n")
    assert resp.status_code == 201, resp.text
    body = resp.json()
    job = await db.get(ImportJob, uuid.UUID(body["job_id"]))
    assert job.kind == "asset_bulk_update"
    assert job.status == "preview" and job.phase == "preview"
    assert job.total_rows == 1
    assert job.created_by == (await db.scalar(
        select(Person.id).where(Person.email == "ada@test.example.com")))
    preview = body["preview"]
    assert [r["row"] for r in preview["rows"]] == [2]
    assert preview["rows"][0]["action"] == "update"
    log = (await db.scalars(select(AuditLog).where(
        AuditLog.action == "asset_bulk_update_job_create"))).all()
    assert len(log) == 1 and log[0].entity_type == "asset" and log[0].entity_id is None


async def test_another_admin_gets_404(client, db, seeded_user, admin_hdrs, other_admin_hdrs):
    resp = await upload(client, admin_hdrs, CSV)
    job_id = resp.json()["job_id"]
    denied = await client.get(base(job_id), headers=other_admin_hdrs)
    assert denied.status_code == 404
    assert denied.json()["detail"]["code"] == "job_not_found"


async def test_unknown_job_is_404(client, db, seeded_user, admin_hdrs):
    resp = await client.get(base(uuid.uuid4()), headers=admin_hdrs)
    assert resp.status_code == 404
    assert resp.json()["detail"]["code"] == "job_not_found"


async def test_preview_override_resolves_an_attention_row(
        client, db, seeded_user, admin_hdrs):
    await mk_asset(db, 100, "DUP-1")
    b = await mk_asset(db, 101, "DUP-1")
    resp = await upload(client, admin_hdrs, CSV + ",dup-1,,,,,,,,,,racked,\n")
    job_id = resp.json()["job_id"]
    first = resp.json()["preview"]
    assert first["rows"][0]["action"] == "attention"

    again = await client.post(
        f"{base(job_id)}/preview", headers=admin_hdrs,
        json={"overrides": {"2": {"asset": str(b.id)}}, "skip": []})
    assert again.status_code == 200, again.text
    [row] = again.json()["rows"]
    assert row["action"] == "update"
    assert row["asset_id"] == str(b.id)
    assert row["matched_by"] == "your pick"


async def test_commit_unresolved_is_422_rows_invalid(client, db, seeded_user, admin_hdrs):
    resp = await upload(client, admin_hdrs, CSV + ",NOPE,,,,,,,,,,,\n")
    job_id = resp.json()["job_id"]
    commit = await client.post(f"{base(job_id)}/commit", headers=admin_hdrs, json={})
    assert commit.status_code == 422
    body = commit.json()["detail"]
    assert body["code"] == "rows_invalid"
    assert [r["row"] for r in body["rows"]] == [2]

    job = await db.get(ImportJob, uuid.UUID(job_id))
    assert job.status == "preview"


async def test_commit_good_queues_job_with_stored_options(
        client, db, seeded_user, admin_hdrs):
    a = await mk_asset(db, 100, "SN-A", name="old-a")
    resp = await upload(client, admin_hdrs, CSV + f"{a.legacy_id},,new-a,,,,,,,,,,\n")
    job_id = resp.json()["job_id"]
    commit = await client.post(
        f"{base(job_id)}/commit", headers=admin_hdrs,
        json={"overrides": {}, "skip": [], "approved_updates": [2], "approve_all": False})
    assert commit.status_code == 200, commit.text
    out = commit.json()
    assert out["status"] == "queued" and out["phase"] == "commit"
    assert out["options"] == {"overrides": {}, "skip": [], "approved_updates": [2],
                              "approve_all": False}
    log = (await db.scalars(select(AuditLog).where(
        AuditLog.action == "asset_bulk_update_queued"))).all()
    assert len(log) == 1


async def test_run_once_completes_queued_job(client, db, seeded_user, admin_hdrs):
    a = await mk_asset(db, 100, "SN-A", name="old-a")
    resp = await upload(client, admin_hdrs, CSV + f"{a.legacy_id},,new-a,,,,,,,,,,\n")
    job_id = resp.json()["job_id"]
    await client.post(
        f"{base(job_id)}/commit", headers=admin_hdrs, json={"approve_all": True})

    from serversherpa.db.engine import get_sessionmaker
    assert await run_once(get_sessionmaker()) is True

    status = await client.get(base(job_id), headers=admin_hdrs)
    assert status.status_code == 200
    body = status.json()
    assert body["status"] == "completed"
    assert body["results"]["summary"] == {"updated": 1, "skipped": 0, "unchanged": 0}


async def test_cancel_from_preview_then_preview_is_409(client, db, seeded_user, admin_hdrs):
    resp = await upload(client, admin_hdrs, CSV)
    job_id = resp.json()["job_id"]
    cancel = await client.post(f"{base(job_id)}/cancel", headers=admin_hdrs)
    assert cancel.status_code == 204

    job = await db.get(ImportJob, uuid.UUID(job_id))
    assert job.status == "cancelled" and job.finished_at is not None

    again = await client.post(f"{base(job_id)}/cancel", headers=admin_hdrs)
    assert again.status_code == 409
    assert again.json()["detail"]["code"] == "job_not_cancellable"

    preview = await client.post(f"{base(job_id)}/preview", headers=admin_hdrs, json={})
    assert preview.status_code == 409
    assert preview.json()["detail"]["code"] == "job_not_editable"


async def test_template_and_export_formats(client, db, seeded_user, admin_hdrs):
    csv_resp = await client.get(
        "/assets/bulk-update/template?format=csv", headers=admin_hdrs)
    assert csv_resp.status_code == 200
    assert csv_resp.headers["content-disposition"] == (
        'attachment; filename="assets-update-template.csv"')

    xlsx_resp = await client.get(
        "/assets/bulk-update/template?format=xlsx", headers=admin_hdrs)
    assert xlsx_resp.status_code == 200
    assert xlsx_resp.headers["content-disposition"] == (
        'attachment; filename="assets-update-template.xlsx"')
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert wb.sheetnames == ["Assets", "Reference"]

    await mk_asset(db, 100, "SN-A", name="a")
    exp_csv = await client.get(
        "/assets/bulk-update/export?format=csv", headers=admin_hdrs)
    assert exp_csv.status_code == 200
    assert exp_csv.headers["content-disposition"] == (
        'attachment; filename="assets-export.csv"')
    assert "100" in exp_csv.text

    exp_xlsx = await client.get(
        "/assets/bulk-update/export?format=xlsx", headers=admin_hdrs)
    assert exp_xlsx.status_code == 200
    assert exp_xlsx.headers["content-disposition"] == (
        'attachment; filename="assets-export.xlsx"')
    wb2 = openpyxl.load_workbook(io.BytesIO(exp_xlsx.content))
    assert wb2.sheetnames == ["Assets", "Reference"]

    bad = await client.get(
        "/assets/bulk-update/template?format=pdf", headers=admin_hdrs)
    assert bad.status_code == 422
    assert bad.json()["detail"]["code"] == "unknown_format"


async def test_empty_upload_is_422(client, db, seeded_user, admin_hdrs):
    resp = await upload(client, admin_hdrs, "")
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"] == "empty_file"


async def test_get_asset_by_id_still_works(client, db, seeded_user, admin_hdrs):
    a = await mk_asset(db, 100, "SN-A", name="a")
    resp = await client.get(f"/assets/{a.id}", headers=admin_hdrs)
    assert resp.status_code == 200
    assert resp.json()["id"] == str(a.id)


async def test_cancel_marks_cancel_requested_and_drops_the_payload(
        client, db, seeded_user, admin_hdrs):
    a = await mk_asset(db, 100, "SN-A", name="old-a")
    resp = await upload(client, admin_hdrs, CSV + f"{a.legacy_id},,new-a,,,,,,,,,,\n")
    job_id = resp.json()["job_id"]
    await client.post(f"{base(job_id)}/commit", headers=admin_hdrs, json={"approve_all": True})

    cancel = await client.post(f"{base(job_id)}/cancel", headers=admin_hdrs)
    assert cancel.status_code == 204
    job = await db.get(ImportJob, uuid.UUID(job_id))
    await db.refresh(job)
    assert job.status == "cancelled" and job.cancel_requested is True
    assert job.payload is None
    # SQL-level check: the ORM decodes a stored JSON `null` back to Python
    # `None` on read regardless, so only this catches a bare JSONB column
    # storing Python `None` as JSON `null` instead of a true SQL NULL.
    assert await db.scalar(
        text("select payload is null from import_jobs where id = :id"),
        {"id": uuid.UUID(job_id)}) is True

    from serversherpa.db.engine import get_sessionmaker
    assert await run_once(get_sessionmaker()) is False          # never claimed
    assert (await db.scalar(select(Asset.name).where(Asset.legacy_id == 100))) == "old-a"


async def test_cancel_waits_for_a_claim_holding_the_row_then_refuses(
        client, db, seeded_user, admin_hdrs):
    """The cancel re-reads the job FOR UPDATE: while a worker's claim holds
    the row it waits, then sees `running` and refuses — it can never flip a
    job the worker already took."""
    import asyncio

    from serversherpa.db.engine import get_sessionmaker

    a = await mk_asset(db, 100, "SN-A", name="old-a")
    resp = await upload(client, admin_hdrs, CSV + f"{a.legacy_id},,new-a,,,,,,,,,,\n")
    job_id = resp.json()["job_id"]
    await client.post(f"{base(job_id)}/commit", headers=admin_hdrs, json={"approve_all": True})

    async with get_sessionmaker()() as worker_db:
        claimed = await worker_db.scalar(
            select(ImportJob).where(ImportJob.id == uuid.UUID(job_id)).with_for_update())
        claimed.status = "running"
        await worker_db.flush()                 # row locked, claim not yet committed
        cancel = asyncio.create_task(
            client.post(f"{base(job_id)}/cancel", headers=admin_hdrs))
        await asyncio.sleep(0.5)
        assert not cancel.done()                # blocked on the claim's row lock
        await worker_db.commit()
    resp = await cancel
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "job_not_cancellable"
    job = await db.get(ImportJob, uuid.UUID(job_id))
    await db.refresh(job)
    assert job.status == "running" and job.cancel_requested is False


@pytest.mark.parametrize("path", ["preview", "commit"])
@pytest.mark.parametrize("payload", [[], "x", 7, None])
async def test_a_non_object_body_is_422_invalid_json(
        client, db, seeded_user, admin_hdrs, path, payload):
    resp = await upload(client, admin_hdrs, CSV)
    job_id = resp.json()["job_id"]
    bad = await client.post(f"{base(job_id)}/{path}", headers=admin_hdrs, json=payload)
    assert bad.status_code == 422, bad.text
    assert bad.json()["detail"]["code"] == "invalid_json"


@pytest.mark.parametrize("value", ["false", "true", 1, "yes"])
async def test_approve_all_must_be_exactly_true(client, db, seeded_user, admin_hdrs, value):
    a = await mk_asset(db, 100, "SN-A", name="old-a")
    resp = await upload(client, admin_hdrs, CSV + f"{a.legacy_id},,new-a,,,,,,,,,,\n")
    job_id = resp.json()["job_id"]
    commit = await client.post(f"{base(job_id)}/commit", headers=admin_hdrs,
                               json={"approve_all": value})
    assert commit.status_code == 200, commit.text
    assert commit.json()["options"]["approve_all"] is False

    from serversherpa.db.engine import get_sessionmaker
    assert await run_once(get_sessionmaker()) is True
    job = (await client.get(base(job_id), headers=admin_hdrs)).json()
    assert job["results"]["summary"] == {"updated": 0, "skipped": 1, "unchanged": 0}
    assert (await db.scalar(select(Asset.name).where(Asset.legacy_id == 100))) == "old-a"


async def test_preview_rows_do_not_ship_the_internal_changes_map(
        client, db, seeded_user, admin_hdrs):
    a = await mk_asset(db, 100, "SN-A", name="old-a")
    resp = await upload(client, admin_hdrs, CSV + f"{a.legacy_id},,new-a,,,,,,,,,,\n")
    [row] = resp.json()["preview"]["rows"]
    assert row["diff"] == {"name": {"old": "old-a", "new": "new-a"}}
    assert "changes" not in row
    again = await client.post(f"{base(resp.json()['job_id'])}/preview",
                              headers=admin_hdrs, json={})
    assert all("changes" not in r for r in again.json()["rows"])
