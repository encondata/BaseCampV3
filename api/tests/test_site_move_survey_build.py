"""End-to-end tests for `reports/site_move_survey/__init__.py`'s
`build()` — gather -> context -> fill -> optional standards/photos
sheets -> xlsx bytes — plus the worker's error-code mapping and inbox
body for this report type. Seeds through the ORM against the real test
database and MinIO bucket, the same pattern
`test_site_move_survey_gather.py` and `test_report_worker.py` use. See
docs/superpowers/specs/2026-09-11-site-move-survey-design.md § "Report
module" (the `build` bullet).
"""

import io
import uuid
import zipfile
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from PIL import Image
from sqlalchemy import select

from serversherpa.db.models import (
    Asset, AssetModel, Attachment, Initiative, InitiativeAsset, Notification, Partner,
    Person, ReportDefinition, ReportRun, Site, SiteSurveyEntry,
)
from serversherpa.reports import worker
from serversherpa.reports.registry import get_module
from serversherpa.reports.site_move_survey import XLSX_CONTENT_TYPE, build
from serversherpa.reports.site_move_survey.gather import SurveyGatherError
from serversherpa.services.storage import get_object, put_object

FIXTURES = Path(__file__).resolve().parent / "fixtures"
TEMPLATE_BYTES = (FIXTURES / "champagne_annotated_template.xlsx").read_bytes()
STANDARDS_DOCX_BYTES = (FIXTURES / "transportation_standards.docx").read_bytes()

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _png_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (12, 8), "red").save(buf, format="PNG")
    return buf.getvalue()


PHOTO_BYTES = _png_bytes()


def _times(n):
    base = datetime(2026, 1, 1, tzinfo=UTC)
    return [base + timedelta(minutes=i) for i in range(n)]


def _partner(**kw):
    base = dict(id=uuid.uuid4(), name="Champagne Logistics", partner_types=["logistics"])
    base.update(kw)
    return Partner(**base)


def _site(**kw):
    base = dict(id=uuid.uuid4(), name="Datacenter West", address_line1="300 Origin St",
               city="Seattle", region="WA", postal_code="98101", country="US")
    base.update(kw)
    return Site(**base)


def _person(**kw):
    base = dict(id=uuid.uuid4(), first_name="Rae", last_name="Requester",
               email="rae@cumulus.example")
    base.update(kw)
    return Person(**base)


def _asset_model(**kw):
    base = dict(id=uuid.uuid4(), make="Dell", model="R740", ru_size=2,
               weight_lbs=Decimal("55.50"))
    base.update(kw)
    return AssetModel(**base)


def _asset(**kw):
    base = dict(id=uuid.uuid4(), name="db-01", serial_number="SN-001")
    base.update(kw)
    return Asset(**base)


async def _attachment(db, *, entity_type, entity_id, kind, filename, content_type,
                      storage_key, content, created_at=None):
    await put_object(storage_key, content, content_type)
    att = Attachment(entity_type=entity_type, entity_id=entity_id, kind=kind,
                     storage_key=storage_key, filename=filename, content_type=content_type,
                     size_bytes=len(content), created_at=created_at)
    db.add(att)
    await db.flush()
    return att


async def _definition(db):
    """The system "Site & Move Survey" definition — migration 0052 seeds
    it, but the test harness TRUNCATEs report_definitions before every
    test (tests/conftest.py), so look it up by report_type first and
    fall back to creating a fresh one shaped the same way, exactly as
    test_site_move_survey_gather.py's own `_run()` helper does."""
    d = await db.scalar(select(ReportDefinition)
                        .where(ReportDefinition.report_type == "site_move_survey"))
    if d is None:
        d = ReportDefinition(
            name="Site & Move Survey", report_type="site_move_survey", is_system=True,
            options={"company_name": "Cumulus Solutions Group",
                     "include_transportation_standards": True,
                     "include_site_photos": True, "condensed_assets": True})
        db.add(d)
        await db.flush()
    return d


def _run(*, definition, partner, initiative=None, requester, options):
    return ReportRun(
        definition_id=definition.id, report_type="site_move_survey",
        initiative_id=initiative.id if initiative is not None else None,
        options={"partner_id": str(partner.id), **options},
        requested_by=requester.id, requested_rank=0)


# ---------------------------------------------------------------------------
# The full happy path — a move with two assets, survey answers, a
# standards docx, and a site photo.
# ---------------------------------------------------------------------------

@pytest.fixture
async def scenario(db):
    """Everything the happy-path tests share: a logistics partner with
    the Champagne template, a contact/requester, origin+destination
    sites with survey answers and one origin photo, a move initiative
    with two different-model assets, and a definition carrying a custom
    `company_name` plus the Transportation Standards docx — so the build
    test can prove the definition's own option value (not this module's
    hardcoded default) is what lands in the workbook."""
    t0, t1 = _times(2)
    partner = _partner()
    requester = _person()
    origin = _site(name="Datacenter West", address_line1="300 Origin St", city="Ashburn",
                   region="VA", postal_code="20147")
    destination = _site(name="Datacenter East", id=uuid.uuid4(),
                        address_line1="900 Destination Ave", city="Sterling",
                        region="VA", postal_code="20164")
    db.add_all([partner, requester, origin, destination])
    await db.flush()

    db.add_all([
        SiteSurveyEntry(site_id=origin.id, field_key="dock_available", value=True),
        SiteSurveyEntry(site_id=origin.id, field_key="dock_hours", value="8am-5pm"),
        SiteSurveyEntry(site_id=destination.id, field_key="dock_available", value=False),
        SiteSurveyEntry(site_id=destination.id, field_key="dock_hours", value="24/7"),
    ])

    model_a = _asset_model(make="Dell", model="R740", ru_size=2)
    model_b = _asset_model(make="HPE", model="DL380", id=uuid.uuid4(), ru_size=4)
    asset1 = _asset(name="db-01", serial_number="SN-001")
    asset2 = _asset(name="db-02", serial_number="SN-002", id=uuid.uuid4())
    db.add_all([model_a, model_b, asset1, asset2])
    await db.flush()
    asset1.model_id = model_a.id
    asset2.model_id = model_b.id

    initiative = Initiative(name="Champagne Move", initiative_type="move", status="planned",
                            origin_site_id=origin.id, destination_site_id=destination.id)
    db.add(initiative)
    await db.flush()

    db.add_all([
        InitiativeAsset(initiative_id=initiative.id, asset_id=asset1.id,
                        source_rack="RACK-12", source_ru=Decimal("20")),
        InitiativeAsset(initiative_id=initiative.id, asset_id=asset2.id,
                        source_rack="RACK-13", source_ru=Decimal("5")),
    ])
    await db.flush()

    await _attachment(db, entity_type="partner", entity_id=partner.id, kind="survey_template",
                      filename="champagne_annotated_template.xlsx", content_type=XLSX_MIME,
                      storage_key=f"test/sms-build/{partner.id}/template.xlsx",
                      content=TEMPLATE_BYTES, created_at=t0)

    await _attachment(db, entity_type="site", entity_id=origin.id, kind="photo",
                      filename="dock.png", content_type="image/png",
                      storage_key=f"test/sms-build/{origin.id}/dock.png",
                      content=PHOTO_BYTES, created_at=t0)

    definition = await _definition(db)
    definition.options = {**definition.options, "company_name": "Acme Test Co"}
    await _attachment(db, entity_type="report_definition", entity_id=definition.id,
                      kind="report_asset", filename="Transportation Standards.docx",
                      content_type=DOCX_MIME,
                      storage_key=f"test/sms-build/{definition.id}/standards.docx",
                      content=STANDARDS_DOCX_BYTES, created_at=t1)
    await db.commit()

    return {"partner": partner, "requester": requester, "origin": origin,
           "destination": destination, "initiative": initiative, "definition": definition}


async def test_build_fills_workbook_per_asset(db, scenario):
    run = _run(definition=scenario["definition"], partner=scenario["partner"],
              initiative=scenario["initiative"], requester=scenario["requester"],
              options={"condensed_assets": False})
    db.add(run)
    await db.flush()

    result = await build(db, run)

    assert result.content_type == XLSX_CONTENT_TYPE == XLSX_MIME
    assert result.filename.startswith("Site & Move Survey - Champagne Logistics - ")
    assert result.filename.endswith(".xlsx")

    wb = openpyxl.load_workbook(io.BytesIO(result.content))
    info = wb["Customer and Site Information"]
    # The definition's OWN company_name option, not this module's
    # hardcoded default — proves build() merges the definition's saved
    # options under the run's, not the other way around.
    assert info["C11"].value == "Acme Test Co"
    assert info["C12"].value == scenario["requester"].display_name
    assert info["C18"].value == "Datacenter West"
    assert info["C21"].value == "Ashburn, VA"
    assert info["C28"].value == "Datacenter East"

    general = wb["General Questions"]
    assert general["D21"].value == "yes"          # origin.survey.dock_available
    assert general["E21"].value == "no"           # destination.survey.dock_available
    assert general["C22"].value == "Origin: 8am-5pm  |  Destination: 24/7"

    equip = wb["Equipment Listing"]
    row8 = (equip["B8"].value, equip["C8"].value, equip["D8"].value, equip["F8"].value)
    row9 = (equip["B9"].value, equip["C9"].value, equip["D9"].value, equip["F9"].value)
    rows = {row8, row9}
    assert ("RACK-12 U20", "Dell", "R740", 1) in rows
    assert ("RACK-13 U5", "HPE", "DL380", 1) in rows

    assert "Transportation Standards" in wb.sheetnames
    assert "Site Photos" in wb.sheetnames


async def test_build_condenses_assets_by_make_model_with_qty(db, scenario):
    run = _run(definition=scenario["definition"], partner=scenario["partner"],
              initiative=scenario["initiative"], requester=scenario["requester"],
              options={"condensed_assets": True})
    db.add(run)
    await db.flush()

    result = await build(db, run)
    wb = openpyxl.load_workbook(io.BytesIO(result.content))
    equip = wb["Equipment Listing"]
    # Two different (make, model) pairs -> one condensed row each, each
    # with its own qty (1) — the fixture's two assets never share a model,
    # so this proves the condensed code path runs (grouping itself is
    # covered by test_site_move_survey_fill.py's own unit tests).
    makes_models_qty = {(equip["C8"].value, equip["D8"].value, equip["F8"].value),
                        (equip["C9"].value, equip["D9"].value, equip["F9"].value)}
    assert makes_models_qty == {("Dell", "R740", 1), ("HPE", "DL380", 1)}
    # Condensed rows carry no per-asset rack/serial.
    assert equip["B8"].value in (None, "") and equip["B9"].value in (None, "")


async def test_build_toggle_off_removes_transport_sheet_and_skips_photos(db, scenario):
    run = _run(definition=scenario["definition"], partner=scenario["partner"],
              initiative=scenario["initiative"], requester=scenario["requester"],
              options={"include_transportation_standards": False,
                      "include_site_photos": False})
    db.add(run)
    await db.flush()

    result = await build(db, run)
    wb = openpyxl.load_workbook(io.BytesIO(result.content))
    assert "Transportation Standards" not in wb.sheetnames
    assert "Site Photos" not in wb.sheetnames


# ---------------------------------------------------------------------------
# No initiative — partner + manually chosen sites, zero assets
# ---------------------------------------------------------------------------

async def test_build_without_initiative_has_no_asset_rows_and_writes_notes(db):
    t0, = _times(1)
    partner = _partner()
    requester = _person()
    origin = _site(name="Standalone Origin")
    destination = _site(name="Standalone Destination", id=uuid.uuid4())
    db.add_all([partner, requester, origin, destination])
    await db.flush()

    await _attachment(db, entity_type="partner", entity_id=partner.id, kind="survey_template",
                      filename="template.xlsx", content_type=XLSX_MIME,
                      storage_key=f"test/sms-build/{partner.id}/t.xlsx",
                      content=TEMPLATE_BYTES, created_at=t0)

    definition = await _definition(db)
    await db.commit()

    run = _run(definition=definition, partner=partner, initiative=None, requester=requester,
              options={"source_site_id": str(origin.id),
                      "destination_site_id": str(destination.id),
                      "asset_notes": "ships separately, handle with care"})
    db.add(run)
    await db.flush()

    result = await build(db, run)
    wb = openpyxl.load_workbook(io.BytesIO(result.content))
    equip = wb["Equipment Listing"]
    assert equip["G8"].value == "ships separately, handle with care"
    assert equip["A8"].value in (None, "")             # blanked, not an asset index

    info = wb["Customer and Site Information"]
    assert info["C18"].value == "Standalone Origin"
    assert info["C28"].value == "Standalone Destination"


# ---------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------

async def test_missing_template_raises_no_survey_template(db):
    partner = _partner()
    requester = _person()
    db.add_all([partner, requester])
    await db.flush()
    definition = await _definition(db)
    await db.commit()

    run = _run(definition=definition, partner=partner, requester=requester, options={})
    db.add(run)
    await db.flush()

    with pytest.raises(SurveyGatherError) as exc_info:
        await build(db, run)
    assert exc_info.value.code == "no_survey_template"


async def test_corrupt_template_raises_template_unreadable(db):
    t0, = _times(1)
    partner = _partner()
    requester = _person()
    db.add_all([partner, requester])
    await db.flush()

    await _attachment(db, entity_type="partner", entity_id=partner.id, kind="survey_template",
                      filename="template.xlsx", content_type=XLSX_MIME,
                      storage_key=f"test/sms-build/{partner.id}/corrupt.xlsx",
                      content=b"not actually a zip file", created_at=t0)

    definition = await _definition(db)
    await db.commit()

    run = _run(definition=definition, partner=partner, requester=requester, options={})
    db.add(run)
    await db.flush()

    with pytest.raises(SurveyGatherError) as exc_info:
        await build(db, run)
    assert exc_info.value.code == "template_unreadable"


# ---------------------------------------------------------------------------
# Registry wiring
# ---------------------------------------------------------------------------

def test_registered_under_its_report_type():
    assert get_module("site_move_survey").report_type == "site_move_survey"


# ---------------------------------------------------------------------------
# Worker: real end-to-end run with no initiative — xlsx content type and
# an inbox body naming the partner instead of "—".
# ---------------------------------------------------------------------------

async def test_worker_end_to_end_no_initiative_names_partner_in_inbox(db):
    from serversherpa.db.engine import get_sessionmaker

    t0, = _times(1)
    partner = _partner(name="Champagne Logistics")
    requester = _person()
    origin = _site(name="Standalone Origin")
    db.add_all([partner, requester, origin])
    await db.flush()

    await _attachment(db, entity_type="partner", entity_id=partner.id, kind="survey_template",
                      filename="template.xlsx", content_type=XLSX_MIME,
                      storage_key=f"test/sms-build/{partner.id}/t.xlsx",
                      content=TEMPLATE_BYTES, created_at=t0)

    definition = await _definition(db)
    await db.commit()

    run = _run(definition=definition, partner=partner, initiative=None, requester=requester,
              options={"source_site_id": str(origin.id)})
    run.notify = True
    db.add(run)
    await db.commit()

    assert await worker.run_once(get_sessionmaker()) is True

    # `run` is still this session's identity-mapped instance for this row
    # (expire_on_commit=False, db/engine.py) — a plain `db.get`/`db.execute`
    # would just hand back the pre-worker cached copy, since the worker
    # commits its changes through entirely separate sessions. `refresh`
    # forces the reload.
    await db.refresh(run)
    assert run.status == "completed"
    assert run.storage_key == f"reports/standalone/{run.id}.xlsx"
    assert run.filename.startswith("Site & Move Survey - Champagne Logistics - ")
    stored = await get_object(run.storage_key)
    zipfile.ZipFile(io.BytesIO(stored))                 # a real xlsx, not fake bytes
    assert run.attachment_id is None                    # no initiative to attach to

    n = await db.scalar(select(Notification).where(Notification.person_id == requester.id))
    assert n is not None and n.kind == "report_ready"
    assert n.body == "Champagne Logistics"              # not the literal "—"


async def test_worker_maps_gather_error_codes_onto_run_error(db):
    from serversherpa.db.engine import get_sessionmaker

    partner = _partner()
    requester = _person()
    db.add_all([partner, requester])
    await db.flush()
    definition = await _definition(db)
    await db.commit()

    run = _run(definition=definition, partner=partner, requester=requester, options={})
    db.add(run)
    await db.commit()

    assert await worker.run_once(get_sessionmaker()) is True
    await db.refresh(run)                    # see the note in the test above
    assert run.status == "failed" and run.error == "no_survey_template"
