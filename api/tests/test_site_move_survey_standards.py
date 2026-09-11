"""Tests for the Site & Move Survey xlsx's two optional sheets:
`reports/site_move_survey/standards.py` (the stdlib docx parser over the
reconstructed `transportation_standards.docx` fixture, and the
"Transportation Standards" sheet writer) and `photos.py` (the "Site
Photos" sheet writer). Pure Python — no DB. See
docs/superpowers/specs/2026-09-11-site-move-survey-design.md § "Report
module" and the V2 reference (read-only) at
/Users/jrh1812/Developer/BaseCampV2-reference/api/reports/site_move_survey.py
for the behavior this ports (`_parse_standards_docx` /
`_append_transportation_standards` / `_append_site_photos`).
"""

import io
from pathlib import Path

import openpyxl
from PIL import Image as PILImage

from serversherpa.reports.site_move_survey import photos, standards

FIXTURE = Path(__file__).parent / "fixtures" / "transportation_standards.docx"


def _png_bytes(color=(200, 30, 30), size=(10, 10)) -> bytes:
    buf = io.BytesIO()
    PILImage.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------------------
# parse_standards_docx — the reconstructed fixture
# ---------------------------------------------------------------------------

def test_parses_expected_item_counts_from_the_fixture():
    items = standards.parse_standards_docx(FIXTURE.read_bytes())
    kinds = [kind for kind, _ in items]
    assert kinds.count("heading") == 7
    assert kinds.count("bullet") == 50
    assert kinds.count("text") == 2
    assert kinds.count("image") == 6


def test_first_item_is_the_expected_heading():
    items = standards.parse_standards_docx(FIXTURE.read_bytes())
    assert items[0] == ("heading", "Transportation Partner Requirements")


def test_bullet_levels_are_present_as_leading_indent():
    """V2 bakes the list level into the payload string itself (`"    "
    * level + "• " + text`) rather than carrying it as separate data —
    see standards.py's docstring. A level-0 bullet starts directly with
    the marker; a deeper one is indented first."""
    items = standards.parse_standards_docx(FIXTURE.read_bytes())
    bullets = [payload for kind, payload in items if kind == "bullet"]
    assert any(b.startswith("• ") for b in bullets), "expected at least one level-0 bullet"
    assert any(b.startswith("    ") for b in bullets), "expected at least one indented bullet"
    # Every bullet still carries the marker somewhere in its indent prefix.
    assert all("• " in b for b in bullets)


def test_images_are_extracted_as_raw_bytes_decodable_by_pillow():
    items = standards.parse_standards_docx(FIXTURE.read_bytes())
    images = [payload for kind, payload in items if kind == "image"]
    assert len(images) == 6
    for raw in images:
        assert isinstance(raw, bytes)
        PILImage.open(io.BytesIO(raw)).load()  # raises if not a real image


def test_parse_standards_cached_reuses_the_result_for_the_same_attachment_id(monkeypatch):
    calls = []
    original_parse = standards.parse_standards_docx

    def _spy_parse(docx_bytes):
        calls.append(docx_bytes)
        return original_parse(docx_bytes)

    monkeypatch.setattr(standards, "parse_standards_docx", _spy_parse)
    docx_bytes = FIXTURE.read_bytes()

    first = standards.parse_standards_cached("att-1", docx_bytes)
    second = standards.parse_standards_cached("att-1", docx_bytes)
    assert first is second
    assert len(calls) == 1

    # A different attachment id is parsed independently.
    third = standards.parse_standards_cached("att-2", docx_bytes)
    assert len(calls) == 2
    assert third == first


# ---------------------------------------------------------------------------
# append_standards_sheet
# ---------------------------------------------------------------------------

def test_append_standards_sheet_writes_headings_bullets_text_and_images():
    wb = openpyxl.Workbook()
    items = [
        ("heading", "Section One"),
        ("bullet", "• First point"),
        ("bullet", "    • Nested point"),
        ("text", "Some free-form paragraph."),
        ("image", _png_bytes()),
        ("image", _png_bytes(color=(30, 120, 200))),
    ]

    standards.append_standards_sheet(wb, items)

    assert "Transportation Standards" in wb.sheetnames
    ws = wb["Transportation Standards"]

    cell_values = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert "Section One" in cell_values
    assert "• First point" in cell_values
    assert "    • Nested point" in cell_values
    assert "Some free-form paragraph." in cell_values

    heading_row = next(r for r in range(2, ws.max_row + 1)
                       if ws.cell(row=r, column=2).value == "Section One")
    assert ws.cell(row=heading_row, column=2).font.bold is True

    assert len(ws._images) == 2


def test_append_standards_sheet_skips_a_corrupt_image_without_raising():
    wb = openpyxl.Workbook()
    items = [("heading", "Intro"), ("image", b"not a real image"), ("text", "After the bad image.")]

    standards.append_standards_sheet(wb, items)

    ws = wb["Transportation Standards"]
    assert len(ws._images) == 0
    cell_values = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert "Intro" in cell_values
    assert "After the bad image." in cell_values


def test_append_standards_sheet_over_the_real_fixture_places_all_six_images():
    wb = openpyxl.Workbook()
    items = standards.parse_standards_docx(FIXTURE.read_bytes())

    standards.append_standards_sheet(wb, items)

    ws = wb["Transportation Standards"]
    assert len(ws._images) == 6
    # Every heading text made it into some cell in column B.
    heading_texts = {payload for kind, payload in items if kind == "heading"}
    cell_texts = {ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
    assert heading_texts <= cell_texts


# ---------------------------------------------------------------------------
# photos.append_site_photos
# ---------------------------------------------------------------------------

def test_append_site_photos_writes_a_heading_per_site_and_all_images():
    wb = openpyxl.Workbook()
    entries = [
        ("Datacenter West", [_png_bytes(), _png_bytes(color=(0, 200, 0))]),
        ("Datacenter East", [_png_bytes(color=(0, 0, 200))]),
    ]

    photos.append_site_photos(wb, entries)

    assert "Site Photos" in wb.sheetnames
    ws = wb["Site Photos"]
    col_a_values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert "Datacenter West" in col_a_values
    assert "Datacenter East" in col_a_values
    west_row = col_a_values.index("Datacenter West") + 1
    assert ws.cell(row=west_row, column=1).font.bold is True
    assert len(ws._images) == 3


def test_append_site_photos_caps_at_whatever_the_caller_passed_in():
    """The 10-per-site cap is `gather.py`'s job (it queries with
    `LIMIT 10`); this module just draws whatever list it's handed, so
    feeding it 10 images places all 10."""
    wb = openpyxl.Workbook()
    entries = [("Datacenter West", [_png_bytes() for _ in range(10)])]

    photos.append_site_photos(wb, entries)

    ws = wb["Site Photos"]
    assert len(ws._images) == 10


def test_append_site_photos_skips_sites_with_no_images_and_a_corrupt_one():
    wb = openpyxl.Workbook()
    entries = [
        ("Empty Site", []),
        ("Real Site", [_png_bytes(), b"not a real image"]),
    ]

    photos.append_site_photos(wb, entries)

    ws = wb["Site Photos"]
    col_a_values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert "Empty Site" not in col_a_values
    assert "Real Site" in col_a_values
    assert len(ws._images) == 1


def test_append_site_photos_creates_no_sheet_when_every_entry_is_empty():
    wb = openpyxl.Workbook()
    photos.append_site_photos(wb, [("Empty Site", [])])
    assert "Site Photos" not in wb.sheetnames
