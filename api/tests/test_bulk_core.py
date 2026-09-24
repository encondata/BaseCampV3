"""The shared bulk-import core (imports/bulk.py): parsing, numbering,
templates. Content-agnostic — every importer passes its own column list."""
import io

import openpyxl
import pytest

from serversherpa.imports import bulk

COLS = ["name", "city", "notes"]


def test_cell_normalizes_scalars_and_strips_one_guard_quote():
    assert bulk.cell(None) == ""
    assert bulk.cell(89501.0) == "89501"
    assert bulk.cell(1.5) == "1.5"
    assert bulk.cell("  x  ") == "x"
    assert bulk.cell("'=SUM(A1)") == "=SUM(A1)"
    assert bulk.cell("''=x") == "''=x"         # only a quote right before a formula lead comes off


def test_guard_cell_prefixes_formula_leads_but_not_numbers():
    assert bulk.guard_cell("=HYPERLINK(x)") == "'=HYPERLINK(x)"
    assert bulk.guard_cell("-119.8") == "-119.8"
    assert bulk.guard_cell("+1") == "+1"
    assert bulk.guard_cell("plain") == "plain"


def test_check_columns_reports_unknowns_sorted():
    bulk.check_columns(["name", "city"], COLS)
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.check_columns(["zip", "name", "aaa"], COLS)
    assert exc.value.code == "unknown_columns"
    assert exc.value.extra["columns"] == ["aaa", "zip"]


def test_number_json_rows_numbers_from_one_and_drops_blank_lines():
    rows = bulk.number_json_rows(
        [{"name": "A"}, {"name": "", "city": ""}, {"city": "Reno"}], COLS)
    assert rows == [(1, {"name": "A", "city": "", "notes": ""}),
                    (3, {"name": "", "city": "Reno", "notes": ""})]
    assert bulk.number_json_rows({"name": "One"}, COLS)[0][0] == 1
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.number_json_rows("nope", COLS)  # type: ignore[arg-type]
    assert exc.value.code == "invalid_json"
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.number_json_rows([{"name": str(i)} for i in range(bulk.MAX_ROWS + 1)], COLS)
    assert exc.value.code == "too_many_rows"


def test_parse_upload_csv_numbers_from_two_and_handles_bom():
    rows = bulk.parse_upload("t.csv", "﻿name,city\nA,Reno\n".encode(), COLS, "Sheet")
    assert rows == [(2, {"name": "A", "city": "Reno", "notes": ""})]


def test_parse_upload_xlsx_prefers_named_sheet_and_skips_blank_headers():
    wb = openpyxl.Workbook()
    other = wb.active
    other.title = "Other"
    other.append(["name"])
    other.append(["Wrong"])
    ws = wb.create_sheet("People")
    ws.append(["name", "", "city"])
    ws.append(["Right", "spacer", "Reno"])
    buf = io.BytesIO()
    wb.save(buf)
    rows = bulk.parse_upload("t.xlsx", buf.getvalue(), COLS, "People")
    assert rows == [(2, {"name": "Right", "city": "Reno", "notes": ""})]


def test_parse_upload_honors_a_per_tool_max_rows_and_reports_it_as_the_limit():
    csv_bytes = b"name,city\nA,Reno\nB,Reno\nC,Reno\n"
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.parse_upload("t.csv", csv_bytes, COLS, "Sheet", max_rows=2)
    assert exc.value.code == "too_many_rows"
    assert exc.value.extra["limit"] == 2
    # the default (1,000) still allows the same 3-row file with no override
    rows = bulk.parse_upload("t.csv", csv_bytes, COLS, "Sheet")
    assert len(rows) == 3


def _xlsx(header: list[str], lines: list[list[str]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(header)
    for line in lines:
        ws.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_an_oversized_file_is_too_many_rows_and_read_only_one_row_past_the_limit(
        fmt, monkeypatch):
    lines = [[f"N{i}", "Reno"] for i in range(10)]
    content = (_xlsx(["name", "city"], lines) if fmt == "xlsx" else
               ("name,city\n" + "".join(f"{a},{b}\n" for a, b in lines)).encode())
    seen: list[int] = []
    real = bulk.numbered

    def spy(rows, first_row, columns, max_rows=bulk.MAX_ROWS):
        seen.append(len(rows))
        return real(rows, first_row, columns, max_rows=max_rows)

    monkeypatch.setattr(bulk, "numbered", spy)
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.parse_upload(f"t.{fmt}", content, COLS, "People", max_rows=3)
    assert exc.value.code == "too_many_rows"
    assert exc.value.extra == {"limit": 3}
    assert seen == [4]                    # stopped one row past the limit, not all 10
    # at the limit exactly, every row is still read and returned
    seen.clear()
    rows = bulk.parse_upload(f"t.{fmt}", content, COLS, "People", max_rows=10)
    assert seen == [10] and len(rows) == 10


def test_parse_upload_xlsx_over_the_default_limit_is_too_many_rows():
    content = _xlsx(["name"], [[f"N{i}"] for i in range(bulk.MAX_ROWS + 5)])
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.parse_upload("t.xlsx", content, COLS, "People")
    assert exc.value.code == "too_many_rows"
    assert exc.value.extra == {"limit": bulk.MAX_ROWS}


def test_parse_upload_honors_a_per_tool_max_bytes():
    csv_bytes = b"name,city\nA,Reno\n"
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.parse_upload("t.csv", csv_bytes, COLS, "Sheet", max_bytes=5)
    assert exc.value.code == "file_too_large"
    assert exc.value.extra["limit"] == 5


def test_numbered_and_number_json_rows_honor_a_per_tool_max_rows():
    rows = [{"name": str(i)} for i in range(3)]
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.numbered(rows, 1, COLS, max_rows=2)
    assert exc.value.code == "too_many_rows"
    assert exc.value.extra["limit"] == 2
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.number_json_rows(rows, COLS, max_rows=2)
    assert exc.value.code == "too_many_rows"
    assert exc.value.extra["limit"] == 2


def test_parse_upload_error_codes():
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.parse_upload("x.txt", b"hi", COLS, "S")
    assert exc.value.code == "unsupported_file"
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.parse_upload("x.json", b"{bad", COLS, "S")
    assert exc.value.code == "invalid_json"
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.parse_upload("x.xlsx", b"not a workbook", COLS, "S")
    assert exc.value.code == "invalid_xlsx"
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.parse_upload("x.csv", b"\xff\xfe\x00bad", COLS, "S")   # not utf-8
    assert exc.value.code == "invalid_csv"
    with pytest.raises(bulk.BulkImportError) as exc:
        bulk.parse_upload("big.csv", b"x" * (bulk.MAX_BYTES + 1), COLS, "S")
    assert exc.value.code == "file_too_large"


def test_csv_and_xlsx_builders_guard_formulas_and_round_trip():
    rows = [{"name": "=EVIL()", "city": "Reno", "notes": "-5"}]
    csv_text = bulk.build_rows_csv(rows, COLS)
    assert csv_text.splitlines() == ["name,city,notes", "'=EVIL(),Reno,-5"]
    assert bulk.parse_upload("t.csv", csv_text.encode(), COLS, "S")[0][1]["name"] == "=EVIL()"

    blob = bulk.build_rows_xlsx(rows, COLS, "People",
                                [("Valid levels", ["L1", "L2"]), ("Valid statuses", ["active"])])
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    assert wb.sheetnames == ["People", "Reference"]
    assert wb["People"]["A2"].data_type == "s"           # pinned to text
    ref = [row[0].value for row in wb["Reference"].iter_rows()]
    assert ref == ["Valid levels", "L1", "L2", None, "Valid statuses", "active"]
    assert bulk.parse_upload("t.xlsx", blob, COLS, "People")[0][1]["name"] == "=EVIL()"
