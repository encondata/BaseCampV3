"""Move-asset import file parsing: CSV/XLSX -> numbered rows, V2 headers."""

import io

import openpyxl
import pytest

from serversherpa.imports.parsing import (
    CANONICAL, MAX_BYTES, SAMPLE_ROWS, TEMPLATE_HEADERS, ImportFileError,
    build_template_csv, build_template_xlsx, parse_upload,
)

CSV = (
    "Serial Number,Asset Name,Asset Make,Asset Model,Source RU,Data 1,"
    "Mgmt 1,Vendor Involvment,Extra Col\n"
    "SN-1,web-01,Dell,R740,12,sw1:e1,m1,,note\n"
    ",,,,,,,,\n"
    "SN-2,db-01,HPE,DL380,3.5,,,yes,\n"
).encode()


def _xlsx(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_csv_parses_to_canonical_and_raw():
    rows = parse_upload("ft.csv", CSV)
    assert [n for n, _, _ in rows] == [2, 4]      # blank line 3 skipped
    n, canonical, raw = rows[0]
    assert canonical["serial_number"] == "SN-1"
    assert canonical["asset_name"] == "web-01"
    assert canonical["asset_make"] == "Dell"
    assert canonical["source_ru"] == "12"
    assert canonical["data_1"] == "sw1:e1"
    assert canonical["mgmt_1"] == "m1"
    assert canonical["vendor_involvement"] == ""
    assert raw["Extra Col"] == "note"             # unknown column preserved
    assert set(canonical) == set(CANONICAL)
    # v2's misspelling maps too
    assert rows[1][1]["vendor_involvement"] == "yes"


def test_headers_case_insensitive():
    content = b"SERIAL NUMBER,asset name\nsn-9,x\n"
    [(_, canonical, _)] = parse_upload("a.csv", content)
    assert canonical["serial_number"] == "sn-9"
    assert canonical["asset_name"] == "x"


def test_source_and_destination_position_headers_map():
    content = (b"Serial Number,Source Position,Destination Position\n"
              b"sn-1,Front,Rear\n")
    [(_, canonical, _)] = parse_upload("a.csv", content)
    assert canonical["source_position"] == "Front"
    assert canonical["destination_position"] == "Rear"


def test_xlsx_matches_csv():
    content = _xlsx([["Serial Number", "Asset Name", "Source RU"],
                     ["SN-1", "web-01", 12]])
    [(n, canonical, raw)] = parse_upload("ft.xlsx", content)
    assert n == 2
    assert canonical["serial_number"] == "SN-1"
    assert canonical["source_ru"] == "12"         # 12.0 -> "12"


def test_missing_serial_column_is_a_file_error():
    with pytest.raises(ImportFileError) as e:
        parse_upload("a.csv", b"Asset Name\nweb-01\n")
    assert e.value.code == "missing_serial_column"


def test_unsupported_and_invalid_files():
    with pytest.raises(ImportFileError) as e:
        parse_upload("a.txt", b"x")
    assert e.value.code == "unsupported_file"
    with pytest.raises(ImportFileError) as e:
        parse_upload("a.xlsx", b"not a zip")
    assert e.value.code == "invalid_xlsx"
    with pytest.raises(ImportFileError) as e:
        parse_upload("a.csv", b"x" * (MAX_BYTES + 1))
    assert e.value.code == "file_too_large"


def test_template_csv_headers_and_roundtrip():
    text = build_template_csv()
    assert text.splitlines()[0] == ",".join(TEMPLATE_HEADERS)
    rows = parse_upload("t.csv", text.encode())
    assert len(rows) == len(SAMPLE_ROWS)
    assert rows[0][1]["serial_number"] == SAMPLE_ROWS[0]["Serial Number"]


def test_template_xlsx_has_reference_sheet():
    wb = openpyxl.load_workbook(io.BytesIO(build_template_xlsx()))
    assert wb.sheetnames == ["Move Assets", "Reference"]
    header = [c.value for c in next(wb["Move Assets"].iter_rows(max_row=1))]
    assert header == TEMPLATE_HEADERS
    rows = parse_upload("t.xlsx", build_template_xlsx())
    assert len(rows) == len(SAMPLE_ROWS)


def test_xlsx_blank_header_cell_does_not_shift_columns():
    content = _xlsx([["Serial Number", "", "Asset Name"],
                     ["SN-1", "spacer", "web-01"]])
    [(_, canonical, raw)] = parse_upload("ft.xlsx", content)
    assert canonical["serial_number"] == "SN-1"
    assert canonical["asset_name"] == "web-01"     # not shifted into the spacer
    assert list(raw.values()) == ["SN-1", "web-01"]


def test_pod_headers_map():
    content = (b"Serial Number,Source Pod #,Destination Pod Number\n"
               b"sn-1,14,9\n")
    [(_, canonical, _)] = parse_upload("a.csv", content)
    assert canonical["source_pod"] == "14"
    assert canonical["destination_pod"] == "9"


def test_bare_pod_header_is_the_source_pod():
    content = b"Serial Number,Pod #\nsn-1,14\n"
    [(_, canonical, _)] = parse_upload("a.csv", content)
    assert canonical["source_pod"] == "14"
    assert canonical["destination_pod"] == ""


def test_template_places_pod_before_rack():
    assert TEMPLATE_HEADERS.index("Source Pod") + 1 == \
        TEMPLATE_HEADERS.index("Source Rack")
    assert TEMPLATE_HEADERS.index("Destination Pod") + 1 == \
        TEMPLATE_HEADERS.index("Destination Rack")
    for sample in SAMPLE_ROWS:
        assert set(sample) == set(TEMPLATE_HEADERS)
