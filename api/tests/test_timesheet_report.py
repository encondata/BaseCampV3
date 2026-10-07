"""Timesheet report module: options validation, run-option resolution,
Excel and PDF output, build() filenames. Data layer is test_timesheet_gather."""

import uuid
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import openpyxl
import pytest

from serversherpa.db.models import Person, ReportDefinition, ReportRun, TimeEntry
from serversherpa.reports import timesheet
from serversherpa.reports.move_scan_history.xlsx import XLSX_MIME
from serversherpa.reports.registry import OptionsError, get_module
from serversherpa.reports.timesheet import pdf as ts_pdf
from serversherpa.reports.timesheet import xlsx as ts_xlsx
from serversherpa.reports.timesheet.gather import (
    EntryRow,
    TimesheetData,
    TimesheetFilters,
    day_rows,
    flags_for,
    rollups,
)

NY = ZoneInfo("America/New_York")
NOW = datetime(2026, 10, 6, 20, 0, tzinfo=UTC)
ANN, BOB = uuid.uuid4(), uuid.uuid4()
JOB = uuid.uuid4()
GOOD = {"from": "2026-10-01", "to": "2026-10-31"}


# ------------------------------------------------------------ options (pure)

def test_report_type_and_registration():
    assert timesheet.report_type == "timesheet"
    assert get_module("timesheet") is timesheet


def test_default_options():
    assert timesheet.default_options() == {
        "default_format": "xlsx", "default_views": ["day", "punch"],
        "default_statuses": ["approved", "pending"]}


def test_validate_options_fills_defaults_and_checks_domains():
    assert timesheet.validate_options({}) == timesheet.default_options()
    assert timesheet.validate_options({"default_format": "pdf"})["default_format"] == "pdf"
    for bad in ({"bogus": 1}, {"default_format": "csv"}, {"default_views": []},
                {"default_views": ["week"]}, {"default_statuses": []},
                {"default_statuses": ["approved", "paid"]},
                {"default_views": "day"}):
        with pytest.raises(OptionsError):
            timesheet.validate_options(bad)


def test_validate_run_options_good_input_passes_through_unchanged():
    opts = {**GOOD, "person_id": str(ANN), "site_id": str(uuid.uuid4()),
            "statuses": ["approved", "rejected"], "views": ["day"], "format": "pdf"}
    assert timesheet.validate_run_options(opts) == opts
    # nothing is defaulted in
    assert timesheet.validate_run_options(GOOD) == GOOD


def test_validate_run_options_drops_empty_filters():
    out = timesheet.validate_run_options({**GOOD, "person_id": None, "site_id": ""})
    assert out == GOOD


def test_validate_run_options_rejects_unknown_keys():
    with pytest.raises(OptionsError) as exc:
        timesheet.validate_run_options({**GOOD, "colour": "red"})
    assert "unknown option 'colour'" in exc.value.problems


def test_validate_run_options_non_dict_and_bad_members_are_problems_not_errors():
    for bad in (None, [], "from=2026-10-01", 5):
        with pytest.raises(OptionsError):
            timesheet.validate_run_options(bad)
    for key, value in (("statuses", [["approved"]]), ("views", [{"a": 1}]),
                       ("statuses", [None]), ("views", [1])):
        with pytest.raises(OptionsError, match=key):
            timesheet.validate_run_options({**GOOD, key: value})
    for key in ("from", "to"):
        with pytest.raises(OptionsError):
            timesheet.validate_run_options({**GOOD, key: ["2026-10-01"]})
    with pytest.raises(OptionsError, match="person_id"):
        timesheet.validate_run_options({**GOOD, "person_id": {"x": 1}})
    with pytest.raises(OptionsError):
        timesheet.validate_options(["x"])
    for key in ("default_views", "default_statuses"):
        with pytest.raises(OptionsError):
            timesheet.validate_options({key: [["day"]]})


def test_validate_run_options_lists_every_problem():
    with pytest.raises(OptionsError) as exc:
        timesheet.validate_run_options({
            "person_id": "nope", "statuses": [], "views": [], "format": "csv"})
    text = " ".join(exc.value.problems)
    for needle in ("'from'", "'to'", "person_id", "statuses", "views", "format"):
        assert needle in text


def test_validate_run_options_date_rules():
    with pytest.raises(OptionsError, match="'from' must be on or before 'to'"):
        timesheet.validate_run_options({"from": "2026-10-31", "to": "2026-10-01"})
    with pytest.raises(OptionsError, match="YYYY-MM-DD"):
        timesheet.validate_run_options({"from": "10/01/2026", "to": "2026-10-31"})
    # 366 days inclusive is the limit
    ok = timesheet.validate_run_options({"from": "2026-01-01", "to": "2027-01-01"})
    assert ok["to"] == "2027-01-01"
    with pytest.raises(OptionsError, match="366"):
        timesheet.validate_run_options({"from": "2026-01-01", "to": "2027-01-02"})


def test_validate_run_options_rejects_unknown_status_and_view():
    with pytest.raises(OptionsError, match="statuses"):
        timesheet.validate_run_options({**GOOD, "statuses": ["approved", "paid"]})
    with pytest.raises(OptionsError, match="views"):
        timesheet.validate_run_options({**GOOD, "views": ["week"]})


class _Run:
    def __init__(self, options, initiative_id=None):
        self.options = options
        self.initiative_id = initiative_id


def test_parse_run_merge_order_run_over_definition_over_defaults():
    # defaults only
    f, views, fmt = timesheet.parse_run(_Run(GOOD), {})
    assert (f.from_day, f.to_day) == (date(2026, 10, 1), date(2026, 10, 31))
    assert f.statuses == ("approved", "pending")
    assert views == ["day", "punch"] and fmt == "xlsx"
    assert f.person_id is None and f.site_id is None and f.initiative_id is None
    # definition beats defaults
    defn = {"default_format": "pdf", "default_views": ["punch"],
            "default_statuses": ["approved"]}
    f, views, fmt = timesheet.parse_run(_Run(GOOD), defn)
    assert (f.statuses, views, fmt) == (("approved",), ["punch"], "pdf")
    # run beats definition
    run = _Run({**GOOD, "statuses": ["open", "rejected"], "views": ["day"],
                "format": "xlsx", "person_id": str(ANN)}, initiative_id=JOB)
    f, views, fmt = timesheet.parse_run(run, defn)
    assert (f.statuses, views, fmt) == (("open", "rejected"), ["day"], "xlsx")
    assert f.person_id == ANN and f.initiative_id == JOB


def test_parse_run_rejects_a_run_without_dates():
    with pytest.raises(OptionsError):
        timesheet.parse_run(_Run({}), {})


# ------------------------------------------------------------ sample data

def _row(*, person=ANN, name="Ann Smith", day=date(2026, 10, 1), start_h=8,
         end_h=16.0, status="approved", source="punch", adjusted=False,
         job=JOB, job_name="Alpha", site="HQ", tz="America/New_York",
         open_=False, reason=None, notes=None, approver=None,
         break_minutes=0) -> EntryRow:
    zone = ZoneInfo(tz)
    start = datetime(day.year, day.month, day.day, tzinfo=zone) + timedelta(hours=start_h)
    out = None if open_ else start + timedelta(hours=end_h - start_h)
    worked = 0 if open_ else int((out - start).total_seconds() // 60) - break_minutes
    return EntryRow(
        id=uuid.uuid4(), person_id=person, person_name=name, initiative_id=job,
        job_name=job_name if job else None, site_name=site, tz_name=tz,
        local_day=day, clock_in=start, clock_out=out,
        break_minutes=break_minutes, worked_minutes=worked,
        status="open" if open_ else status,
        status_label="On the clock" if open_ else status.capitalize(),
        source=source, approved_by_name=approver, adjusted=adjusted,
        adjust_reason=reason, notes=notes)


def _data(entries=None, *, person_label="Everyone", job_label="All jobs",
          person_id=None, initiative_id=None) -> TimesheetData:
    if entries is None:
        entries = [
            _row(),                                                  # 8h approved
            _row(start_h=17, end_h=19.5, status="pending",
                 adjusted=True, reason="Forgot to clock out", notes="late <b>"),
            _row(person=BOB, name="Bob Jones", day=date(2026, 10, 2), job=None,
                 start_h=6, end_h=17, source="manual", approver="Boss Lady"),
            _row(person=BOB, name="Bob Jones", day=date(2026, 10, 3),
                 open_=True, tz="America/Los_Angeles", site="West"),
        ]
    entries = sorted(entries, key=lambda e: (e.local_day, e.person_name, e.clock_in))
    flags_for(entries, NOW)
    by_person, by_job, approved, pending = rollups(entries)
    days = day_rows(entries)
    filters = TimesheetFilters(
        from_day=date(2026, 10, 1), to_day=date(2026, 10, 31),
        person_id=person_id, initiative_id=initiative_id, site_id=None,
        statuses=("approved", "pending", "open"))
    return TimesheetData(
        filters=filters, person_label=person_label, job_label=job_label,
        site_label="All sites", entries=entries, days=days, by_person=by_person,
        by_job=by_job, approved_minutes=approved, pending_minutes=pending,
        people=len(by_person), day_count=len({d.day for d in days}),
        flagged_entries=sum(1 for e in entries if e.flags),
        default_tz="America/New_York")


def _sheet_rows(ws):
    return [[c.value for c in r] for r in ws.iter_rows()]


def _find(rows, first):
    return next(i for i, r in enumerate(rows) if r and r[0] == first)


# ------------------------------------------------------------ xlsx

def _wb(views=("day", "punch"), data=None):
    raw = ts_xlsx.build_workbook(data or _data(), list(views), NOW)
    return openpyxl.load_workbook(BytesIO(raw))


def test_xlsx_sheet_names_follow_the_views():
    assert _wb().sheetnames == ["Summary", "By day", "Punches"]
    assert _wb(["day"]).sheetnames == ["Summary", "By day"]
    assert _wb(["punch"]).sheetnames == ["Summary", "Punches"]


def test_xlsx_summary_title_filters_and_kpis():
    ws = _wb()["Summary"]
    rows = _sheet_rows(ws)
    assert rows[0][0] == "Timesheet: 2026-10-01 to 2026-10-31"
    assert ws["A1"].font.bold
    head = rows[:_find(rows, "By person")]
    lines = {r[0]: r[1] for r in head if len(r) > 1 and r[0]}
    assert lines["Person"] == "Everyone"
    assert lines["Job"] == "All jobs"
    assert lines["Site"] == "All sites"
    assert lines["Statuses"] == "Approved, Pending, On the clock"
    assert lines["Generated"].endswith("EDT")
    assert "America/New_York" in lines["Time zone"]
    # KPIs: 4 entries, 2 people, 3 distinct dates, 8h approved + 11h approved
    assert lines["Entries"] == 4
    assert lines["People"] == 2
    assert lines["Days"] == 3
    assert lines["Approved hours"] == pytest.approx(19.0)
    assert lines["Pending hours"] == pytest.approx(2.5)
    assert lines["Flagged entries"] == 3   # adjusted, manual+Over 10 h, still clocked in


def test_xlsx_summary_by_person_and_by_job_tables():
    rows = _sheet_rows(_wb()["Summary"])
    i = _find(rows, "By person")
    assert rows[i + 1][:7] == ["Person", "Days", "Entries", "Approved", "Pending",
                               "Total", "Flagged"]
    ann = rows[i + 2]
    assert ann[0] == "Ann Smith" and ann[1] == 1 and ann[2] == 2
    assert ann[3] == pytest.approx(8.0) and ann[4] == pytest.approx(2.5)
    assert ann[5] == pytest.approx(10.5) and ann[6] == 1
    j = _find(rows, "By job")
    assert rows[j + 1][:6] == ["Job", "People", "Entries", "Approved", "Pending", "Total"]
    names = [r[0] for r in rows[j + 2:j + 4]]
    assert names == ["Alpha", "No job"]


def test_xlsx_by_day_headers_and_rows():
    ws = _wb()["By day"]
    rows = _sheet_rows(ws)
    assert rows[0] == ["Date", "Person", "Entries", "First in", "Last out",
                       "Worked", "Hours", "Status", "Flags"]
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"
    first = rows[1]
    assert first[0].date() == date(2026, 10, 1)
    assert first[1] == "Ann Smith" and first[2] == 2
    assert first[3] == "08:00 EDT" and first[4] == "19:30 EDT"
    assert first[5] == "10h 30m"
    assert first[6] == pytest.approx(10.5)
    assert first[7] == "Mixed"
    assert first[8] == "Adjusted"
    bob_open = rows[3]
    assert bob_open[4] in (None, "")          # still on the clock
    assert bob_open[7] == "On the clock" and bob_open[8] == "Still clocked in"
    assert bob_open[3] == "08:00 PDT"          # the entry's own zone


def test_xlsx_hours_are_numeric_with_two_places():
    ws = _wb()["By day"]
    cell = ws["G2"]
    assert cell.value == pytest.approx(10.5) and cell.number_format == "0.00"
    pc = _wb()["Punches"]["I2"]     # 8.00 reads back as an int; still a number
    assert isinstance(pc.value, (int, float)) and pc.value == 8
    assert pc.number_format == "0.00"


def test_xlsx_punches_headers_and_flags_join():
    ws = _wb()["Punches"]
    rows = _sheet_rows(ws)
    assert rows[0] == ["Date", "Person", "Job", "Site", "Clock in", "Clock out",
                       "Break (min)", "Worked", "Hours", "Status", "Source",
                       "Approved by", "Flags", "Adjust reason", "Notes"]
    assert ws.freeze_panes == "A2"
    adj = rows[2]
    assert adj[13] == "Forgot to clock out" and adj[14] == "late <b>"
    manual = rows[3]
    assert manual[12] == "Manual entry, Over 10 h"
    assert manual[11] == "Boss Lady" and manual[10] == "manual"
    assert manual[2] in (None, "")             # no job
    assert rows[1][4] == "08:00 EDT" and rows[1][5] == "16:00 EDT"


def test_xlsx_column_width_is_capped_at_forty():
    long = _row(notes="x" * 200)
    ws = _wb(data=_data([long]))["Punches"]
    assert ws.column_dimensions["O"].width <= 40


def test_xlsx_free_text_is_stored_as_text_never_a_formula():
    evil = _row(name="+cmd", notes="=1+1", reason="@SUM(A1)", job_name="-job",
                site="=HYPERLINK(\"http://x\")", adjusted=True, approver="=boss")
    wb = _wb(data=_data([evil], person_label="=who", job_label="+job"))
    punches = wb["Punches"]
    row = {h: c for h, c in zip(
        [c.value for c in punches[1]], punches[2], strict=True)}
    assert row["Notes"].value == "=1+1" and row["Notes"].data_type == "s"
    assert row["Person"].value == "+cmd" and row["Person"].data_type == "s"
    assert row["Adjust reason"].value == "@SUM(A1)"
    assert row["Job"].value == "-job" and row["Job"].data_type == "s"
    assert row["Site"].data_type == "s" and row["Approved by"].data_type == "s"
    assert wb["By day"]["B2"].value == "+cmd" and wb["By day"]["B2"].data_type == "s"
    summary = wb["Summary"]
    texts = {r[0].value: r[1] for r in summary.iter_rows(min_row=3, max_row=4)}
    assert texts["Person"].value == "=who" and texts["Person"].data_type == "s"
    assert texts["Job"].value == "+job" and texts["Job"].data_type == "s"
    for ws in wb.worksheets:
        assert not any(c.data_type == "f" for r in ws.iter_rows() for c in r)


def test_xlsx_illegal_control_characters_are_stripped():
    bad = _row(name="An\x01n", notes="line\x0bbreak\ttab\nnew", reason="r\x1fs")
    wb = _wb(data=_data([bad], person_label="Ev\x02ery"))
    row = {h: c.value for h, c in zip(
        [c.value for c in wb["Punches"][1]], wb["Punches"][2], strict=True)}
    assert row["Person"] == "Ann"
    assert row["Notes"] == "linebreak\ttab\nnew"
    assert row["Adjust reason"] == "rs"
    assert wb["By day"]["B2"].value == "Ann"
    assert wb["Summary"]["B3"].value == "Every"


def test_xlsx_large_workbook_is_written_in_order_and_position():
    # ws.max_row is O(cells) in openpyxl: asking it per row was quadratic.
    n = 3000
    entries = [_row(day=date(2026, 10, 1) + timedelta(days=i % 28), start_h=6 + (i % 5),
                    end_h=7 + (i % 5), name=f"Person {i:04d}", person=uuid.uuid4())
               for i in range(n)]
    wb = _wb(data=_data(entries))
    punches = wb["Punches"]
    assert punches.max_row == n + 1
    names = [punches.cell(row=r, column=2).value for r in range(2, n + 2)]
    assert sorted(names) == sorted(f"Person {i:04d}" for i in range(n))
    assert punches.cell(row=n + 1, column=1).number_format == "yyyy-mm-dd"
    assert punches.cell(row=n + 1, column=9).number_format == "0.00"
    assert wb["By day"].max_row == n + 1


def test_xlsx_empty_result_still_builds():
    wb = _wb(data=_data([]))
    assert wb.sheetnames == ["Summary", "By day", "Punches"]
    assert _sheet_rows(wb["By day"])[0][0] == "Date"


# ------------------------------------------------------------ pdf

def test_pdf_html_has_sections_for_the_chosen_views_and_tracking_id():
    tid = str(uuid.uuid4())
    html = ts_pdf.render_html(_data(), ["day", "punch"], generated_at=NOW,
                              tracking_id=tid)
    for heading in ("Timesheet", "By person", "By job", "By day", "Punches",
                    "End of Report"):
        assert heading in html
    assert tid in html
    assert "data:image/png;base64," in html           # PDF417
    assert "Page " in html and "counter(pages)" in html
    assert "08:00 EDT" in html and "10h 30m" in html
    assert "late &lt;b&gt;" in html                   # autoescaped


def test_pdf_html_omits_unchosen_views_and_by_job_when_job_filtered():
    day_only = ts_pdf.render_html(_data(), ["day"], generated_at=NOW, tracking_id="t")
    assert "<h2>By day</h2>" in day_only and "<h2>Punches</h2>" not in day_only
    punch_only = ts_pdf.render_html(_data(), ["punch"], generated_at=NOW, tracking_id="t")
    assert "<h2>Punches</h2>" in punch_only and "<h2>By day</h2>" not in punch_only
    filtered = ts_pdf.render_html(
        _data(job_label="Alpha", initiative_id=JOB), ["day"],
        generated_at=NOW, tracking_id="t")
    assert "<h2>By job</h2>" not in filtered
    assert "<h2>By person</h2>" in filtered


async def test_pdf_renders_real_pdf_bytes():
    html = ts_pdf.render_html(_data(), ["day", "punch"], generated_at=NOW,
                              tracking_id=str(uuid.uuid4()))
    content = await ts_pdf.render_pdf(html)
    assert content.startswith(b"%PDF")


# ------------------------------------------------------------ build()

async def _seed(db, *, run_options, definition_options=None, with_job=False):
    from serversherpa.db.models import Initiative
    requester = Person(first_name="Rae", last_name="Requester")
    ann = Person(first_name="Ann", last_name="Smith")
    db.add_all([requester, ann])
    await db.flush()
    job = None
    if with_job:
        job = Initiative(name="Alpha: Phase 1/2", initiative_type="project")
        db.add(job)
        await db.flush()
    definition = ReportDefinition(
        name="Timesheet", report_type="timesheet", is_system=True,
        options=definition_options if definition_options is not None
        else timesheet.default_options())
    db.add(definition)
    await db.flush()
    start = datetime(2026, 10, 2, 12, 0, tzinfo=UTC)
    db.add(TimeEntry(person_id=ann.id, clock_in_at=start, status="approved",
                     clock_out_at=start + timedelta(hours=8),
                     initiative_id=job.id if job else None))
    run = ReportRun(definition_id=definition.id, report_type="timesheet",
                    initiative_id=job.id if job else None,
                    options={k: (str(ann.id) if v == "<ann>" else v)
                             for k, v in run_options.items()},
                    requested_by=requester.id, requested_rank=0)
    db.add(run)
    await db.commit()
    return run


async def test_build_xlsx_default_filename(db):
    run = await _seed(db, run_options=dict(GOOD))
    result = await timesheet.build(db, run)
    assert result.content_type == XLSX_MIME
    import re
    assert re.fullmatch(
        r"Timesheet - 2026-10-01 to 2026-10-31 - \d{4}-\d{2}-\d{2} \d{4}\.xlsx",
        result.filename), result.filename
    wb = openpyxl.load_workbook(BytesIO(result.content))
    assert wb.sheetnames == ["Summary", "By day", "Punches"]
    assert _sheet_rows(wb["Punches"])[1][1] == "Ann Smith"


async def test_build_filename_carries_person_and_job_with_unsafe_stripped(db):
    run = await _seed(db, run_options={**GOOD, "person_id": "<ann>", "views": ["day"]},
                      with_job=True)
    result = await timesheet.build(db, run)
    assert result.filename.startswith(
        "Timesheet - 2026-10-01 to 2026-10-31 - Ann Smith - Alpha- Phase 1-2 - ")
    for ch in '\\/:*?"<>|':
        assert ch not in result.filename
    wb = openpyxl.load_workbook(BytesIO(result.content))
    assert wb.sheetnames == ["Summary", "By day"]


async def test_build_falls_back_to_definition_defaults(db):
    run = await _seed(db, run_options=dict(GOOD), definition_options={
        "default_format": "pdf", "default_views": ["punch"],
        "default_statuses": ["approved"]})
    result = await timesheet.build(db, run)
    assert result.filename.endswith(".pdf")
    assert result.content_type == "application/pdf"
    assert result.content.startswith(b"%PDF")


async def test_build_too_many_entries_propagates(db, monkeypatch):
    from serversherpa.reports.timesheet import gather as g
    from serversherpa.reports.timesheet.gather import TimesheetTooLarge
    run = await _seed(db, run_options=dict(GOOD))
    real = g.gather
    monkeypatch.setattr(g, "gather", lambda d, f, **kw: real(d, f, limit=0, **kw))
    with pytest.raises(TimesheetTooLarge) as exc:
        await timesheet.build(db, run)
    assert exc.value.code == "too_many_entries"


async def test_build_pdf_over_the_pdf_limit_fails_with_its_own_code(db, monkeypatch):
    from serversherpa.reports.timesheet import gather as g
    from serversherpa.reports.timesheet.gather import TimesheetTooLarge
    run = await _seed(db, run_options={**GOOD, "format": "pdf"})
    monkeypatch.setattr(g, "MAX_PDF_ENTRIES", 0)
    with pytest.raises(TimesheetTooLarge) as exc:
        await timesheet.build(db, run)
    assert exc.value.code == "too_many_for_pdf"
