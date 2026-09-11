#!/usr/bin/env python3
"""Rebuild the Transportation Standards docx from a generated survey.

The generated Champagne survey Jimmy supplied has a "Transportation
Standards" sheet appended to it (V2 used to append that sheet at render
time from a docx it kept in object storage; that source docx was lost).
This script reverses the append: it reads the sheet's rows top to bottom,
classifies each as a heading (bold text), a bullet (leading `•`, with
4-space indentation per list level), or plain text, pulls the sheet's
embedded images out in anchor-row order, and writes a minimal `.docx`
that reproduces the same content — so the report definition can carry a
real `report_asset` and `reports/site_move_survey/standards.py`'s parser
(a stdlib port of V2's `_parse_standards_docx`) can read it back.

Built with the stdlib `zipfile`/`xml` modules only — no `python-docx`
dependency, matching the parser it feeds.

Usage:
    python scripts/rebuild_transportation_standards.py <generated.xlsx> <out.docx>
"""

from __future__ import annotations

import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl

SHEET_NAME = "Transportation Standards"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"
PR_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

EMU_PER_PIXEL = 9525  # 96 dpi


@dataclass
class TextItem:
    kind: str  # "heading" | "text" | "bullet"
    text: str
    level: int = 0


@dataclass
class ImageItem:
    row: int  # 1-indexed anchor row, used only for ordering
    data: bytes
    width: int
    height: int


def extract_items(xlsx_path: Path) -> list[TextItem | ImageItem]:
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[SHEET_NAME]

    text_events: list[tuple[int, TextItem]] = []
    for row in range(1, ws.max_row + 1):
        cell = ws.cell(row=row, column=2)  # column B
        raw = cell.value
        if not raw or not str(raw).strip():
            continue
        raw = str(raw)
        stripped = raw.lstrip(" ")
        leading = len(raw) - len(stripped)
        if stripped.startswith("•"):
            content = stripped[1:].strip()
            level = leading // 4
            text_events.append((row, TextItem(kind="bullet", text=content, level=level)))
        elif bool(cell.font and cell.font.bold):
            text_events.append((row, TextItem(kind="heading", text=raw.strip())))
        else:
            text_events.append((row, TextItem(kind="text", text=raw.strip())))

    image_events: list[tuple[int, ImageItem]] = []
    for img in getattr(ws, "_images", []):
        anchor_row = img.anchor._from.row + 1  # anchor rows are 0-indexed
        image_events.append(
            (
                anchor_row,
                ImageItem(row=anchor_row, data=img._data(), width=int(img.width), height=int(img.height)),
            )
        )

    # Merge in row order; a text row and an image never share a row in the
    # Champagne source, but sort text first as a tiebreak just in case.
    combined = [(row, 0, item) for row, item in text_events] + [
        (row, 1, item) for row, item in image_events
    ]
    combined.sort(key=lambda entry: (entry[0], entry[1]))
    return [item for _row, _prio, item in combined]


def _content_types_xml(image_count: int) -> bytes:
    defaults = [
        f'<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
        f'<Default Extension="xml" ContentType="application/xml"/>',
    ]
    if image_count:
        defaults.append('<Default Extension="png" ContentType="image/png"/>')
    overrides = [
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>',
        '<Override PartName="/word/numbering.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>',
    ]
    body = "".join(defaults) + "".join(overrides)
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Types xmlns="{CT_NS}">{body}</Types>'
    )
    return xml.encode("utf-8")


def _package_rels_xml() -> bytes:
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{PR_NS}">'
        f'<Relationship Id="rId1" '
        f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        f'Target="word/document.xml"/>'
        f"</Relationships>"
    )
    return xml.encode("utf-8")


def _document_rels_xml(image_rel_ids: list[str]) -> bytes:
    rels = [
        f'<Relationship Id="rIdNumbering" '
        f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" '
        f'Target="numbering.xml"/>'
    ]
    for idx, rel_id in enumerate(image_rel_ids, start=1):
        rels.append(
            f'<Relationship Id="{rel_id}" '
            f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            f'Target="media/image{idx}.png"/>'
        )
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{PR_NS}">' + "".join(rels) + "</Relationships>"
    )
    return xml.encode("utf-8")


def _numbering_xml() -> bytes:
    levels = "".join(
        f'<w:lvl w:ilvl="{lvl}"><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/></w:lvl>'
        for lvl in range(4)
    )
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:numbering xmlns:w="{W_NS}">'
        f'<w:abstractNum w:abstractNumId="0">{levels}</w:abstractNum>'
        '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
        "</w:numbering>"
    )
    return xml.encode("utf-8")


def _heading_paragraph(text: str) -> str:
    return (
        "<w:p><w:pPr><w:pStyle w:val=\"Heading1\"/></w:pPr>"
        f'<w:r><w:rPr><w:b/></w:rPr><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'
    )


def _bullet_paragraph(text: str, level: int) -> str:
    return (
        '<w:p><w:pPr><w:pStyle w:val="ListParagraph"/>'
        f'<w:numPr><w:ilvl w:val="{level}"/><w:numId w:val="1"/></w:numPr></w:pPr>'
        f'<w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'
    )


def _text_paragraph(text: str) -> str:
    return f'<w:p><w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'


def _image_paragraph(rel_id: str, doc_pr_id: int, width_px: int, height_px: int) -> str:
    cx = max(1, width_px) * EMU_PER_PIXEL
    cy = max(1, height_px) * EMU_PER_PIXEL
    return (
        "<w:p><w:r><w:drawing>"
        f'<wp:inline xmlns:wp="{WP_NS}" distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:docPr id="{doc_pr_id}" name="Picture {doc_pr_id}"/>'
        f'<a:graphic xmlns:a="{A_NS}">'
        f'<a:graphicData uri="{PIC_NS}">'
        f'<pic:pic xmlns:pic="{PIC_NS}">'
        f'<pic:nvPicPr><pic:cNvPr id="{doc_pr_id}" name="Picture {doc_pr_id}"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip xmlns:r="{R_NS}" r:embed="{rel_id}"/>'
        "<a:stretch><a:fillRect/></a:stretch></pic:blipFill>"
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
        "</pic:pic></a:graphicData></a:graphic></wp:inline>"
        "</w:drawing></w:r></w:p>"
    )


def build_docx(items: list[TextItem | ImageItem]) -> bytes:
    paragraphs: list[str] = []
    image_rel_ids: list[str] = []
    image_data: list[bytes] = []
    doc_pr_id = 1

    for item in items:
        if isinstance(item, ImageItem):
            rel_id = f"rIdImg{len(image_rel_ids) + 1}"
            image_rel_ids.append(rel_id)
            image_data.append(item.data)
            paragraphs.append(_image_paragraph(rel_id, doc_pr_id, item.width, item.height))
            doc_pr_id += 1
        elif item.kind == "heading":
            paragraphs.append(_heading_paragraph(item.text))
        elif item.kind == "bullet":
            paragraphs.append(_bullet_paragraph(item.text, item.level))
        else:
            paragraphs.append(_text_paragraph(item.text))

    body = "".join(paragraphs) + "<w:sectPr/>"
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{W_NS}">'
        f"<w:body>{body}</w:body>"
        "</w:document>"
    ).encode("utf-8")

    buf_path_data = {
        "[Content_Types].xml": _content_types_xml(len(image_rel_ids)),
        "_rels/.rels": _package_rels_xml(),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _document_rels_xml(image_rel_ids),
        "word/numbering.xml": _numbering_xml(),
    }
    for idx, data in enumerate(image_data, start=1):
        buf_path_data[f"word/media/image{idx}.png"] = data

    import io

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in buf_path_data.items():
            zf.writestr(name, data)
    return out.getvalue()


def rebuild(src_path: Path, dest_path: Path) -> None:
    items = extract_items(src_path)
    docx_bytes = build_docx(items)
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    dest_path.write_bytes(docx_bytes)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(f"usage: {argv[0]} <generated.xlsx> <out.docx>", file=sys.stderr)
        return 2
    src_path = Path(argv[1])
    dest_path = Path(argv[2])
    rebuild(src_path, dest_path)
    print(f"wrote {dest_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
