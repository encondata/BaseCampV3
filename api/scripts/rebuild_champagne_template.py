#!/usr/bin/env python3
"""Rebuild the Champagne annotated survey template from a generated survey.

Jimmy supplied a *generated* Champagne survey (a filled-in copy of the real
partner questionnaire). This script reverses that fill: it drops the
appended "Transportation Standards" sheet (that content is now sourced from
the report definition's `report_asset` attachment, see
`rebuild_transportation_standards.py`) and writes the `{{...}}` placeholders
from the V2 template-annotation guide (section 10, the Champagne annotation
map) back into their cells, so the result is usable both as a test fixture
and as the real survey_template attachment on the Champagne partner.

Usage:
    python scripts/rebuild_champagne_template.py <generated.xlsx> <out.xlsx>
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

CUSTOMER_SHEET = "Customer and Site Information"
QUESTIONS_SHEET = "General Questions"
EQUIPMENT_SHEET = "Equipment Listing"

# Cells on "Customer and Site Information".
CUSTOMER_PLACEHOLDERS = {
    "C11": "{{customer.company}}",
    "C12": "{{customer.contact_name}}",
    "C13": "{{customer.address}}",
    "C14": "{{customer.phone}}",
    "C15": "{{customer.email}}",
    "C18": "{{origin.name}}",
    "C28": "{{destination.name}}",
    "C19": "{{origin.address}}",
    "C29": "{{destination.address}}",
    "C21": "{{origin.city}}, {{origin.state}}",
    "C31": "{{destination.city}}, {{destination.state}}",
    "C22": "{{origin.zip}}",
    "C32": "{{destination.zip}}",
    "C23": "{{origin.contact_name}}",
    "C33": "{{destination.contact_name}}",
    "C24": "{{origin.contact_phone}}",
    "C34": "{{destination.contact_phone}}",
    "C25": "{{origin.contact_email}}",
    "C35": "{{destination.contact_email}}",
}

# Stray leftover cells on "Customer and Site Information" from the generated
# copy that must not survive into the blank template.
CUSTOMER_STRAY_CELLS = ["E35", "B40", "C40"]

# Cells on "General Questions". D = Origin column, E = Destination column.
QUESTIONS_PLACEHOLDERS = {
    "D3": "{{move.scheduled_start_date}}",
    "D4": "{{move.scheduled_start_time}}",
    "D19": "{{origin.survey.security_clearance_required}}",
    "E19": "{{destination.survey.security_clearance_required}}",
    "C20": "{{origin.survey.security_details}}",
    "D21": "{{origin.survey.dock_available}}",
    "E21": "{{destination.survey.dock_available}}",
    "C22": "Origin: {{origin.survey.dock_hours}}  |  Destination: {{destination.survey.dock_hours}}",
    "D23": "{{origin.survey.dock_75ft_accessible}}",
    "E23": "{{destination.survey.dock_75ft_accessible}}",
    "D24": "{{origin.survey.ground_level_entrance}}",
    "E24": "{{destination.survey.ground_level_entrance}}",
    "C25": "Origin: {{origin.survey.ground_level_details}}  |  Destination: {{destination.survey.ground_level_details}}",
    "D26": "{{origin.survey.floor}}",
    "E26": "{{destination.survey.floor}}",
    "D27": "{{origin.survey.elevator_available}}",
    "E27": "{{destination.survey.elevator_available}}",
    "D28": "{{origin.survey.dock_to_dc_distance_ft}}",
    "E28": "{{destination.survey.dock_to_dc_distance_ft}}",
    "D29": "{{origin.survey.floor_covering_required}}",
    "E29": "{{destination.survey.floor_covering_required}}",
}

# D8/D9 stay static ("Yes"/"No" company-default answers) — not placeholders.
QUESTIONS_STATIC = {"D8": "Yes", "D9": "No"}

# Row 8 on "Equipment Listing" is the repeating asset template row.
EQUIPMENT_ROW = 8
EQUIPMENT_PLACEHOLDERS = {
    "A": "{{asset.index}}",
    "B": "{{asset.rack}}",
    "C": "{{asset.manufacturer}}",
    "D": "{{asset.model}}",
    "E": "{{asset.u_size}}",
    "F": "{{asset.qty}}",
    "G": "{{asset.comments}}",
}


def rebuild(src_path: Path, dest_path: Path) -> None:
    wb = openpyxl.load_workbook(src_path)

    # Drop any sheet whose title mentions transportation standards; that
    # content now lives in the definition's report_asset docx instead of a
    # baked-in sheet.
    for title in list(wb.sheetnames):
        if "transport" in title.lower():
            del wb[title]

    customer_ws = wb[CUSTOMER_SHEET]
    for cell, placeholder in CUSTOMER_PLACEHOLDERS.items():
        customer_ws[cell] = placeholder
    for cell in CUSTOMER_STRAY_CELLS:
        customer_ws[cell] = None

    questions_ws = wb[QUESTIONS_SHEET]
    for cell, placeholder in QUESTIONS_PLACEHOLDERS.items():
        questions_ws[cell] = placeholder
    for cell, value in QUESTIONS_STATIC.items():
        questions_ws[cell] = value

    equipment_ws = wb[EQUIPMENT_SHEET]
    for col, placeholder in EQUIPMENT_PLACEHOLDERS.items():
        equipment_ws[f"{col}{EQUIPMENT_ROW}"] = placeholder

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest_path)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(f"usage: {argv[0]} <generated.xlsx> <out.xlsx>", file=sys.stderr)
        return 2
    src_path = Path(argv[1])
    dest_path = Path(argv[2])
    rebuild(src_path, dest_path)
    print(f"wrote {dest_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
