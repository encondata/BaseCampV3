"""Move-asset import file parsing: CSV/XLSX bytes -> numbered rows.

Headers are V2's upload-ft template verbatim (case-insensitive, including
v2's 'Vendor Involvment' misspelling) so existing facility-tracker
spreadsheets keep working. Unknown columns are NOT errors — they ride
along in the per-row `raw` dict and land in initiative_assets.raw_ft.
"""

import csv
import io
from typing import Any

MAX_BYTES = 20 * 1024 * 1024

CANONICAL = [
    "serial_number", "asset_name", "asset_make", "asset_model", "rfid_tag",
    "priority", "disposition", "owner", "source_rack", "source_ru",
    "destination_rack", "destination_ru",
    "data_1", "data_2", "data_3", "data_4", "data_5", "data_6",
    "mgmt_1", "mgmt_2", "vendor_involvement",
]

HEADER_MAP = {
    "serial number": "serial_number", "asset name": "asset_name",
    "asset make": "asset_make", "asset model": "asset_model",
    "rfid tag": "rfid_tag", "priority": "priority",
    "disposition": "disposition", "owner": "owner",
    "source rack": "source_rack", "source ru": "source_ru",
    "destination rack": "destination_rack",
    "destination ru": "destination_ru",
    "data 1": "data_1", "data 2": "data_2", "data 3": "data_3",
    "data 4": "data_4", "data 5": "data_5", "data 6": "data_6",
    "mgmt 1": "mgmt_1", "mgmt 2": "mgmt_2",
    "vendor involvement": "vendor_involvement",
    "vendor involvment": "vendor_involvement",   # v2 template's spelling
}

TEMPLATE_HEADERS = [
    "Serial Number", "Asset Name", "Asset Make", "Asset Model", "RFID Tag",
    "Priority", "Disposition", "Owner", "Source Rack", "Source RU",
    "Destination Rack", "Destination RU", "Data 1", "Data 2", "Data 3",
    "Data 4", "Data 5", "Data 6", "Mgmt 1", "Mgmt 2", "Vendor Involvement",
]

SAMPLE_ROWS: list[dict] = [
    {"Serial Number": "SN-0001", "Asset Name": "web-01", "Asset Make": "Dell",
     "Asset Model": "PowerEdge R740", "RFID Tag": "", "Priority": "Wave 1",
     "Disposition": "Relocate", "Owner": "Platform",
     "Source Rack": "11.01.01.01A.02", "Source RU": "12",
     "Destination Rack": "BJ08", "Destination RU": "24",
     "Data 1": "sw1:eth1/1", "Data 2": "", "Data 3": "", "Data 4": "",
     "Data 5": "", "Data 6": "", "Mgmt 1": "mgmt-sw:1", "Mgmt 2": "",
     "Vendor Involvement": ""},
    {"Serial Number": "SN-0002", "Asset Name": "san-01", "Asset Make": "HPE",
     "Asset Model": "Nimble HF20", "RFID Tag": "", "Priority": "Wave 2",
     "Disposition": "", "Owner": "", "Source Rack": "BJ01",
     "Source RU": "3.5", "Destination Rack": "", "Destination RU": "",
     "Data 1": "", "Data 2": "", "Data 3": "", "Data 4": "", "Data 5": "",
     "Data 6": "", "Mgmt 1": "", "Mgmt 2": "", "Vendor Involvement": "yes"},
]


class ImportFileError(Exception):
    """Whole-file failure (not a per-row error)."""

    def __init__(self, code: str, **extra: Any) -> None:
        super().__init__(code)
        self.code = code
        self.extra = extra


def _cell(value: Any) -> str:
    """Spreadsheet cells arrive as str/float/int/bool/None — normalize to
    trimmed text. Integral floats (openpyxl's 12.0) drop the .0 so
    numeric-looking text columns round-trip."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _check_serial_header(headers: list[str]) -> None:
    if not any(HEADER_MAP.get(h.lower()) == "serial_number" for h in headers):
        raise ImportFileError("missing_serial_column")


def _numbered(raw_rows: list[dict], first_row: int,
              ) -> list[tuple[int, dict, dict]]:
    out: list[tuple[int, dict, dict]] = []
    for i, raw in enumerate(raw_rows):
        text_row = {(k or "").strip(): _cell(v) for k, v in raw.items()
                    if (k or "").strip()}
        canonical = {c: "" for c in CANONICAL}
        for header, value in text_row.items():
            field = HEADER_MAP.get(header.lower())
            if field:
                canonical[field] = value
        if any(v != "" for v in text_row.values()):   # skip fully blank lines
            out.append((first_row + i, canonical, text_row))
    return out


def parse_upload(filename: str, content: bytes,
                 ) -> list[tuple[int, dict, dict]]:
    if len(content) > MAX_BYTES:
        raise ImportFileError("file_too_large", limit=MAX_BYTES)
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        except UnicodeDecodeError:
            raise ImportFileError("invalid_csv") from None
        if reader.fieldnames is None:
            raise ImportFileError("invalid_csv")
        _check_serial_header([f.strip() for f in reader.fieldnames if f])
        rows = [{(k or "").strip(): v for k, v in r.items() if k}
                for r in reader]
        return _numbered(rows, first_row=2)
    if name.endswith((".xlsx", ".xls")):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content),
                                        read_only=True, data_only=True)
        except Exception:
            raise ImportFileError("invalid_xlsx") from None
        ws = wb.worksheets[0]
        lines = ws.iter_rows(values_only=True)
        header = [_cell(h) for h in (next(lines, None) or tuple())]
        header = [h for h in header if h]
        if not header:
            raise ImportFileError("invalid_xlsx")
        _check_serial_header(header)
        rows = [dict(zip(header, line)) for line in lines]
        return _numbered(rows, first_row=2)
    raise ImportFileError("unsupported_file")


def build_template_csv() -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=TEMPLATE_HEADERS,
                            lineterminator="\n")
    writer.writeheader()
    writer.writerows(SAMPLE_ROWS)
    return buf.getvalue()


def build_template_xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Move Assets"
    ws.append(TEMPLATE_HEADERS)
    for row in SAMPLE_ROWS:
        ws.append([row[c] for c in TEMPLATE_HEADERS])
    ref = wb.create_sheet("Reference")
    ref.append(["Required column"])
    ref.append(["Serial Number (or enable serial generation with a "
                "non-blank Asset Name)"])
    ref.append([])
    ref.append(["Make/model modes"])
    ref.append(["fuzzy — match catalog + aliases; unmatched rows need review"])
    ref.append(["force — always create missing make/models"])
    ref.append(["hybrid — match first, create when unmatched"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
