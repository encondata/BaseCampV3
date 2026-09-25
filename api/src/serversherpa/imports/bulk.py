"""Shared bulk-import core: parse (csv/xlsx/json) → numbered rows, plus the
csv/xlsx builders used for templates and exports. Content-agnostic — each
importer (sites, containers, workers) passes its own column list and sheet
name, and keeps its own matching/diff/commit logic.

Row numbering: a file's header is line 1 so data starts at 2; a JSON
payload has no header so rows are numbered from 1.
"""

import csv
import io
import json
import re
from itertools import islice
from typing import Any

MAX_ROWS = 1000
MAX_BYTES = 5 * 1024 * 1024

FORMULA_LEAD = ("=", "+", "-", "@")
_NUMERIC = re.compile(r"^[+-]?\d+(\.\d+)?$")


class BulkImportError(Exception):
    """Whole-payload failure (not a per-row error)."""

    def __init__(self, code: str, **extra: Any) -> None:
        super().__init__(code)
        self.code = code
        self.extra = extra


def cell(value: Any) -> str:
    """Spreadsheet cells arrive as str/float/int/bool/None — normalize to
    trimmed text. Integral floats (openpyxl's 89501.0) drop the .0 so
    numeric-looking text columns round-trip. One leading apostrophe in front
    of a formula lead character is dropped, so our own guarded CSV export
    (see `guard_cell`) re-uploads as the value it started as."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text[:1] == "'" and text[1:2] in FORMULA_LEAD:
        text = text[1:]
    return text


def guard_cell(text: str) -> str:
    """OWASP CSV-injection guard, mirroring the portal's `csvCell`: a cell
    that opens with = + - @ gets a leading apostrophe so Excel/Sheets read it
    as text — unless the whole cell is a number, so -119.8138 stays a
    coordinate."""
    if text[:1] in FORMULA_LEAD and not _NUMERIC.match(text):
        return f"'{text}"
    return text


def check_columns(keys: list[str], columns: list[str]) -> None:
    unknown = sorted({k for k in keys if k not in columns})
    if unknown:
        raise BulkImportError("unknown_columns", columns=unknown)


def numbered(rows: list[dict], first_row: int, columns: list[str],
             max_rows: int = MAX_ROWS) -> list[tuple[int, dict]]:
    if len(rows) > max_rows:
        raise BulkImportError("too_many_rows", limit=max_rows)
    out: list[tuple[int, dict]] = []
    for i, raw in enumerate(rows):
        check_columns(list(raw.keys()), columns)
        row = {col: cell(raw.get(col)) for col in columns}
        if any(v != "" for v in row.values()):        # skip fully blank lines
            out.append((first_row + i, row))
    return out


def number_json_rows(rows: Any, columns: list[str],
                     max_rows: int = MAX_ROWS) -> list[tuple[int, dict]]:
    if isinstance(rows, dict):
        rows = [rows]          # a single bare object is a one-row import
    if not isinstance(rows, list) or not all(isinstance(r, dict) for r in rows):
        raise BulkImportError("invalid_json")
    return numbered(rows, 1, columns, max_rows=max_rows)


def parse_upload(filename: str, content: bytes, columns: list[str], sheet: str,
                 max_rows: int = MAX_ROWS, max_bytes: int = MAX_BYTES,
                 ) -> list[tuple[int, dict]]:
    if len(content) > max_bytes:
        raise BulkImportError("file_too_large", limit=max_bytes)
    name = filename.lower()
    if name.endswith(".json"):
        try:
            return number_json_rows(json.loads(content.decode("utf-8-sig")),
                                    columns, max_rows=max_rows)
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise BulkImportError("invalid_json") from None
    if name.endswith(".csv"):
        try:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        except UnicodeDecodeError:
            raise BulkImportError("invalid_csv") from None
        if reader.fieldnames is None:
            raise BulkImportError("invalid_csv")
        check_columns([f.strip() for f in reader.fieldnames if f], columns)
        # read one row past the limit and stop: `numbered` still raises
        # too_many_rows exactly as before, without building every row first
        rows = [{(k or "").strip(): v for k, v in r.items() if k}
                for r in islice(reader, max_rows + 1)]
        return numbered(rows, 2, columns, max_rows=max_rows)
    if name.endswith(".xlsx"):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content),
                                        read_only=True, data_only=True)
        except Exception:
            raise BulkImportError("invalid_xlsx") from None
        try:
            ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
            lines = ws.iter_rows(values_only=True)
            header = [cell(h) for h in (next(lines, None) or tuple())]
            if not any(header):
                raise BulkImportError("invalid_xlsx")
            check_columns([h for h in header if h], columns)
            # one row past the limit is enough for `numbered` to raise
            # too_many_rows; never build the rest of an oversized sheet
            rows = [{h: v for h, v in zip(header, line) if h}
                    for line in islice(lines, max_rows + 1)]
        finally:
            wb.close()
        return numbered(rows, 2, columns, max_rows=max_rows)
    raise BulkImportError("unsupported_file")


def build_rows_csv(rows: list[dict], columns: list[str]) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=columns, lineterminator="\n")
    writer.writeheader()
    writer.writerows({c: guard_cell(str(row[c])) for c in columns}
                     for row in rows)
    return buf.getvalue()


def build_rows_xlsx(rows: list[dict], columns: list[str], sheet: str,
                    reference: list[tuple[str, list[str]]]) -> bytes:
    """`sheet` + a Reference sheet: each (title, keys) block is a title row,
    one key per row, and a blank row between blocks."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(columns)
    for row in rows:
        ws.append([row[c] for c in columns])
        # a cell openpyxl would otherwise store as a formula is pinned to
        # text, so a name like "=HYPERLINK(…)" can never execute on open
        for c in ws[ws.max_row]:
            if isinstance(c.value, str) and c.value[:1] in FORMULA_LEAD:
                c.data_type = "s"
    ref = wb.create_sheet("Reference")
    for i, (title, keys) in enumerate(reference):
        if i:
            ref.append([])
        ref.append([title])
        for key in keys:
            ref.append([key])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def renumber(numbered: list[tuple[int, dict]], row_numbers: Any) -> list[tuple[int, dict]]:
    """JSON rows re-posted after a file preview carry the spreadsheet line
    numbers the preview assigned, so overrides / skips keyed by those numbers
    still line up. None keeps the JSON numbering (from 1)."""
    if row_numbers is None:
        return numbered
    if (not isinstance(row_numbers, list) or len(row_numbers) != len(numbered)
            or not all(isinstance(n, int) and not isinstance(n, bool) for n in row_numbers)
            or len(set(row_numbers)) != len(row_numbers)):
        raise BulkImportError("invalid_row_numbers")
    return [(n, row) for n, (_, row) in zip(row_numbers, numbered, strict=True)]


def parse_overrides(raw: Any, fields: tuple[str, ...]) -> dict[int, dict[str, str]]:
    """`{"<row>": {"<field>": "<picked id>"}}` → {row: {field: id}}; only the
    importer's own `fields`, each a non-empty string."""
    if raw is None:
        return {}
    if not isinstance(raw, dict):
        raise BulkImportError("invalid_overrides")
    out: dict[int, dict[str, str]] = {}
    for key, picks in raw.items():
        try:
            row = int(key)
        except (TypeError, ValueError):
            raise BulkImportError("invalid_overrides") from None
        if (not isinstance(picks, dict)
                or not all(f in fields and isinstance(v, str) and v for f, v in picks.items())):
            raise BulkImportError("invalid_overrides")
        out[row] = dict(picks)
    return out


def parse_row_list(raw: Any, code: str) -> set[int]:
    """A list of row numbers (skips, approvals); `code` names the error."""
    if raw is None:
        return set()
    if not isinstance(raw, list) or not all(
            isinstance(n, int) and not isinstance(n, bool) for n in raw):
        raise BulkImportError(code)
    return set(raw)
