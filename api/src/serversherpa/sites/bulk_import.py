"""Sites bulk import: parse (csv/xlsx/json) → validate → preview/commit.
Pure row pipeline; routes stay thin. All-or-nothing semantics live here.

Blank-cell rule: on create rows, blank status/country take the defaults
(active/US); on update rows a blank cell means "no change", never a clear —
the original cell blankness is tracked out-of-band (`_blank`) because the
normalized `data` dict has already had defaults applied by diff time.
"""

import csv
import io
import json
import uuid
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from serversherpa.db.models import (
    Client, Partner, Site, SiteClient, SiteType, StatusValue,
)

COLUMNS = [
    "name", "code", "type", "status", "address_line1", "address_line2",
    "city", "region", "postal_code", "country", "latitude", "longitude",
    "timezone", "dc_provider", "partner", "clients", "notes",
]
# template column → Site attribute (identity except type; partner/clients are
# relations, handled separately)
SITE_ATTR = {c: ("site_type" if c == "type" else c) for c in COLUMNS
             if c not in ("partner", "clients")}
MAX_ROWS = 1000
MAX_BYTES = 5 * 1024 * 1024

SAMPLE_ROWS: list[dict] = [
    {"name": "Example DC West", "code": "DCW", "type": "datacenter",
     "status": "active", "address_line1": "100 Server Way", "address_line2": "",
     "city": "Reno", "region": "NV", "postal_code": "89501", "country": "US",
     "latitude": "39.5296", "longitude": "-119.8138",
     "timezone": "America/Los_Angeles", "dc_provider": "Switch",
     "partner": "", "clients": "Acme Co; Globex", "notes": "Sample row — replace me"},
    {"name": "Example Office", "code": "", "type": "office", "status": "planned",
     "address_line1": "", "address_line2": "", "city": "Zurich", "region": "",
     "postal_code": "", "country": "CH", "latitude": "", "longitude": "",
     "timezone": "Europe/Zurich", "dc_provider": "", "partner": "",
     "clients": "", "notes": ""},
]


class BulkImportError(Exception):
    """Whole-payload failure (not a per-row error)."""

    def __init__(self, code: str, **extra: Any) -> None:
        super().__init__(code)
        self.code = code
        self.extra = extra


# ── parsing ─────────────────────────────────────────────────────────

def _cell(value: Any) -> str:
    """Spreadsheet cells arrive as str/float/int/bool/None — normalize to
    trimmed text. Integral floats (openpyxl's 89501.0) drop the .0 so
    numeric-looking text columns round-trip."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _check_columns(keys: list[str]) -> None:
    unknown = sorted({k for k in keys if k not in COLUMNS})
    if unknown:
        raise BulkImportError("unknown_columns", columns=unknown)


def _numbered(rows: list[dict], first_row: int) -> list[tuple[int, dict]]:
    if len(rows) > MAX_ROWS:
        raise BulkImportError("too_many_rows", limit=MAX_ROWS)
    out: list[tuple[int, dict]] = []
    for i, raw in enumerate(rows):
        _check_columns(list(raw.keys()))
        row = {col: _cell(raw.get(col)) for col in COLUMNS}
        if any(v != "" for v in row.values()):        # skip fully blank lines
            out.append((first_row + i, row))
    return out


def number_json_rows(rows: Any) -> list[tuple[int, dict]]:
    if isinstance(rows, dict):
        rows = [rows]          # a single bare object is a one-row import
    if not isinstance(rows, list) or not all(isinstance(r, dict) for r in rows):
        raise BulkImportError("invalid_json")
    return _numbered(rows, first_row=1)


def parse_upload(filename: str, content: bytes) -> list[tuple[int, dict]]:
    if len(content) > MAX_BYTES:
        raise BulkImportError("file_too_large", limit=MAX_BYTES)
    name = filename.lower()
    if name.endswith(".json"):
        try:
            return number_json_rows(json.loads(content.decode("utf-8-sig")))
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise BulkImportError("invalid_json") from None
    if name.endswith(".csv"):
        try:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        except UnicodeDecodeError:
            raise BulkImportError("invalid_csv") from None
        if reader.fieldnames is None:
            raise BulkImportError("invalid_csv")
        _check_columns([f.strip() for f in reader.fieldnames if f])
        rows = [{(k or "").strip(): v for k, v in r.items() if k}
                for r in reader]
        return _numbered(rows, first_row=2)
    if name.endswith(".xlsx"):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content),
                                        read_only=True, data_only=True)
        except Exception:
            raise BulkImportError("invalid_xlsx") from None
        ws = wb["Sites"] if "Sites" in wb.sheetnames else wb.worksheets[0]
        lines = ws.iter_rows(values_only=True)
        header = [_cell(h) for h in (next(lines, None) or tuple())]
        if not any(header):
            raise BulkImportError("invalid_xlsx")
        _check_columns([h for h in header if h])
        rows = [{h: v for h, v in zip(header, line) if h} for line in lines]
        return _numbered(rows, first_row=2)
    raise BulkImportError("unsupported_file")


# ── templates ───────────────────────────────────────────────────────

def build_template_csv() -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=COLUMNS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(SAMPLE_ROWS)
    return buf.getvalue()


def build_template_xlsx(type_keys: list[str], status_keys: list[str]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sites"
    ws.append(COLUMNS)
    for row in SAMPLE_ROWS:
        ws.append([row[c] for c in COLUMNS])
    ref = wb.create_sheet("Reference")
    ref.append(["Valid type keys"])
    for key in type_keys:
        ref.append([key])
    ref.append([])
    ref.append(["Valid status keys"])
    for key in status_keys:
        ref.append([key])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── validation + preview ────────────────────────────────────────────

async def _reference_data(db: AsyncSession) -> dict:
    type_keys = {t.key for t in await db.scalars(select(SiteType))}
    status_keys = set(await db.scalars(select(StatusValue.key).where(
        StatusValue.record_type == "site")))
    partners: dict[str, list[Partner]] = {}
    for p in await db.scalars(select(Partner)):
        partners.setdefault(p.name.lower(), []).append(p)
    clients: dict[str, list[Client]] = {}
    for c in await db.scalars(select(Client)):
        clients.setdefault(c.name.lower(), []).append(c)
    return {"types": type_keys, "statuses": status_keys,
            "partners": partners, "clients": clients}


def _split_clients(cell: str) -> list[str]:
    return [part.strip() for part in cell.split(";") if part.strip()]


def _coord(value: str, lo: float, hi: float,
           errors: list[str], label: str) -> float | None:
    if value == "":
        return None
    try:
        num = float(value)
    except ValueError:
        errors.append(f"{label} is not a number")
        return None
    if not (lo <= num <= hi):
        errors.append(f"{label} out of range")
        return None
    return num


async def preview_rows(db: AsyncSession, numbered: list[tuple[int, dict]],
                       *, allow_updates: bool) -> dict:
    ref = await _reference_data(db)
    names_seen: dict[str, list[int]] = {}
    for n, row in numbered:
        if row["name"]:
            names_seen.setdefault(row["name"].lower(), []).append(n)

    existing: dict[str, list[Site]] = {}
    wanted = [row["name"] for _, row in numbered if row["name"]]
    if wanted:
        for site in await db.scalars(select(Site).where(Site.name.in_(wanted))):
            existing.setdefault(site.name.lower(), []).append(site)

    dup_sites = [s for sites in existing.values() for s in sites]
    current_clients: dict[uuid.UUID, dict[uuid.UUID, str]] = {}
    partner_names: dict[uuid.UUID, str] = {}
    if dup_sites:
        links = (await db.execute(
            select(SiteClient.site_id, Client.id, Client.name)
            .join(Client, Client.id == SiteClient.client_id)
            .where(SiteClient.site_id.in_([s.id for s in dup_sites])))).all()
        for site_id, client_id, cname in links:
            current_clients.setdefault(site_id, {})[client_id] = cname
        pids = {s.partner_id for s in dup_sites if s.partner_id}
        if pids:
            partner_names = dict((await db.execute(
                select(Partner.id, Partner.name).where(Partner.id.in_(pids))
            )).all())

    results = []
    for n, row in numbered:
        errors: list[str] = []
        name = row["name"]
        if not name:
            errors.append("name is required")
        elif len(names_seen[name.lower()]) > 1:
            errors.append(f"duplicate name '{name}' within the import")

        if row["type"] and row["type"] not in ref["types"]:
            errors.append(f"unknown type '{row['type']}'")
        if row["status"] and row["status"] not in ref["statuses"]:
            errors.append(f"unknown status '{row['status']}'")

        lat = _coord(row["latitude"], -90, 90, errors, "latitude")
        lon = _coord(row["longitude"], -180, 180, errors, "longitude")
        if (row["latitude"] == "") != (row["longitude"] == ""):
            errors.append("latitude and longitude must both be set")

        partner_obj: Partner | None = None
        if row["partner"]:
            matches = ref["partners"].get(row["partner"].lower(), [])
            if len(matches) == 0:
                errors.append(f"unknown partner '{row['partner']}'")
            elif len(matches) > 1:
                errors.append(f"ambiguous partner '{row['partner']}'")
            else:
                partner_obj = matches[0]

        client_objs: list[Client] = []
        for cname in _split_clients(row["clients"]):
            matches = ref["clients"].get(cname.lower(), [])
            if len(matches) == 0:
                errors.append(f"unknown client '{cname}'")
            elif len(matches) > 1:
                errors.append(f"ambiguous client '{cname}'")
            else:
                client_objs.append(matches[0])

        blank = {"status": row["status"] == "", "country": row["country"] == ""}
        data = dict(row)
        data["latitude"], data["longitude"] = lat, lon
        data["clients"] = _split_clients(row["clients"])
        if blank["status"]:
            data["status"] = "active"
        if blank["country"]:
            data["country"] = "US"

        dupes = existing.get(name.lower(), []) if name else []
        action, diff_out, site_id = "create", None, None
        if errors:
            action = "error"
        elif dupes:
            if len(dupes) > 1:
                action = "error"
                errors.append(f"multiple existing sites named '{name}'")
            elif not allow_updates:
                action = "error"
                errors.append(f"site '{name}' already exists")
            else:
                site = dupes[0]
                site_id = str(site.id)
                changes = _diff_row(
                    site, data, blank, partner_obj, client_objs,
                    current_clients.get(site.id, {}), partner_names)
                action = "update" if changes else "unchanged"
                diff_out = changes or None

        results.append({"row": n, "name": name or None, "action": action,
                        "errors": errors, "diff": diff_out, "site_id": site_id,
                        "data": data if action != "error" else None})

    can_commit = bool(results) and all(r["action"] != "error" for r in results)
    return {"rows": results, "can_commit": can_commit,
            "update_allowed": allow_updates}


def _diff_row(site: Site, data: dict, blank: dict,
              partner_obj: Partner | None, client_objs: list[Client],
              linked: dict, partner_names: dict) -> dict:
    """Changed fields only; blank in the row = no change (the create-only
    status/country defaults in `data` must not read as edits — `blank`
    remembers the original cells)."""
    out: dict = {}
    for col, attr in SITE_ATTR.items():
        raw = data[col]
        if col in ("status", "country") and blank[col]:
            continue
        if col in ("latitude", "longitude"):
            if raw is None:
                continue
            old = getattr(site, attr)
            old = float(old) if old is not None else None
            if old != raw:
                out[col] = {"old": old, "new": raw}
            continue
        if raw == "":
            continue
        old = getattr(site, attr)
        if (old or "") != raw:
            out[col] = {"old": old, "new": raw}
    if data["partner"] and partner_obj is not None:
        if site.partner_id != partner_obj.id:
            out["partner"] = {"old": partner_names.get(site.partner_id),
                              "new": partner_obj.name}
    if data["clients"]:
        want = {c.id: c.name for c in client_objs}
        add = sorted(n for i, n in want.items() if i not in linked)
        remove = sorted(n for i, n in linked.items() if i not in want)
        if add or remove:
            out["clients"] = {"add": add, "remove": remove}
    return out


# ── commit ──────────────────────────────────────────────────────────

async def commit_rows(db: AsyncSession, actor_person_id: uuid.UUID,
                      numbered: list[tuple[int, dict]], *, allow_updates: bool,
                      approved_updates: set[str], source_label: str) -> dict:
    """All-or-nothing: re-validates everything, then commits creates plus
    APPROVED updates in one transaction. Raises rows_invalid (carrying the
    full preview payload) if anything blocks — nothing is written."""
    from serversherpa.services.audit import audit

    preview = await preview_rows(db, numbered, allow_updates=allow_updates)
    blocked = [r for r in preview["rows"] if r["action"] == "error"]
    unapproved = [r for r in preview["rows"]
                  if r["action"] == "update" and r["site_id"] not in approved_updates]
    if blocked or unapproved or not preview["rows"]:
        for r in unapproved:
            r["errors"] = [*r["errors"], "update not approved"]
            r["action"] = "error"
        raise BulkImportError("rows_invalid", rows=preview["rows"])

    ref = await _reference_data(db)
    created = updated = unchanged = 0
    for r in preview["rows"]:
        data = r["data"]
        if r["action"] == "unchanged":
            unchanged += 1
        elif r["action"] == "create":
            await _create_site(db, actor_person_id, data, ref)
            created += 1
        else:
            await _apply_update(db, actor_person_id, r, ref)
            updated += 1
    audit(db, actor_id=actor_person_id, entity_type="site_bulk_import",
          entity_id=None, action="bulk_import",
          changes={"created": created, "updated": updated,
                   "unchanged": unchanged, "source": source_label})
    await db.commit()
    return {"created": created, "updated": updated, "unchanged": unchanged}


def _resolve_partner(ref: dict, name: str) -> Partner | None:
    matches = ref["partners"].get(name.lower(), []) if name else []
    return matches[0] if len(matches) == 1 else None


async def _create_site(db: AsyncSession, actor_person_id: uuid.UUID,
                       data: dict, ref: dict) -> None:
    from serversherpa.services.audit import audit

    fields = {attr: data[col] for col, attr in SITE_ATTR.items()
              if data[col] not in ("", None)}
    partner = _resolve_partner(ref, data["partner"])
    if partner is not None:
        fields["partner_id"] = partner.id
    site = Site(**fields, created_by=actor_person_id)
    db.add(site)
    await db.flush()
    for cname in data["clients"]:
        client = ref["clients"][cname.lower()][0]
        db.add(SiteClient(site_id=site.id, client_id=client.id,
                          linked_by=actor_person_id))
    changes = {key: {"from": None, "to": _audit_value(value)}
               for key, value in fields.items()}
    if data["clients"]:
        changes["clients"] = {"from": [], "to": sorted(data["clients"])}
    audit(db, actor_id=actor_person_id, entity_type="site",
          entity_id=str(site.id), action="create", changes=changes)


def _audit_value(value: Any) -> Any:
    return str(value) if isinstance(value, uuid.UUID) else value


async def _apply_update(db: AsyncSession, actor_person_id: uuid.UUID,
                        r: dict, ref: dict) -> None:
    from datetime import UTC, datetime

    from serversherpa.services.audit import audit

    site = await db.get(Site, uuid.UUID(r["site_id"]))
    changes: dict = {}
    for col, change in (r["diff"] or {}).items():
        if col == "clients":
            continue
        if col == "partner":
            partner = _resolve_partner(ref, change["new"])
            site.partner_id = partner.id if partner else site.partner_id
            changes["partner"] = {"from": change["old"], "to": change["new"]}
            continue
        setattr(site, SITE_ATTR[col], change["new"])
        changes[col] = {"from": change["old"], "to": change["new"]}
    client_change = (r["diff"] or {}).get("clients")
    if client_change:
        current = set(await db.scalars(select(SiteClient.client_id).where(
            SiteClient.site_id == site.id)))
        want_ids = {ref["clients"][n.lower()][0].id
                    for n in r["data"]["clients"]}
        for client_id in current - want_ids:
            await db.execute(SiteClient.__table__.delete().where(
                SiteClient.site_id == site.id,
                SiteClient.client_id == client_id))
        for client_id in want_ids - current:
            db.add(SiteClient(site_id=site.id, client_id=client_id,
                              linked_by=actor_person_id))
        changes["clients"] = client_change
    site.updated_at = datetime.now(UTC)
    audit(db, actor_id=actor_person_id, entity_type="site",
          entity_id=str(site.id), action="update", changes=changes)
