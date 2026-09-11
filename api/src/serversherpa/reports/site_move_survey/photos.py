"""Site Photos sheet for the Site & Move Survey xlsx.

Port of V2's `_append_site_photos` (api/reports/site_move_survey.py,
read-only reference at /Users/jrh1812/Developer/BaseCampV2-reference) —
minus the database read, which is `gather.py`'s job in V3 (attachments
`kind='photo'` on `entity_type='site'`, newest first, capped at 10 per
site). This module only ever sees already-fetched image bytes.
"""

import io


def append_site_photos(wb, entries: list[tuple[str, list[bytes]]]) -> None:
    """Append a "Site Photos" sheet: one bold site-name heading per
    entry, followed by a thumbnailed (max 640x640, RGB PNG) image for
    each of that site's photo bytes. Entries with no images are skipped
    (nothing to show); when every entry is empty, no sheet is created at
    all — mirrors V2, which never appended an empty "Site Photos" sheet.
    A single corrupt image is skipped rather than aborting the rest.
    """
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font
    from PIL import Image as PILImage

    entries = [(label, images) for label, images in entries if images]
    if not entries:
        return

    ws = wb.create_sheet("Site Photos")
    row = 1
    for label, images in entries:
        ws.cell(row=row, column=1).value = label
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 2
        for raw in images:
            try:
                pil_img = PILImage.open(io.BytesIO(raw))
                pil_img.thumbnail((640, 640))
                buf = io.BytesIO()
                pil_img.convert("RGB").save(buf, format="PNG")
                buf.seek(0)
                xl_img = XLImage(buf)
                xl_img.anchor = f"A{row}"
                ws.add_image(xl_img)
                row += max(2, int(pil_img.height / 19) + 2)
            except Exception:
                continue
        row += 2
