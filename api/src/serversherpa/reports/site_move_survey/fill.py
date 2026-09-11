"""openpyxl fill engine for the Site & Move Survey template.

Ported verbatim (behavior, not just shape) from V2's
api/reports/site_move_survey.py (`_resolve_path` / `_substitute` /
`_fill_workbook` / `_expand_asset_rows`) — see
/Users/jrh1812/Developer/BaseCampV2-reference/api/reports/site_move_survey.py
(read-only reference) and
docs/site-move-survey/template-annotation-guide.md for the placeholder
rules this implements. Pure functions over an openpyxl workbook; no DB,
no I/O.
"""

import copy as copy_module
import re

from openpyxl.cell.cell import MergedCell

# {{dotted.path}} anywhere in a cell; whitespace right inside the braces
# is ignored (`{{ origin.city }}` == `{{origin.city}}`).
PLACEHOLDER_RE = re.compile(r"\{\{\s*([^}]+?)\s*\}\}")
# A cell whose ENTIRE value is a single placeholder. These resolve to the
# raw context value (int/bool/None/etc.) instead of always producing a
# string, so e.g. `{{move.asset_count}}` keeps rendering as a number.
WHOLE_CELL_RE = re.compile(r"^\s*\{\{\s*([^}]+?)\s*\}\}\s*$")


def resolve_path(context: dict, path: str):
    """Walk a dotted path through nested dicts. Any missing key, or a
    node along the way that isn't a dict, resolves to '' rather than
    raising — a template referencing data that doesn't exist just
    renders blank."""
    node = context
    for part in path.split("."):
        part = part.strip()
        if isinstance(node, dict) and part in node:
            node = node[part]
        else:
            return ""
    return "" if node is None else node


def substitute(value, context: dict):
    """Fill placeholders in one cell value per the annotation guide.

    A whole-cell placeholder (e.g. a cell containing only
    `{{move.asset_count}}`) returns the raw resolved value so numeric/
    boolean types survive into the xlsx cell. A cell that mixes literal
    text with one or more placeholders always renders as a string (with
    missing values substituted as '', so `"{{a}}, {{b}}"` with both empty
    renders `", "` rather than disappearing).
    """
    if not isinstance(value, str) or "{{" not in value:
        return value
    whole = WHOLE_CELL_RE.match(value)
    if whole:
        return resolve_path(context, whole.group(1))
    return PLACEHOLDER_RE.sub(lambda m: str(resolve_path(context, m.group(1))), value)


def fill_workbook(wb, context: dict, assets: list[dict], asset_notes: str | None, *,
                  include_transportation_standards: bool) -> None:
    """Fill every sheet's placeholders in place.

    When the toggle is off, any sheet whose title contains "transport"
    (case-insensitive) is dropped first — covers a template that ships
    its own Transportation Standards sheet, since the toggle always wins.
    A workbook is never emptied entirely by this (a single-sheet workbook
    keeps its last sheet even if it happens to match).

    Each worksheet gets at most one "template row" — the first row
    containing any `{{asset.*}}` placeholder — which is expanded via
    `expand_asset_rows` instead of substituted like every other cell.
    """
    if not include_transportation_standards:
        for title in list(wb.sheetnames):
            if "transport" in title.lower() and len(wb.sheetnames) > 1:
                del wb[title]

    for ws in wb.worksheets:
        template_row_idx = None
        for row in ws.iter_rows():
            has_asset_ph = False
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                v = cell.value
                if isinstance(v, str) and re.search(r"\{\{\s*asset\.", v):
                    has_asset_ph = True
                    break
            if has_asset_ph:
                template_row_idx = row[0].row
                break

        for row in ws.iter_rows():
            if template_row_idx is not None and row and row[0].row == template_row_idx:
                continue
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if isinstance(cell.value, str) and "{{" in cell.value:
                    cell.value = substitute(cell.value, context)

        if template_row_idx is not None:
            expand_asset_rows(ws, template_row_idx, context, assets, asset_notes)


def expand_asset_rows(ws, row_idx: int, context: dict, assets: list[dict],
                      asset_notes: str | None) -> None:
    """Expand the asset template row at `row_idx` into one row per entry
    in `assets` (already condensed-or-not by `assets.asset_rows`), copying
    the template row's style (font, border, alignment, fill, number
    format — everything `_style` carries) onto each newly written row.

    With zero assets, the template row is blanked and, when `asset_notes`
    is provided, that free-form text is written into the row's last
    placeholder column (typically Comments) so existing templates surface
    it without any re-annotation.

    Leftover template content below the written rows is cleared, column
    by column, stopping at the first row that's already empty across all
    of the template row's columns.
    """
    template_cells = []
    for cell in ws[row_idx]:
        if isinstance(cell, MergedCell):
            continue
        if cell.value is not None:
            template_cells.append((cell.column, cell.value, cell))
    if not template_cells:
        return

    placeholder_cols = [c for c, v, _ in template_cells
                        if isinstance(v, str) and "{{" in v]
    all_cols = [c for c, _, _ in template_cells]

    def write_row(target_idx, asset_ctx):
        for col, tmpl_value, tmpl_cell in template_cells:
            target = ws.cell(row=target_idx, column=col)
            if isinstance(target, MergedCell):
                continue
            local_ctx = dict(context)
            local_ctx["asset"] = asset_ctx
            target.value = substitute(tmpl_value, local_ctx)
            if target_idx != row_idx:
                target._style = copy_module.copy(tmpl_cell._style)

    if assets:
        for i, asset in enumerate(assets):
            write_row(row_idx + i, asset)
        first_clear = row_idx + len(assets)
    else:
        write_row(row_idx, {})
        if asset_notes and placeholder_cols:
            ws.cell(row=row_idx, column=placeholder_cols[-1]).value = asset_notes
        first_clear = row_idx + 1

    r = first_clear
    while r <= ws.max_row:
        row_has_content = False
        for col in all_cols:
            cell = ws.cell(row=r, column=col)
            if not isinstance(cell, MergedCell) and cell.value not in (None, ""):
                row_has_content = True
                cell.value = None
        if not row_has_content:
            break
        r += 1
