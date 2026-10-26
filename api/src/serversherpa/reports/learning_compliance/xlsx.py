"""Learning compliance workbook: Summary and Compliance (one row per person
and required document)."""

from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook

from serversherpa.reports.learning_compliance.gather import ComplianceData
from serversherpa.reports.timesheet.xlsx import (
    _BOLD,
    _TITLE,
    _as_text,
    _autofit,
    _bold_row,
    _Sheet,
)

_DATE_FMT = "yyyy-mm-dd"
_STAMP_FMT = "yyyy-mm-dd hh:mm"


def _naive_utc(at: datetime | None) -> datetime | None:
    """Excel has no time zones: store the UTC wall time (the Summary says so)."""
    return None if at is None else at.astimezone(UTC).replace(tzinfo=None)


def _summary(wb: Workbook, data: ComplianceData, generated_at: datetime) -> None:
    ws = wb.active
    ws.title = "Summary"
    sh = _Sheet(ws)
    sh.append(["Learning compliance"])
    ws["A1"].font = _TITLE
    sh.append([])
    for label, value in (
            ("Scope", data.initiative_name or "All active initiatives and the general set"),
            ("Generated", f"{generated_at.astimezone(UTC):%Y-%m-%d %H:%M} UTC"),
            ("Times", "Signed at is in UTC.")):
        sh.cell(sh.append([label, value]), 1).font = _BOLD
    sh.append([])
    for label, value in (
            ("People", data.people), ("Required documents", len(data.rows)),
            ("Valid", data.count("valid")),
            ("Expires during initiative", data.count("expires_during")),
            ("Missing", data.count("missing"))):
        sh.cell(sh.append([label, value]), 1).font = _BOLD
    _as_text(ws)
    _autofit(ws)


def _compliance(wb: Workbook, data: ComplianceData) -> None:
    ws = wb.create_sheet("Compliance")
    sh = _Sheet(ws)
    everyone = data.initiative_id is None
    headers = (["Person", "Document"] + (["Initiative"] if everyone else [])
               + ["Status", "Version signed", "Language", "Signed at", "Expires on",
                  "Overrides"])
    _bold_row(ws, sh.append(headers), len(headers))
    stamp_col = headers.index("Signed at") + 1
    for r in data.rows:
        row = sh.append([r.person, r.document] + ([r.initiative] if everyone else [])
                        + [r.status, r.version, r.language or "", _naive_utc(r.signed_at),
                           r.expires_on, r.overrides])
        sh.cell(row, stamp_col).number_format = _STAMP_FMT
        sh.cell(row, stamp_col + 1).number_format = _DATE_FMT
    ws.freeze_panes = "A2"
    _as_text(ws)
    _autofit(ws)


def build_workbook(data: ComplianceData, generated_at: datetime) -> bytes:
    wb = Workbook()
    _summary(wb, data, generated_at)
    _compliance(wb, data)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
