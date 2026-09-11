"""Move Scan History XLSX — Overview + Scan History sheets. Port of V2's
`generate_xlsx()` (api/reports/scan_history_report.py) with three extra
Overview rows (Scheduled Start, Source, Destination — V2's PDF had them,
the XLSX didn't — see the design spec's "Deliberate differences") and
timezone-aware timestamps (V2 stored/rendered naive local time).
"""

from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font

from serversherpa.reports.move_scan_history.timefmt import timezone_label, zone_abbrev
from serversherpa.reports.move_scan_history.gather import ScanHistoryData, StatusCol

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_OVERVIEW_DT_FMT = "%m/%d/%Y, %I:%M:%S %p"
_CELL_DT_FMT = "%m/%d/%Y %I:%M:%S %p"
_DATE_FMT = "%m/%d/%Y"

_BOLD = Font(bold=True)


def _fmt(dt: datetime | None, tz: ZoneInfo, fmt: str) -> str:
    if dt is None:
        return ""
    return dt.astimezone(tz).strftime(fmt)


def _autofit(ws) -> None:
    for col in ws.columns:
        max_length = 0
        for cell in col:
            value = str(cell.value) if cell.value is not None else ""
            max_length = max(max_length, len(value))
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 30)


def _build_overview(wb: Workbook, data: ScanHistoryData, columns: list[StatusCol],
                    generated_at: datetime, tz: ZoneInfo) -> None:
    ws = wb.active
    ws.title = "Overview"

    scheduled_start = (_fmt(data.scheduled_start, tz, _DATE_FMT)
                       if data.scheduled_start else "Not scheduled")
    block_rows = [
        ["Move Name", data.name],
        ["Client", data.client_name or "N/A"],
        ["Scheduled Start", scheduled_start],
        ["Source", data.source_name or "N/A"],
        ["Destination", data.destination_name or "N/A"],
        ["Date Generated", f"{_fmt(generated_at, tz, _OVERVIEW_DT_FMT)} {zone_abbrev(tz, generated_at)}"],
        # Named once here — every timestamp on both sheets is in this zone.
        ["Time Zone", timezone_label(tz, generated_at)],
        ["Total Assets", data.total_assets],
        ["Completion", f"{data.completion_pct}%"],
    ]
    for row in block_rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=1, max_row=len(block_rows), max_col=1):
        for cell in row:
            cell.font = _BOLD
    ws.append([])  # blank separator row

    if not data.assets:
        ws.append(["No assets found in this move"])
        return

    headers = ["Asset ID", "Serial Number", "Asset Name"] + [c.label for c in columns]
    ws.append(headers)
    header_row = ws.max_row
    for col_idx in range(1, len(headers) + 1):
        ws.cell(row=header_row, column=col_idx).font = _BOLD

    for asset in data.assets:
        by_status = {h.status_key: h.at for h in data.scan_progress.get(asset.asset_id, ())}
        row_data = [asset.asset_id, asset.serial_number, asset.name]
        for col in columns:
            at = by_status.get(col.key)
            row_data.append(_fmt(at, tz, _CELL_DT_FMT) if at else "")
        ws.append(row_data)

    _autofit(ws)


def _build_scan_history(wb: Workbook, data: ScanHistoryData, tz: ZoneInfo,
                        generated_at: datetime) -> None:
    ws = wb.create_sheet("Scan History")
    ws.append(["Asset ID", "Serial Number", "Asset Name", "Status Name", f"Timestamp ({zone_abbrev(tz, generated_at)})"])
    for col_idx in range(1, 6):
        ws.cell(row=1, column=col_idx).font = _BOLD

    asset_lookup = {a.asset_id: a for a in data.assets}
    flat: list[tuple[int, str, str, str, datetime]] = []
    for asset_id, hits in data.scan_progress.items():
        asset = asset_lookup.get(asset_id)
        serial = asset.serial_number if asset else ""
        name = asset.name if asset else ""
        for hit in hits:
            flat.append((asset_id, serial, name, hit.status_label, hit.at))

    flat.sort(key=lambda row: (row[0], row[4]))

    for asset_id, serial, name, status_label, at in flat:
        ws.append([asset_id, serial, name, status_label, _fmt(at, tz, _CELL_DT_FMT)])

    if not flat:
        ws.append(["No scan history found"])

    _autofit(ws)


def build_workbook(data: ScanHistoryData, columns: list[StatusCol],
                    generated_at: datetime, tz: ZoneInfo) -> bytes:
    wb = Workbook()
    _build_overview(wb, data, columns, generated_at, tz)
    _build_scan_history(wb, data, tz, generated_at)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
